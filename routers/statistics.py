# routers/statistics.py
import json
import csv
import io
import pytz
import pyzipper
import openpyxl
from openpyxl.styles import Alignment, PatternFill, Font, Border, Side
from datetime import datetime
from typing import Optional
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import func, or_, text

from app import models
from app.database import get_db, get_logs_db
from auth import get_current_user
from routers.activity_logger import record_neon_activity

router = APIRouter(prefix="/api/v1", tags=["Statistics & Operations"])

REGIONAL_HIERARCHY = {
    "KMP NORTH": ["KMP NORTH HEADQUARTERS", "KMP NORTH", "KAWEMPE", "KAKIRI", "KASANGATI", "MATUGGA", "NANSANA", "OLD KAMPALA", "WAKISO", "WANDEGEYA"],
    "KMP EAST": ["KMP EAST HEADQUARTERS", "KMP EAST", "JINJA ROAD", "KIRA", "KIRA DIV", "KIRA ROAD", "MUKONO", "NAGGALAMA", "SEETA"],
    "KMP SOUTH": ["KMP SOUTH HEADQUARTERS", "KMP SOUTH", "NATEETE", "CPS KAMPALA", "PARLIAMENT", "ENTEBBE", "KABALAGALA", "KAJJANSI", "KASENYI", "KATWE", "KYENGERA", "NSANGI"],
    "KMP HEADQUARTERS": ["KMP HEADQUARTERS", "KMP CID", "KMP TRAFFIC", "KMP ICT", "KMP FLYING SQUAD", "KMP CRIME INTELLIGENCE"],
    "POLICE HEADQUARTERS": ["NAGURU", "OPERATIONS", "CRIME INTELLIGENCE", "CID", "LOGISTICS & ENGINEERING", "ICT", "CT", "FIRE & RESCUE"]
}

def safe_int(val) -> int:
    try:
        return int(float(val)) if val else 0
    except (ValueError, TypeError):
        return 0

def get_model_safe(*names):
    for name in names:
        if hasattr(models, name):
            return getattr(models, name)
    target_slugs = [n.lower() for n in names]
    for attr_name, cls in models.__dict__.items():
        if isinstance(cls, type) and hasattr(cls, '__tablename__'):
            if cls.__tablename__.lower() in target_slugs or attr_name.lower() in target_slugs:
                return cls
    return None

def clean_model_dict(obj):
    if not obj:
        return {}
    if isinstance(obj, dict):
        return obj
    d = obj.__dict__.copy()
    d.pop('_sa_instance_state', None)
    clean = {}
    for k, v in d.items():
        if hasattr(v, 'isoformat'):
            clean[k] = v.isoformat()
        else:
            clean[k] = v
    return clean

def get_officer_signature(user):
    if not user:
        return "UNKNOWN COMMANDER"
    fnum = (user.fnum or "").strip()
    rank = (user.rank or "").strip()
    name = (user.name or "").strip()
    return f"{fnum} {rank} {name}".strip().upper()

def apply_opsec_scope(current_user, query, ModelClass):
    if not ModelClass or not current_user:
        return query.filter(text("1=0"))
    
    user_role = str(current_user.role).strip().upper() if current_user.role else ""
    user_pos = str(current_user.position).strip().upper() if current_user.position else ""
    user_reg = str(current_user.region).strip().upper() if current_user.region else ""
    user_stn = str(current_user.station).strip().upper() if current_user.station else ""

    perms = current_user.permissions or {}
    if isinstance(perms, str):
        try:
            perms = json.loads(perms)
        except Exception:
            perms = {}

    is_absolute_global = (
        user_role in ["SUPER_ADMIN", "ADMIN", "ASSISTANT_SUPER_ADMIN", "RPC", "DEPUTY COMMANDER", "SUPERVISOR"] or
        "KMP COMMANDER" in user_pos or
        "DEPUTY KMP COMMANDER" in user_pos or
        "KMP ADMIN" in user_pos or
        perms.get("view_global_roster") is True or
        perms.get("global_observer") is True or
        user_reg in ["POLICE HEADQUARTERS", "KMP HEADQUARTERS"]
    )

    is_kmp_sys_mgr = (
        user_role == "SYSTEM_MANAGER" and
        user_reg in ["KMP HEADQUARTERS", "POLICE HEADQUARTERS"] and
        "KMP" in user_pos
    )

    is_regional_command = (
        user_role in ["RPC", "DEPUTY_RPC", "SYSTEM_MANAGER", "ASSISTANT_SYSTEM_MANAGER", "REGIONAL_ADMIN", "ASSISTANT_REGIONAL_ADMIN", "DIVISION_ADMIN"] or
        "HR" in user_pos
    )

    if is_absolute_global or is_kmp_sys_mgr:
        return query
        
    elif is_regional_command or user_reg in REGIONAL_HIERARCHY:
        conds = []
        if hasattr(ModelClass, 'region'):
            conds.append(func.upper(ModelClass.region) == user_reg)
        
        if hasattr(ModelClass, 'station') and user_reg in REGIONAL_HIERARCHY:
            expanded_stns = set()
            for s in REGIONAL_HIERARCHY[user_reg]:
                expanded_stns.add(s)
                expanded_stns.add(s.replace(' HEADQUARTERS', '').replace(' HQ', ''))
                expanded_stns.add(s + ' HEADQUARTERS')
                expanded_stns.add(s + ' HQ')
            conds.append(func.upper(ModelClass.station).in_(list(expanded_stns)))
            
        if conds:
            return query.filter(or_(*conds))
        return query.filter(text("1=0"))
        
    elif hasattr(ModelClass, 'station'):
        clean_stn = user_stn.replace(" HEADQUARTERS", "").replace(" HQ", "")
        return query.filter(
            or_(
                func.upper(ModelClass.station) == user_stn,
                func.upper(ModelClass.station) == f"{clean_stn} HEADQUARTERS",
                func.upper(ModelClass.station) == f"{clean_stn} HQ"
            )
        )
        
    return query.filter(text("1=0"))

def build_filtered_query(query, ModelClass, region, station, start_date, end_date, search):
    clean_stn = ""
    
    if region and region.upper() not in ["ALL REGIONS", "ALL", "UNDEFINED", "NULL", ""]:
        if hasattr(ModelClass, 'region'):
            query = query.filter(func.upper(ModelClass.region) == region.upper())
            
    if station and station.upper() not in ["ALL STATIONS", "ALL STATIONS / DIVISIONS", "ALL", "UNDEFINED", "NULL", ""]:
        if hasattr(ModelClass, 'station'):
            clean_stn = station.upper().replace(' HEADQUARTERS', '').replace(' HQ', '')
            query = query.filter(
                or_(
                    func.upper(ModelClass.station) == station.upper(),
                    func.upper(ModelClass.station) == clean_stn,
                    func.upper(ModelClass.station) == f"{clean_stn} HEADQUARTERS"
                )
            )

    invalid_date_strings = ["ALL TIME", "ALL", "UNDEFINED", "NULL", ""]
    try:
        if start_date and str(start_date).strip().upper() not in invalid_date_strings and hasattr(ModelClass, 'date'):
            query = query.filter(func.cast(ModelClass.date, text("VARCHAR")) >= start_date)
    
        if end_date and str(end_date).strip().upper() not in invalid_date_strings and hasattr(ModelClass, 'date'):
            query = query.filter(func.cast(ModelClass.date, text("VARCHAR")) <= end_date)
    except Exception:
        pass

    if search and str(search).strip().upper() not in ["UNDEFINED", "NULL", ""]:
        term = f"%{search.strip().upper()}%"
        search_conds = []
        if hasattr(ModelClass, 'station'): search_conds.append(ModelClass.station.ilike(term))
        if hasattr(ModelClass, 'date'): search_conds.append(func.cast(ModelClass.date, text("VARCHAR")).ilike(term))
        if hasattr(ModelClass, 'id'): search_conds.append(func.cast(ModelClass.id, text("VARCHAR")).ilike(term))
        if hasattr(ModelClass, 'sn'): search_conds.append(func.cast(ModelClass.sn, text("VARCHAR")).ilike(term))
        if hasattr(ModelClass, 'agric_crime_report'): search_conds.append(ModelClass.agric_crime_report.ilike(term))
        if hasattr(ModelClass, 'nature_of_accident'): search_conds.append(ModelClass.nature_of_accident.ilike(term))
        if hasattr(ModelClass, 'nature_of_fire'): search_conds.append(ModelClass.nature_of_fire.ilike(term))
        if search_conds:
            query = query.filter(or_(*search_conds))
            
    return query


# ==========================================
# 1. DISRUPTIVE OPS STATS ENDPOINTS (`operational_statistics`)
# ==========================================
@router.get("/stats")
def get_stats(
    region: Optional[str] = Query(default=None),
    station: Optional[str] = Query(default=None),
    start_date: Optional[str] = Query(default=None),
    end_date: Optional[str] = Query(default=None),
    search: Optional[str] = Query(default=None),
    limit: int = Query(default=300),
    db: Session = Depends(get_db),
    current_user: models.Users = Depends(get_current_user)
):
    try:
        StatsModel = get_model_safe('operational_statistics', 'OperationalStatistics', 'OperationalStats', 'operational_stats', 'Stats', 'stats')
        if not StatsModel:
            sql_query = text("SELECT id, date, region, station, arrested, given_bond, cautioned, pending_court, taken_to_court, released, remanded, convicted FROM operational_statistics ORDER BY id DESC LIMIT :limit")
            raw_res = db.execute(sql_query, {"limit": limit}).fetchall()
            return [{
                "id": row[0],
                "date": str(row[1] or ''),
                "region": str(row[2] or ''),
                "station": str(row[3] or ''),
                "arrested": safe_int(row[4]),
                "given_bond": safe_int(row[5]),
                "cautioned": safe_int(row[6]),
                "pending_court": safe_int(row[7]),
                "taken_to_court": safe_int(row[8]),
                "released": safe_int(row[9]),
                "remanded": safe_int(row[10]),
                "convicted": safe_int(row[11])
            } for row in raw_res]
        
        query = db.query(StatsModel)
        query = apply_opsec_scope(current_user, query, StatsModel)
        query = build_filtered_query(query, StatsModel, region, station, start_date, end_date, search)
        
        pk_col = getattr(StatsModel, 'id', getattr(StatsModel, 'sn', None))
        if pk_col is not None and hasattr(pk_col, 'desc'):
            records = query.order_by(pk_col.desc()).limit(limit).all()
        else:
            records = query.limit(limit).all()
        
        results = []
        for r in records:
            results.append({
                "id": getattr(r, 'id', getattr(r, 'sn', 0)),
                "date": str(getattr(r, 'date', '') or ''),
                "region": str(getattr(r, 'region', '') or ''),
                "station": str(getattr(r, 'station', '') or ''),
                "arrested": safe_int(getattr(r, 'arrested', 0)),
                "given_bond": safe_int(getattr(r, 'given_bond', 0)),
                "cautioned": safe_int(getattr(r, 'cautioned', 0)),
                "pending_court": safe_int(getattr(r, 'pending_court', 0)),
                "taken_to_court": safe_int(getattr(r, 'taken_to_court', 0)),
                "released": safe_int(getattr(r, 'released', 0)),
                "remanded": safe_int(getattr(r, 'remanded', 0)),
                "convicted": safe_int(getattr(r, 'convicted', 0))
            })
        return results
    except Exception as e:
        db.rollback()
        print(f"Stats Endpoint Critical Exception: {e}")
        return []

@router.get("/stats/export")
def export_stats_secure_zip(
    region: Optional[str] = Query(default=None),
    station: Optional[str] = Query(default=None),
    start_date: Optional[str] = Query(default=None),
    end_date: Optional[str] = Query(default=None),
    search: Optional[str] = Query(default=None),
    db: Session = Depends(get_db),
    logs_db: Session = Depends(get_logs_db),
    current_user: models.Users = Depends(get_current_user)
):
    StatsModel = get_model_safe('operational_statistics', 'OperationalStatistics', 'OperationalStats', 'operational_stats', 'Stats', 'stats')
    if not StatsModel: raise HTTPException(status_code=500, detail="Model not found")

    query = db.query(StatsModel)
    query = apply_opsec_scope(current_user, query, StatsModel)
    query = build_filtered_query(query, StatsModel, region, station, start_date, end_date, search)
    pk_col = getattr(StatsModel, 'id', getattr(StatsModel, 'sn', None))
    records = query.order_by(pk_col.desc()).all() if pk_col is not None else query.all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Disruptive OPS Statistics"
    ws.views.sheetView[0].showGridLines = True

    eat_tz = pytz.timezone("Africa/Nairobi")
    eat_time = datetime.now(eat_tz).replace(tzinfo=None)
    
    officer_name = (current_user.name or "UNKNOWN").strip().upper()
    officer_fnum = (current_user.fnum or "UNKNOWN").strip().upper()
    officer_rank = (current_user.rank or "").strip().upper()
    officer_station = (current_user.station or "HQ").strip().upper()
    timestamp_str = eat_time.strftime("%Y-%m-%d %H:%M:%S EAT")

    wb.properties.creator = f"{officer_fnum} {officer_rank} {officer_name}"
    wb.properties.lastModifiedBy = f"KMP-CSDMS SECURITY AUDIT ENGINE [{officer_station}]"
    wb.properties.title = "Official Law Enforcement Restricted Disruptive OPS Statistics Export"

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell_font = Font(name="Calibri", size=10)
    thin_border = Border(left=Side(style='thin', color='D1D5DB'), right=Side(style='thin', color='D1D5DB'), top=Side(style='thin', color='D1D5DB'), bottom=Side(style='thin', color='D1D5DB'))

    ws.merge_cells("A1:L1")
    banner_cell = ws.cell(row=1, column=1)
    banner_cell.value = f"RESTRICTED LAW ENFORCEMENT RECORD — FORENSICALLY STAMPED TO OFFICER: {officer_fnum} {officer_rank} {officer_name} ({officer_station}) — TIMESTAMP: {timestamp_str}"
    banner_cell.font = Font(name="Calibri", size=9, bold=True, color="991B1B")
    banner_cell.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    banner_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 22

    ws.cell(row=2, column=1, value="UGANDA POLICE FORCE - DISRUPTIVE OPS STATISTICS LEDGER").font = Font(bold=True, size=11)
    ws.row_dimensions[2].height = 18

    headers = ["ID", "Date", "Region", "Station", "Arrested", "Given Bond", "Cautioned", "Pending Court", "Taken to Court", "Released by Court", "Remanded", "Convicted"]
    ws.row_dimensions[4].height = 26
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    current_row = 5
    for r in records:
        row_data = [
            getattr(r, 'id', getattr(r, 'sn', '')), getattr(r, 'date', ''), getattr(r, 'region', ''), getattr(r, 'station', ''),
            getattr(r, 'arrested', 0), getattr(r, 'given_bond', 0), getattr(r, 'cautioned', 0),
            getattr(r, 'pending_court', 0), getattr(r, 'taken_to_court', 0), getattr(r, 'released', 0),
            getattr(r, 'remanded', 0), getattr(r, 'convicted', 0)
        ]
        ws.append(row_data)
        ws.row_dimensions[current_row].height = 20
        for col_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.font = cell_font
            cell.border = thin_border
            if col_idx in [1, 2, 5, 6, 7, 8, 9, 10, 11, 12]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
        current_row += 1

    for col in ws.columns:
        col_letter = col[0].column_letter
        max_len = max([len(str(cell.value or '')) for cell in col], default=0)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 35)

    excel_stream = io.BytesIO()
    wb.save(excel_stream)
    excel_stream.seek(0)

    zip_stream = io.BytesIO()
    zip_password = str(current_user.fnum).strip().encode('utf-8')
    fnum_clean = str(current_user.fnum).replace('/', '_').upper()
    excel_filename = f"{fnum_clean}_Disruptive_OPS_Stats_{eat_time.strftime('%Y%m%d')}.xlsx"
    zip_filename = f"SECURE_DISRUPTIVE_OPS_STATS_{eat_time.strftime('%Y%m%d')}.zip"

    with pyzipper.AESZipFile(zip_stream, 'w', compression=pyzipper.ZIP_DEFLATED, encryption=pyzipper.WZ_AES) as zf:
        zf.setpassword(zip_password)
        zf.writestr(excel_filename, excel_stream.getvalue())

    zip_stream.seek(0)

    record_neon_activity(logs_db=logs_db, fnum=current_user.fnum, action_type="EXPORT", module="DISRUPTIVE_OPS_STATS", target_id="SECURE_ZIP_EXPORT", changes_summary=f"{current_user.fnum} exported Disruptive Ops Stats.")

    return StreamingResponse(
        zip_stream, media_type="application/zip",
        headers={'Content-Disposition': f'attachment; filename="{zip_filename}"', 'Access-Control-Expose-Headers': 'Content-Disposition'}
    )

@router.post("/stats")
def create_stat(data: dict, db: Session = Depends(get_db), logs_db: Session = Depends(get_logs_db), current_user: models.Users = Depends(get_current_user)):
    StatsModel = get_model_safe('operational_statistics', 'OperationalStatistics', 'OperationalStats', 'operational_stats', 'Stats', 'stats')
    if not StatsModel: raise HTTPException(status_code=500, detail="Model not configured.")
    try:
        data.pop('sn', None); data.pop('id', None)
        valid_cols = [c.key for c in StatsModel.__table__.columns]
        safe_data = {k: v for k, v in data.items() if k in valid_cols}
        new_record = StatsModel(**safe_data)
        db.add(new_record); db.commit(); db.refresh(new_record)

        rec_id = getattr(new_record, 'id', getattr(new_record, 'sn', '0'))
        record_neon_activity(logs_db=logs_db, fnum=current_user.fnum, action_type="CREATE", module="DISRUPTIVE_OPS_STATS", target_id=str(rec_id), changes_summary=f"{current_user.fnum} created Disruptive Ops Stats ID {rec_id}.")
        return clean_model_dict(new_record)
    except Exception as e:
        db.rollback(); logs_db.rollback(); raise HTTPException(status_code=500, detail=str(e))

@router.put("/stats/{record_id}")
def update_stat(record_id: int, data: dict, db: Session = Depends(get_db), logs_db: Session = Depends(get_logs_db), current_user: models.Users = Depends(get_current_user)):
    StatsModel = get_model_safe('operational_statistics', 'OperationalStatistics', 'OperationalStats', 'operational_stats', 'Stats', 'stats')
    if not StatsModel: raise HTTPException(status_code=500, detail="Model not configured.")
    try:
        pk_col = getattr(StatsModel, 'id', getattr(StatsModel, 'sn', None))
        record = db.query(StatsModel).filter(pk_col == record_id).first()
        if not record: raise HTTPException(status_code=404, detail="Record not found.")
        data.pop('id', None); data.pop('sn', None)
        for key, value in data.items():
            if hasattr(record, key): setattr(record, key, value)
        db.commit(); db.refresh(record)

        record_neon_activity(logs_db=logs_db, fnum=current_user.fnum, action_type="UPDATE", module="DISRUPTIVE_OPS_STATS", target_id=str(record_id), changes_summary=f"{current_user.fnum} updated Disruptive Ops Stats ID {record_id}.")
        return clean_model_dict(record)
    except Exception as e:
        db.rollback(); logs_db.rollback(); raise HTTPException(status_code=500, detail=str(e))


# ==========================================
# 2. AGRICULTURAL CRIME STATS (`agric_crimes_statistics` - 8-FIELD DISRUPTIVE OPS MATRIX)
# ==========================================
@router.get("/agric-stats")
def get_agric_stats(
    region: Optional[str] = Query(default=None), station: Optional[str] = Query(default=None),
    start_date: Optional[str] = Query(default=None), end_date: Optional[str] = Query(default=None),
    search: Optional[str] = Query(default=None), limit: int = Query(default=300),
    db: Session = Depends(get_db), current_user: models.Users = Depends(get_current_user)
):
    try:
        AgricStatsModel = get_model_safe('agric_crimes_statistics', 'Agric_Crimes_Statistics', 'AgricCrimesStatistics', 'AgricStats', 'agric_stats', 'Agric_Stats')
        if not AgricStatsModel:
            sql_query = text("SELECT id, date, region, station, arrested, given_bond, cautioned, pending_court, taken_to_court, released, remanded, convicted FROM agric_crimes_statistics ORDER BY id DESC LIMIT :limit")
            raw_res = db.execute(sql_query, {"limit": limit}).fetchall()
            return [{
                "id": row[0],
                "date": str(row[1] or ''),
                "region": str(row[2] or ''),
                "station": str(row[3] or ''),
                "arrested": safe_int(row[4]),
                "given_bond": safe_int(row[5]),
                "cautioned": safe_int(row[6]),
                "pending_court": safe_int(row[7]),
                "taken_to_court": safe_int(row[8]),
                "released": safe_int(row[9]),
                "remanded": safe_int(row[10]),
                "convicted": safe_int(row[11])
            } for row in raw_res]

        query = db.query(AgricStatsModel)
        query = apply_opsec_scope(current_user, query, AgricStatsModel)
        query = build_filtered_query(query, AgricStatsModel, region, station, start_date, end_date, search)
        
        pk_col = getattr(AgricStatsModel, 'id', getattr(AgricStatsModel, 'sn', None))
        records = query.order_by(pk_col.desc()).limit(limit).all() if pk_col is not None and hasattr(pk_col, 'desc') else query.limit(limit).all()
        
        return [{
            "id": getattr(r, 'id', getattr(r, 'sn', 0)),
            "date": str(getattr(r, 'date', '') or ''),
            "region": str(getattr(r, 'region', '') or ''),
            "station": str(getattr(r, 'station', '') or ''),
            "arrested": safe_int(getattr(r, 'arrested', 0)),
            "given_bond": safe_int(getattr(r, 'given_bond', 0)),
            "cautioned": safe_int(getattr(r, 'cautioned', 0)),
            "pending_court": safe_int(getattr(r, 'pending_court', 0)),
            "taken_to_court": safe_int(getattr(r, 'taken_to_court', 0)),
            "released": safe_int(getattr(r, 'released', 0)),
            "remanded": safe_int(getattr(r, 'remanded', 0)),
            "convicted": safe_int(getattr(r, 'convicted', 0))
        } for r in records]
    except Exception as e:
        db.rollback()
        print(f"Agric Stats Endpoint Error: {e}")
        return []

@router.get("/agric-stats/export")
def export_agric_stats_secure_zip(
    region: Optional[str] = Query(default=None), station: Optional[str] = Query(default=None),
    start_date: Optional[str] = Query(default=None), end_date: Optional[str] = Query(default=None),
    search: Optional[str] = Query(default=None), db: Session = Depends(get_db),
    logs_db: Session = Depends(get_logs_db), current_user: models.Users = Depends(get_current_user)
):
    AgricStatsModel = get_model_safe('agric_crimes_statistics', 'Agric_Crimes_Statistics', 'AgricCrimesStatistics', 'AgricStats', 'agric_stats', 'Agric_Stats')
    if not AgricStatsModel: raise HTTPException(status_code=500, detail="Model not found")

    query = db.query(AgricStatsModel)
    query = apply_opsec_scope(current_user, query, AgricStatsModel)
    query = build_filtered_query(query, AgricStatsModel, region, station, start_date, end_date, search)
    pk_col = getattr(AgricStatsModel, 'id', getattr(AgricStatsModel, 'sn', None))
    records = query.order_by(pk_col.desc()).all() if pk_col is not None else query.all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Agricultural Crimes Stats"
    ws.views.sheetView[0].showGridLines = True

    eat_tz = pytz.timezone("Africa/Nairobi")
    eat_time = datetime.now(eat_tz).replace(tzinfo=None)
    
    officer_name = (current_user.name or "UNKNOWN").strip().upper()
    officer_fnum = (current_user.fnum or "UNKNOWN").strip().upper()
    officer_rank = (current_user.rank or "").strip().upper()
    officer_station = (current_user.station or "HQ").strip().upper()
    timestamp_str = eat_time.strftime("%Y-%m-%d %H:%M:%S EAT")

    wb.properties.creator = f"{officer_fnum} {officer_rank} {officer_name}"
    wb.properties.lastModifiedBy = f"KMP-CSDMS SECURITY AUDIT ENGINE [{officer_station}]"
    wb.properties.title = "Official Law Enforcement Restricted Agricultural Statistics Export"

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="047857", end_color="047857", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell_font = Font(name="Calibri", size=10)
    thin_border = Border(left=Side(style='thin', color='D1D5DB'), right=Side(style='thin', color='D1D5DB'), top=Side(style='thin', color='D1D5DB'), bottom=Side(style='thin', color='D1D5DB'))

    ws.merge_cells("A1:L1")
    banner_cell = ws.cell(row=1, column=1)
    banner_cell.value = f"RESTRICTED LAW ENFORCEMENT RECORD — FORENSICALLY STAMPED TO OFFICER: {officer_fnum} {officer_rank} {officer_name} ({officer_station}) — TIMESTAMP: {timestamp_str}"
    banner_cell.font = Font(name="Calibri", size=9, bold=True, color="991B1B")
    banner_cell.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    banner_cell.alignment = Alignment(horizontal="center", vertical="center")
    
    ws.cell(row=2, column=1, value="UGANDA POLICE FORCE - AGRICULTURAL CRIMES STATISTICS LEDGER").font = Font(bold=True, size=11)
    ws.row_dimensions[2].height = 18

    headers = ["ID", "Date", "Region", "Station", "Arrested", "Given Bond", "Cautioned", "Pending Court", "Taken to Court", "Released by Court", "Remanded", "Convicted"]
    ws.row_dimensions[4].height = 26
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    current_row = 5
    for r in records:
        row_data = [
            getattr(r, 'id', getattr(r, 'sn', '')), getattr(r, 'date', ''), getattr(r, 'region', ''), getattr(r, 'station', ''),
            getattr(r, 'arrested', 0), getattr(r, 'given_bond', 0), getattr(r, 'cautioned', 0),
            getattr(r, 'pending_court', 0), getattr(r, 'taken_to_court', 0), getattr(r, 'released', 0),
            getattr(r, 'remanded', 0), getattr(r, 'convicted', 0)
        ]
        ws.append(row_data)
        for col_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.font = cell_font
            cell.border = thin_border
            if col_idx in [1, 2, 5, 6, 7, 8, 9, 10, 11, 12]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
        current_row += 1

    for col in ws.columns:
        col_letter = col[0].column_letter
        max_len = max([len(str(cell.value or '')) for cell in col], default=0)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 35)

    excel_stream = io.BytesIO()
    wb.save(excel_stream)
    excel_stream.seek(0)

    zip_stream = io.BytesIO()
    zip_password = str(current_user.fnum).strip().encode('utf-8')
    fnum_clean = str(current_user.fnum).replace('/', '_').upper()
    excel_filename = f"{fnum_clean}_Agric_Stats_{eat_time.strftime('%Y%m%d')}.xlsx"
    zip_filename = f"SECURE_AGRIC_STATS_{eat_time.strftime('%Y%m%d')}.zip"

    with pyzipper.AESZipFile(zip_stream, 'w', compression=pyzipper.ZIP_DEFLATED, encryption=pyzipper.WZ_AES) as zf:
        zf.setpassword(zip_password)
        zf.writestr(excel_filename, excel_stream.getvalue())
    zip_stream.seek(0)

    record_neon_activity(
        logs_db=logs_db, fnum=current_user.fnum, action_type="EXPORT", module="AGRIC_STATS", target_id="SECURE_ZIP_EXPORT",
        changes_summary=f"{current_user.fnum} securely exported AES-256 encrypted Agricultural Statistics archive."
    )

    return StreamingResponse(
        zip_stream, media_type="application/zip",
        headers={'Content-Disposition': f'attachment; filename="{zip_filename}"', 'Access-Control-Expose-Headers': 'Content-Disposition'}
    )

@router.post("/agric-stats")
def create_agric_stat(data: dict, db: Session = Depends(get_db), logs_db: Session = Depends(get_logs_db), current_user: models.Users = Depends(get_current_user)):
    AgricStatsModel = get_model_safe('agric_crimes_statistics', 'Agric_Crimes_Statistics', 'AgricCrimesStatistics', 'AgricStats', 'agric_stats', 'Agric_Stats')
    if not AgricStatsModel: raise HTTPException(status_code=500, detail="Model not configured.")
    try:
        data.pop('sn', None); data.pop('id', None)
        valid_cols = [c.key for c in AgricStatsModel.__table__.columns]
        safe_data = {k: v for k, v in data.items() if k in valid_cols}
        new_record = AgricStatsModel(**safe_data)
        db.add(new_record); db.commit(); db.refresh(new_record)

        rec_id = getattr(new_record, 'id', getattr(new_record, 'sn', '0'))
        record_neon_activity(logs_db=logs_db, fnum=current_user.fnum, action_type="CREATE", module="AGRIC_STATS", target_id=str(rec_id), changes_summary=f"{current_user.fnum} created Agricultural Statistics ID {rec_id}.")
        return clean_model_dict(new_record)
    except Exception as e:
        db.rollback(); logs_db.rollback(); raise HTTPException(status_code=500, detail=str(e))

@router.put("/agric-stats/{record_id}")
def update_agric_stat(record_id: int, data: dict, db: Session = Depends(get_db), logs_db: Session = Depends(get_logs_db), current_user: models.Users = Depends(get_current_user)):
    AgricStatsModel = get_model_safe('agric_crimes_statistics', 'Agric_Crimes_Statistics', 'AgricCrimesStatistics', 'AgricStats', 'agric_stats', 'Agric_Stats')
    if not AgricStatsModel: raise HTTPException(status_code=500, detail="Model not configured.")
    try:
        pk_col = getattr(AgricStatsModel, 'id', getattr(AgricStatsModel, 'sn', None))
        record = db.query(AgricStatsModel).filter(pk_col == record_id).first()
        if not record: raise HTTPException(status_code=404, detail="Record not found.")
        data.pop('id', None); data.pop('sn', None)
        for key, value in data.items():
            if hasattr(record, key): setattr(record, key, value)
        db.commit(); db.refresh(record)

        record_neon_activity(logs_db=logs_db, fnum=current_user.fnum, action_type="UPDATE", module="AGRIC_STATS", target_id=str(record_id), changes_summary=f"{current_user.fnum} updated Agricultural Statistics ID {record_id}.")
        return clean_model_dict(record)
    except Exception as e:
        db.rollback(); logs_db.rollback(); raise HTTPException(status_code=500, detail=str(e))


# ==========================================
# 3. AGRICULTURAL SUMMARY LEDGER ENDPOINTS (`agricultural_crime_summary`)
# ==========================================
@router.get("/agric-summary")
def get_agric_summary(
    region: Optional[str] = Query(default=None), station: Optional[str] = Query(default=None),
    start_date: Optional[str] = Query(default=None), end_date: Optional[str] = Query(default=None),
    search: Optional[str] = Query(default=None), limit: int = Query(default=300),
    db: Session = Depends(get_db), current_user: models.Users = Depends(get_current_user)
):
    try:
        SummaryModel = get_model_safe('Agricultural_Crime_Summary', 'AgriculturalCrimeSummary', 'agricultural_crime_summary')
        if not SummaryModel: return []
        query = db.query(SummaryModel)
        query = apply_opsec_scope(current_user, query, SummaryModel)
        query = build_filtered_query(query, SummaryModel, region, station, start_date, end_date, search)
        
        pk_col = getattr(SummaryModel, 'id', getattr(SummaryModel, 'sn', None))
        records = query.order_by(pk_col.desc()).limit(limit).all() if pk_col is not None and hasattr(pk_col, 'desc') else query.limit(limit).all()
        
        return [{
            "id": getattr(r, 'id', getattr(r, 'sn', 0)),
            "date": str(getattr(r, 'date', '') or ''),
            "region": str(getattr(r, 'region', '') or ''),
            "station": str(getattr(r, 'station', '') or ''),
            "agric_crime_report": str(getattr(r, 'agric_crime_report', '') or ''),
            "number_count": safe_int(getattr(r, 'number_count', 0)),
            "recoveries": safe_int(getattr(r, 'recoveries', 0)),
            "status": str(getattr(r, 'status', '') or '')
        } for r in records]
    except Exception as e:
        db.rollback()
        print(f"Agric Summary Endpoint Error: {e}")
        return []

@router.post("/agric-summary")
def create_agric_summary(data: dict, db: Session = Depends(get_db), logs_db: Session = Depends(get_logs_db), current_user: models.Users = Depends(get_current_user)):
    SummaryModel = get_model_safe('Agricultural_Crime_Summary', 'AgriculturalCrimeSummary', 'agricultural_crime_summary')
    if not SummaryModel: raise HTTPException(status_code=500, detail="Model not configured.")
    try:
        data.pop('sn', None); data.pop('id', None)
        valid_cols = [c.key for c in SummaryModel.__table__.columns]
        safe_data = {k: v for k, v in data.items() if k in valid_cols}
        new_record = SummaryModel(**safe_data)
        db.add(new_record); db.commit(); db.refresh(new_record)

        rec_id = getattr(new_record, 'id', getattr(new_record, 'sn', '0'))
        record_neon_activity(logs_db=logs_db, fnum=current_user.fnum, action_type="CREATE", module="AGRIC_SUMMARY", target_id=str(rec_id), changes_summary=f"{current_user.fnum} created Agric Summary ID {rec_id}.")
        return clean_model_dict(new_record)
    except Exception as e:
        db.rollback(); logs_db.rollback(); raise HTTPException(status_code=500, detail=str(e))


# ==========================================
# 4. TRAFFIC & ACCIDENTS STATISTICS ENDPOINTS (`accidents_statistics`)
# ==========================================
@router.get("/accidents")
def get_accidents_stats(
    region: Optional[str] = Query(default=None), station: Optional[str] = Query(default=None),
    start_date: Optional[str] = Query(default=None), end_date: Optional[str] = Query(default=None),
    search: Optional[str] = Query(default=None), limit: int = Query(default=300),
    db: Session = Depends(get_db), current_user: models.Users = Depends(get_current_user)
):
    try:
        AccidentsModel = get_model_safe('Accidents_Statistics', 'AccidentsStatistics', 'accidents_statistics')
        if not AccidentsModel: return []
        
        query = db.query(AccidentsModel)
        query = apply_opsec_scope(current_user, query, AccidentsModel)
        query = build_filtered_query(query, AccidentsModel, region, station, start_date, end_date, search)
        
        pk_col = getattr(AccidentsModel, 'id', getattr(AccidentsModel, 'sn', None))
        records = query.order_by(pk_col.desc()).limit(limit).all() if pk_col is not None and hasattr(pk_col, 'desc') else query.limit(limit).all()
        
        return [{
            "id": getattr(r, 'id', getattr(r, 'sn', 0)),
            "date": str(getattr(r, 'date', '') or ''),
            "region": str(getattr(r, 'region', '') or ''),
            "station": str(getattr(r, 'station', '') or ''),
            "nature_of_accident": str(getattr(r, 'nature_of_accident', '') or ''),
            "minor": safe_int(getattr(r, 'minor', 0)),
            "serious": safe_int(getattr(r, 'serious', 0)),
            "fatal": safe_int(getattr(r, 'fatal', 0)),
            "victims": safe_int(getattr(r, 'victims', 0)),
            "total": safe_int(getattr(r, 'total', 0)),
            "entered_by": str(getattr(r, 'entered_by', getattr(r, 'last_updated_by', '')) or ''),
            "created_at": r.created_at.isoformat() if hasattr(r, 'created_at') and r.created_at else ''
        } for r in records]
    except Exception as e:
        db.rollback()
        print(f"Accidents Endpoint Error: {e}")
        return []

@router.post("/accidents")
def create_accidents_stat(data: dict, db: Session = Depends(get_db), logs_db: Session = Depends(get_logs_db), current_user: models.Users = Depends(get_current_user)):
    AccidentsModel = get_model_safe('Accidents_Statistics', 'AccidentsStatistics', 'accidents_statistics')
    if not AccidentsModel: raise HTTPException(status_code=500, detail="Accidents model not configured in database.")
    try:
        data.pop('sn', None); data.pop('id', None)
        valid_cols = [c.key for c in AccidentsModel.__table__.columns]
        safe_data = {k: v for k, v in data.items() if k in valid_cols}
        
        safe_data['total'] = safe_int(safe_data.get('minor', 0)) + safe_int(safe_data.get('serious', 0)) + safe_int(safe_data.get('fatal', 0))
        safe_data['entered_by'] = get_officer_signature(current_user)
        
        new_record = AccidentsModel(**safe_data)
        db.add(new_record); db.commit(); db.refresh(new_record)

        rec_id = getattr(new_record, 'id', getattr(new_record, 'sn', '0'))
        record_neon_activity(
            logs_db=logs_db,
            fnum=current_user.fnum,
            action_type="CREATE",
            module="ACCIDENTS_STATS",
            target_id=str(rec_id),
            changes_summary=f"{current_user.fnum} recorded Traffic Accident stats for {new_record.station}."
        )

        clean = clean_model_dict(new_record)
        if hasattr(new_record, 'created_at') and new_record.created_at:
            clean['created_at'] = new_record.created_at.isoformat() if hasattr(new_record.created_at, 'isoformat') else str(new_record.created_at)
        return clean
    except Exception as e:
        db.rollback(); logs_db.rollback(); raise HTTPException(status_code=500, detail=str(e))

@router.put("/accidents/{record_id}")
def update_accidents_stat(record_id: int, data: dict, db: Session = Depends(get_db), logs_db: Session = Depends(get_logs_db), current_user: models.Users = Depends(get_current_user)):
    AccidentsModel = get_model_safe('Accidents_Statistics', 'AccidentsStatistics', 'accidents_statistics')
    if not AccidentsModel: raise HTTPException(status_code=500, detail="Accidents model not configured.")
    try:
        pk_col = getattr(AccidentsModel, 'id', getattr(AccidentsModel, 'sn', None))
        record = db.query(AccidentsModel).filter(pk_col == record_id).first()
        if not record: raise HTTPException(status_code=404, detail="Accident record not found.")
        
        data.pop('id', None); data.pop('sn', None)
        for key, value in data.items():
            if hasattr(record, key): setattr(record, key, value)
            
        record.total = safe_int(record.minor) + safe_int(record.serious) + safe_int(record.fatal)
        record.last_updated_by = get_officer_signature(current_user)
            
        db.commit(); db.refresh(record)

        record_neon_activity(
            logs_db=logs_db,
            fnum=current_user.fnum,
            action_type="UPDATE",
            module="ACCIDENTS_STATS",
            target_id=str(record_id),
            changes_summary=f"{current_user.fnum} updated Traffic Accident stats ID {record_id}."
        )

        clean = clean_model_dict(record)
        if hasattr(record, 'created_at') and record.created_at:
            clean['created_at'] = record.created_at.isoformat() if hasattr(record.created_at, 'isoformat') else str(record.created_at)
        return clean
    except Exception as e:
        db.rollback(); logs_db.rollback(); raise HTTPException(status_code=500, detail=str(e))

@router.get("/accidents/export")
def export_accidents_stats_secure_zip(
    region: Optional[str] = Query(default=None), station: Optional[str] = Query(default=None),
    start_date: Optional[str] = Query(default=None), end_date: Optional[str] = Query(default=None),
    search: Optional[str] = Query(default=None), db: Session = Depends(get_db),
    logs_db: Session = Depends(get_logs_db), current_user: models.Users = Depends(get_current_user)
):
    AccidentsModel = get_model_safe('Accidents_Statistics', 'AccidentsStatistics', 'accidents_statistics')
    if not AccidentsModel: raise HTTPException(status_code=500, detail="Model not found")

    query = db.query(AccidentsModel)
    query = apply_opsec_scope(current_user, query, AccidentsModel)
    query = build_filtered_query(query, AccidentsModel, region, station, start_date, end_date, search)
    pk_col = getattr(AccidentsModel, 'id', getattr(AccidentsModel, 'sn', None))
    records = query.order_by(pk_col.desc()).all() if pk_col is not None else query.all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Traffic & Accidents Stats"
    ws.views.sheetView[0].showGridLines = True

    eat_tz = pytz.timezone("Africa/Nairobi")
    eat_time = datetime.now(eat_tz).replace(tzinfo=None)
    
    officer_name = (current_user.name or "UNKNOWN").strip().upper()
    officer_fnum = (current_user.fnum or "UNKNOWN").strip().upper()
    officer_rank = (current_user.rank or "").strip().upper()
    officer_station = (current_user.station or "HQ").strip().upper()
    timestamp_str = eat_time.strftime("%Y-%m-%d %H:%M:%S EAT")

    wb.properties.creator = f"{officer_fnum} {officer_rank} {officer_name}"
    wb.properties.lastModifiedBy = f"KMP-CSDMS SECURITY AUDIT ENGINE [{officer_station}]"
    wb.properties.title = "Official Law Enforcement Restricted Traffic & Accidents Export"

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="991B1B", end_color="991B1B", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell_font = Font(name="Calibri", size=10)
    thin_border = Border(left=Side(style='thin', color='D1D5DB'), right=Side(style='thin', color='D1D5DB'), top=Side(style='thin', color='D1D5DB'), bottom=Side(style='thin', color='D1D5DB'))

    ws.merge_cells("A1:K1")
    banner_cell = ws.cell(row=1, column=1)
    banner_cell.value = f"RESTRICTED LAW ENFORCEMENT RECORD — FORENSICALLY STAMPED TO OFFICER: {officer_fnum} {officer_rank} {officer_name} ({officer_station}) — TIMESTAMP: {timestamp_str}"
    banner_cell.font = Font(name="Calibri", size=9, bold=True, color="991B1B")
    banner_cell.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    banner_cell.alignment = Alignment(horizontal="center", vertical="center")
    
    ws.cell(row=2, column=1, value="UGANDA POLICE FORCE - TRAFFIC & ACCIDENTS STATISTICS LEDGER").font = Font(bold=True, size=11)
    ws.row_dimensions[2].height = 18

    headers = ["ID", "Date", "Division/Station", "Nature of Accident", "Minor", "Serious", "Fatal", "Victims", "Total", "Created By", "Date Created"]
    ws.row_dimensions[4].height = 26
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    current_row = 5
    for r in records:
        created_str = r.created_at.strftime("%Y-%m-%d %H:%M") if hasattr(r, 'created_at') and r.created_at else ''
        row_data = [
            getattr(r, 'id', getattr(r, 'sn', '')), getattr(r, 'date', ''), getattr(r, 'station', ''), getattr(r, 'nature_of_accident', ''),
            getattr(r, 'minor', 0), getattr(r, 'serious', 0), getattr(r, 'fatal', 0), getattr(r, 'victims', 0), getattr(r, 'total', 0),
            getattr(r, 'entered_by', getattr(r, 'last_updated_by', '')), created_str
        ]
        ws.append(row_data)
        for col_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.font = cell_font
            cell.border = thin_border
            if col_idx in [1, 2, 5, 6, 7, 8, 9]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
        current_row += 1

    for col in ws.columns:
        col_letter = col[0].column_letter
        max_len = max([len(str(cell.value or '')) for cell in col], default=0)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 35)

    excel_stream = io.BytesIO()
    wb.save(excel_stream)
    excel_stream.seek(0)

    zip_stream = io.BytesIO()
    zip_password = str(current_user.fnum).strip().encode('utf-8')
    fnum_clean = str(current_user.fnum).replace('/', '_').upper()
    excel_filename = f"{fnum_clean}_Accidents_Stats_{eat_time.strftime('%Y%m%d')}.xlsx"
    zip_filename = f"SECURE_ACCIDENTS_STATS_{eat_time.strftime('%Y%m%d')}.zip"

    with pyzipper.AESZipFile(zip_stream, 'w', compression=pyzipper.ZIP_DEFLATED, encryption=pyzipper.WZ_AES) as zf:
        zf.setpassword(zip_password)
        zf.writestr(excel_filename, excel_stream.getvalue())
    zip_stream.seek(0)

    record_neon_activity(
        logs_db=logs_db, fnum=current_user.fnum, action_type="EXPORT", module="ACCIDENTS_STATS", target_id="SECURE_ZIP_EXPORT",
        changes_summary=f"{current_user.fnum} securely exported AES-256 encrypted Accidents Statistics archive."
    )

    return StreamingResponse(
        zip_stream, media_type="application/zip",
        headers={'Content-Disposition': f'attachment; filename="{zip_filename}"', 'Access-Control-Expose-Headers': 'Content-Disposition'}
    )


# ==========================================
# 5. FIRE & RESCUE INCIDENTS STATISTICS ENDPOINTS (`fire_statistics`)
# ==========================================
@router.get("/fires")
def get_fires_stats(
    region: Optional[str] = Query(default=None), station: Optional[str] = Query(default=None),
    start_date: Optional[str] = Query(default=None), end_date: Optional[str] = Query(default=None),
    search: Optional[str] = Query(default=None), limit: int = Query(default=300),
    db: Session = Depends(get_db), current_user: models.Users = Depends(get_current_user)
):
    try:
        FiresModel = get_model_safe('fire_statistics', 'Fire_Statistics', 'FireStatistics', 'FiresStats')
        if not FiresModel: return []
        
        query = db.query(FiresModel)
        query = apply_opsec_scope(current_user, query, FiresModel)
        query = build_filtered_query(query, FiresModel, region, station, start_date, end_date, search)
        
        pk_col = getattr(FiresModel, 'id', getattr(FiresModel, 'sn', None))
        records = query.order_by(pk_col.desc()).limit(limit).all() if pk_col is not None and hasattr(pk_col, 'desc') else query.limit(limit).all()
        
        results = []
        for r in records:
            narrative_val = str(getattr(r, 'narrative', getattr(r, 'nature_of_fire', '')) or '')
            offence_val = str(getattr(r, 'nature_of_fire', '') or '')
            
            text_blob = f"{offence_val} {narrative_val}".lower()
            classification = "RESIDENTIAL PROPERTY"
            if any(k in text_blob for k in ["dormitory", "dorm", "hostel"]):
                classification = "SCHOOL DORMITORY"
            elif any(k in text_blob for k in ["school", "college", "academy", "nursery", "student"]):
                classification = "SCHOOL FACILITY"
            elif any(k in text_blob for k in ["factory", "industrial", "plant", "manufacturing"]):
                classification = "INDUSTRIAL / FACTORY"
            elif any(k in text_blob for k in ["shop", "market", "store", "commercial", "hardware", "tailoring", "facility", "business"]):
                classification = "COMMERCIAL / BUSINESS ROOM"
            elif any(k in text_blob for k in ["church", "mosque", "chapel", "temple"]):
                classification = "RELIGIOUS CENTRE"
            elif any(k in text_blob for k in ["vehicle", "car", "taxi", "bus", "lorry", "truck", "motorcycle", "boda"]):
                classification = "MOTOR VEHICLE / TRANSPORT"
            elif any(k in text_blob for k in ["rental", "room", "house", "residence", "residential"]):
                classification = "RESIDENTIAL ROOM / HOUSE"

            status_findings = "UNDER INVESTIGATION"
            if "cause of fire was still under investigation" in text_blob or "still under investigation" in text_blob:
                status_findings = "UNDER INVESTIGATION"
            elif "spark" in text_blob or "electrical" in text_blob or "short circuit" in text_blob:
                status_findings = "CAUSE ESTABLISHED: ELECTRICAL SHORT CIRCUIT"
            elif "arson" in text_blob or "set the dormitory on fire" in text_blob or "malicious" in text_blob or "student" in text_blob:
                status_findings = "CAUSE ESTABLISHED: ARSON / DELIBERATE ACT"
            elif "charcoal" in text_blob or "stove" in text_blob or "candle" in text_blob or "cook" in text_blob:
                status_findings = "CAUSE ESTABLISHED: UNATTENDED HEAT SOURCE"

            results.append({
                "id": getattr(r, 'id', getattr(r, 'sn', 0)),
                "date": str(getattr(r, 'date', '') or ''),
                "region": str(getattr(r, 'region', '') or ''),
                "station": str(getattr(r, 'station', '') or ''),
                "nature_of_fire": classification,
                "status_findings": status_findings,
                "minor_outbreaks": safe_int(getattr(r, 'minor_outbreaks', 0)),
                "major_disasters": safe_int(getattr(r, 'major_disasters', 0)),
                "fatal_fires": safe_int(getattr(r, 'fatal_fires', 0)),
                "victims": safe_int(getattr(r, 'victims', 0)),
                "total": safe_int(getattr(r, 'total', 0)),
                "summary": f"Findings: {status_findings} | {narrative_val}",
                "entered_by": str(getattr(r, 'entered_by', getattr(r, 'last_updated_by', '')) or ''),
                "created_at": r.created_at.isoformat() if hasattr(r, 'created_at') and r.created_at else ''
            })
        return results
    except Exception as e:
        db.rollback()
        print(f"Fires Endpoint Error: {e}")
        return []

@router.post("/fires")
def create_fires_stat(data: dict, db: Session = Depends(get_db), logs_db: Session = Depends(get_logs_db), current_user: models.Users = Depends(get_current_user)):
    FiresModel = get_model_safe('Fire_Statistics', 'FireStatistics', 'fire_statistics', 'FiresStats')
    if not FiresModel: raise HTTPException(status_code=500, detail="Fire statistics model not configured in database.")
    try:
        data.pop('sn', None); data.pop('id', None)
        valid_cols = [c.key for c in FiresModel.__table__.columns]
        safe_data = {k: v for k, v in data.items() if k in valid_cols}
        
        safe_data['total'] = safe_int(safe_data.get('minor_outbreaks', 0)) + safe_int(safe_data.get('major_disasters', 0)) + safe_int(safe_data.get('fatal_fires', 0))
        safe_data['entered_by'] = get_officer_signature(current_user)
        
        new_record = FiresModel(**safe_data)
        db.add(new_record); db.commit(); db.refresh(new_record)

        rec_id = getattr(new_record, 'id', getattr(new_record, 'sn', '0'))
        record_neon_activity(
            logs_db=logs_db,
            fnum=current_user.fnum,
            action_type="CREATE",
            module="FIRE_STATS",
            target_id=str(rec_id),
            changes_summary=f"{current_user.fnum} recorded Fire Incident stats for {new_record.station}."
        )

        clean = clean_model_dict(new_record)
        if hasattr(new_record, 'created_at') and new_record.created_at:
            clean['created_at'] = new_record.created_at.isoformat() if hasattr(new_record.created_at, 'isoformat') else str(new_record.created_at)
        return clean
    except Exception as e:
        db.rollback(); logs_db.rollback(); raise HTTPException(status_code=500, detail=str(e))

@router.put("/fires/{record_id}")
def update_fires_stat(record_id: int, data: dict, db: Session = Depends(get_db), logs_db: Session = Depends(get_logs_db), current_user: models.Users = Depends(get_current_user)):
    FiresModel = get_model_safe('Fire_Statistics', 'FireStatistics', 'fire_statistics', 'FiresStats')
    if not FiresModel: raise HTTPException(status_code=500, detail="Fire statistics model not configured.")
    try:
        pk_col = getattr(FiresModel, 'id', getattr(FiresModel, 'sn', None))
        record = db.query(FiresModel).filter(pk_col == record_id).first()
        if not record: raise HTTPException(status_code=404, detail="Fire record not found.")
        
        data.pop('id', None); data.pop('sn', None)
        for key, value in data.items():
            if hasattr(record, key): setattr(record, key, value)
            
        record.total = safe_int(record.minor_outbreaks) + safe_int(record.major_disasters) + safe_int(record.fatal_fires)
        record.last_updated_by = get_officer_signature(current_user)
            
        db.commit(); db.refresh(record)

        record_neon_activity(
            logs_db=logs_db,
            fnum=current_user.fnum,
            action_type="UPDATE",
            module="FIRE_STATS",
            target_id=str(record_id),
            changes_summary=f"{current_user.fnum} updated Fire Incident stats ID {record_id}."
        )

        clean = clean_model_dict(record)
        if hasattr(record, 'created_at') and record.created_at:
            clean['created_at'] = record.created_at.isoformat() if hasattr(record.created_at, 'isoformat') else str(record.created_at)
        return clean
    except Exception as e:
        db.rollback(); logs_db.rollback(); raise HTTPException(status_code=500, detail=str(e))

@router.get("/fires/export")
def export_fires_stats_secure_zip(
    region: Optional[str] = Query(default=None), station: Optional[str] = Query(default=None),
    start_date: Optional[str] = Query(default=None), end_date: Optional[str] = Query(default=None),
    search: Optional[str] = Query(default=None), db: Session = Depends(get_db),
    logs_db: Session = Depends(get_logs_db), current_user: models.Users = Depends(get_current_user)
):
    FiresModel = get_model_safe('Fire_Statistics', 'FireStatistics', 'fire_statistics', 'FiresStats')
    if not FiresModel: raise HTTPException(status_code=500, detail="Model not found")

    query = db.query(FiresModel)
    query = apply_opsec_scope(current_user, query, FiresModel)
    query = build_filtered_query(query, FiresModel, region, station, start_date, end_date, search)
    pk_col = getattr(FiresModel, 'id', getattr(FiresModel, 'sn', None))
    records = query.order_by(pk_col.desc()).all() if pk_col is not None else query.all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fire & Rescue Stats"
    ws.views.sheetView[0].showGridLines = True

    eat_tz = pytz.timezone("Africa/Nairobi")
    eat_time = datetime.now(eat_tz).replace(tzinfo=None)
    
    officer_name = (current_user.name or "UNKNOWN").strip().upper()
    officer_fnum = (current_user.fnum or "UNKNOWN").strip().upper()
    officer_rank = (current_user.rank or "").strip().upper()
    officer_station = (current_user.station or "HQ").strip().upper()
    timestamp_str = eat_time.strftime("%Y-%m-%d %H:%M:%S EAT")

    wb.properties.creator = f"{officer_fnum} {officer_rank} {officer_name}"
    wb.properties.lastModifiedBy = f"KMP-CSDMS SECURITY AUDIT ENGINE [{officer_station}]"
    wb.properties.title = "Official Law Enforcement Restricted Fire Incidents Export"

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="B45309", end_color="B45309", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell_font = Font(name="Calibri", size=10)
    thin_border = Border(left=Side(style='thin', color='D1D5DB'), right=Side(style='thin', color='D1D5DB'), top=Side(style='thin', color='D1D5DB'), bottom=Side(style='thin', color='D1D5DB'))

    ws.merge_cells("A1:K1")
    banner_cell = ws.cell(row=1, column=1)
    banner_cell.value = f"RESTRICTED LAW ENFORCEMENT RECORD — FORENSICALLY STAMPED TO OFFICER: {officer_fnum} {officer_rank} {officer_name} ({officer_station}) — TIMESTAMP: {timestamp_str}"
    banner_cell.font = Font(name="Calibri", size=9, bold=True, color="991B1B")
    banner_cell.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    banner_cell.alignment = Alignment(horizontal="center", vertical="center")
    
    ws.cell(row=2, column=1, value="UGANDA POLICE FORCE - FIRE & RESCUE INCIDENTS STATISTICS LEDGER").font = Font(bold=True, size=11)
    ws.row_dimensions[2].height = 18

    headers = ["ID", "Date", "Division/Station", "Nature of Fire", "Minor Outbreaks", "Major Disasters", "Fatal Fires", "Victims", "Total", "Created By", "Date Created"]
    ws.row_dimensions[4].height = 26
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    current_row = 5
    for r in records:
        created_str = r.created_at.strftime("%Y-%m-%d %H:%M") if hasattr(r, 'created_at') and r.created_at else ''
        row_data = [
            getattr(r, 'id', getattr(r, 'sn', '')), getattr(r, 'date', ''), getattr(r, 'station', ''), getattr(r, 'nature_of_fire', ''),
            getattr(r, 'minor_outbreaks', 0), getattr(r, 'major_disasters', 0), getattr(r, 'fatal_fires', 0), getattr(r, 'victims', 0), getattr(r, 'total', 0),
            getattr(r, 'entered_by', getattr(r, 'last_updated_by', '')), created_str
        ]
        ws.append(row_data)
        for col_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.font = cell_font
            cell.border = thin_border
            if col_idx in [1, 2, 5, 6, 7, 8, 9]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
        current_row += 1

    for col in ws.columns:
        col_letter = col[0].column_letter
        max_len = max([len(str(cell.value or '')) for cell in col], default=0)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 35)

    excel_stream = io.BytesIO()
    wb.save(excel_stream)
    excel_stream.seek(0)

    zip_stream = io.BytesIO()
    zip_password = str(current_user.fnum).strip().encode('utf-8')
    fnum_clean = str(current_user.fnum).replace('/', '_').upper()
    excel_filename = f"{fnum_clean}_Fire_Stats_{eat_time.strftime('%Y%m%d')}.xlsx"
    zip_filename = f"SECURE_FIRE_STATS_{eat_time.strftime('%Y%m%d')}.zip"

    with pyzipper.AESZipFile(zip_stream, 'w', compression=pyzipper.ZIP_DEFLATED, encryption=pyzipper.WZ_AES) as zf:
        zf.setpassword(zip_password)
        zf.writestr(excel_filename, excel_stream.getvalue())
    zip_stream.seek(0)

    record_neon_activity(
        logs_db=logs_db, fnum=current_user.fnum, action_type="EXPORT", module="FIRE_STATS", target_id="SECURE_ZIP_EXPORT",
        changes_summary=f"{current_user.fnum} securely exported AES-256 encrypted Fire & Rescue Statistics archive."
    )

    return StreamingResponse(
        zip_stream, media_type="application/zip",
        headers={'Content-Disposition': f'attachment; filename="{zip_filename}"', 'Access-Control-Expose-Headers': 'Content-Disposition'}
    )