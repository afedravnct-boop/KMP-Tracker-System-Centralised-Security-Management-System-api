import io
import os
import json
import base64
import urllib.parse
from datetime import datetime
from typing import Optional, List

import boto3
import openpyxl
import pymupdf
import pytz
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from docx.oxml import parse_xml # Required for floating vertical text
from pptx import Presentation
from pptx.util import Inches, Pt as PPTXPt
from pptx.dml.color import RGBColor as PPTXRGBColor

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form, status
from fastapi.responses import StreamingResponse, JSONResponse
from sqlalchemy.orm import Session
from sqlalchemy import text

from app import models
from app.database import get_db
from auth import get_current_user

router = APIRouter(prefix="/api/v1/templates", tags=["Command Templates"])

# AWS S3 Client Configuration
AWS_REGION = os.getenv("AWS_REGION", "eu-central-1")
BUCKET_NAME = os.getenv("AWS_BUCKET_NAME", "kmp-centralised-security-storage")

s3_client = boto3.client(
    "s3",
    aws_access_key_id=os.getenv("AWS_ACCESS_KEY_ID"),
    aws_secret_access_key=os.getenv("AWS_SECRET_ACCESS_KEY"),
    region_name=AWS_REGION
)

def get_eat_now():
    eat_tz = pytz.timezone("Africa/Nairobi")
    return datetime.now(eat_tz).replace(tzinfo=None)

def get_template_model():
    for model_name in ['CommandTemplate', 'Command_Template', 'command_template', 'command_templates']:
        if hasattr(models, model_name):
            return getattr(models, model_name)
    return None

@router.get("/list")
def get_command_templates(db: Session = Depends(get_db), current_user = Depends(get_current_user)):
    try:
        TemplateModel = get_template_model()
        if not TemplateModel:
            return []
            
        templates = db.query(TemplateModel).all()
        results = []
        for t in templates:
            filename = getattr(t, 'file_name', getattr(t, 'filename', getattr(t, 'file_path', 'document')))
            ext = filename.split('.')[-1].lower() if '.' in filename else "unknown"
            
            if ext in ['docx', 'doc']: file_type = "Word Document"
            elif ext in ['xlsx', 'xls']: file_type = "Excel Spreadsheet"
            elif ext in ['pptx', 'ppt']: file_type = "PowerPoint Presentation"
            elif ext == 'pdf': file_type = "PDF Document"
            else: file_type = getattr(t, 'doc_type', 'Command Template')

            upload_dt = getattr(t, 'upload_date', getattr(t, 'uploaded_at', None))
            date_str = upload_dt.strftime("%Y-%m-%d") if isinstance(upload_dt, datetime) else str(upload_dt or "").split(' ')[0]

            results.append({
                "id": getattr(t, 'id', getattr(t, 'sn', 1)),
                "name": filename,
                "type": file_type,
                "date": date_str,
                "size": getattr(t, 'file_size', 'N/A'),
                "file_path": getattr(t, 'file_path', getattr(t, 'filepath', '')),
                "region": getattr(t, 'region', 'KMP HEADQUARTERS'),
                "station": getattr(t, 'station', 'HQ')
            })
        return results
    except Exception as e:
        print(f"Templates List Notice: {str(e)}")
        return []

@router.post("/upload/{template_id_key}")
async def upload_command_template(
    template_id_key: str,
    file: Optional[UploadFile] = File(None),
    files: Optional[List[UploadFile]] = File(None),
    doc_type: str = Form("Command Template"),
    db: Session = Depends(get_db),
    current_user = Depends(get_current_user)
):
    TemplateModel = get_template_model()
    if not TemplateModel:
        raise HTTPException(status_code=500, detail="Command Template table model not initialized.")

    file_list = [f for f in [file, *(files or [])] if f is not None]
    if not file_list:
        raise HTTPException(status_code=400, detail="No template file provided.")

    eat_time = get_eat_now()
    uploaded_count = 0

    try:
        for single_file in file_list:
            contents = await single_file.read()
            file_size_kb = max(1, round(len(contents) / 1024))
            file_size_str = f"{file_size_kb} KB" if file_size_kb < 1024 else f"{round(file_size_kb / 1024, 1)} MB"

            timestamp = eat_time.strftime("%Y%m%d_%H%M%S")
            safe_filename = f"{timestamp}_{single_file.filename.replace(' ', '_')}"
            s3_key = f"command_templates/{safe_filename}"
            
            s3_client.put_object(
                Bucket=BUCKET_NAME, Key=s3_key, Body=contents,
                ContentType=single_file.content_type or "application/octet-stream", ServerSideEncryption="AES256"
            )
            
            full_s3_url = f"https://{BUCKET_NAME}.s3.{AWS_REGION}.amazonaws.com/{s3_key}"
            
            new_template = TemplateModel(
                file_name=single_file.filename,
                doc_type=doc_type,
                file_size=file_size_str,
                file_path=full_s3_url,
                region=getattr(current_user, 'region', 'KMP HEADQUARTERS'),
                station=getattr(current_user, 'station', 'HQ'),
                uploaded_by=current_user.fnum,
                upload_date=eat_time
            )
            db.add(new_template)
            uploaded_count += 1

        db.commit()
        return {"status": "success", "message": f"Successfully uploaded {uploaded_count} template(s)."}
    except Exception as e:
        print(f"Fetch error: {e}")
        raise HTTPException(status_code=500, detail=f"Database Error: {str(e)}")


# 🟢 ADDED: Master Download & Forensic Stamping for Templates
@router.get("/download/{doc_id}")
def download_template_file(
    doc_id: int, 
    return_url: bool = False,
    db: Session = Depends(get_db),
    current_user = Depends(get_current_user)
):
    TemplateModel = get_template_model()
    if not TemplateModel:
        raise HTTPException(status_code=500, detail="Command Template table model not initialized.")

    doc_record = db.query(TemplateModel).filter(TemplateModel.id == doc_id).first()
    if not doc_record:
        raise HTTPException(status_code=404, detail="Template record not found in system database.")
        
    file_path = getattr(doc_record, 'file_path', getattr(doc_record, 'url', ''))
    file_name = getattr(doc_record, 'file_name', getattr(doc_record, 'name', 'template'))

    if not str(file_path).startswith("http"):
        raise HTTPException(status_code=404, detail="File path invalid or missing S3 storage link.")

    parsed_url = urllib.parse.urlparse(file_path)
    original_s3_key = parsed_url.path.lstrip('/') 
    file_extension = file_name.lower().split('.')[-1] if '.' in file_name else "bin"

    try:
        file_stream = io.BytesIO()
        s3_client.download_fileobj(BUCKET_NAME, original_s3_key, file_stream)
        file_stream.seek(0)
        raw_bytes = file_stream.getvalue()

        eat_time = get_eat_now()
        timestamp_eat = eat_time.strftime("%Y-%m-%d %H:%M:%S EAT")

        officer_fnum = (current_user.fnum or "HQ-UNKNOWN").strip().upper()
        officer_rank = (current_user.rank or "OFFICER").strip().upper()
        officer_name = (current_user.name or "UNKNOWN").strip().upper()
        officer_signature = f"{officer_fnum} {officer_rank} {officer_name}"
        command_post = f"{current_user.station or 'KMP HEADQUARTERS'}, {current_user.region or 'KMP HEADQUARTERS'}"
        stamp_id = f"KMP-TMP-{officer_fnum}-{eat_time.strftime('%Y%m%d%H%M%S')}"

        compact_payload = {"f": officer_fnum, "s": stamp_id}
        encoded_token = base64.b64encode(json.dumps(compact_payload).encode('utf-8')).decode('utf-8')
        keywords_str = f"KMP_AUDIT;{encoded_token}"[:250]
        comments_str = f"Export: {officer_signature} [{command_post}]. ID: {stamp_id}"

        # 🟢 Vertical Left Margin Stamp Text
        vertical_stamp_text = f"SECURE ACCESS BY: {officer_signature}  |  CLEARANCE: {current_user.role}  |  STAMP ID: {stamp_id}  |  TIMESTAMP: {timestamp_eat}"

        output_stream = io.BytesIO()
        content_type = "application/octet-stream"

        if file_extension == 'docx':
            word_doc = Document(io.BytesIO(raw_bytes))
            core_props = word_doc.core_properties
            core_props.author = officer_signature
            core_props.last_modified_by = officer_signature
            core_props.keywords = keywords_str
            core_props.comments = comments_str
            core_props.category = "RESTRICTED / LAW ENFORCEMENT RECORD"

            # 🟢 VML Injection for Floating Vertical Text on the Left Margin
            section = word_doc.sections[0]
            header = section.header
            if not header.paragraphs:
                header_p = header.add_paragraph()
            else:
                header_p = header.paragraphs[0]
            
            try:
                vml_xml = f'''
                <w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" 
                     xmlns:v="urn:schemas-microsoft-com:vml" 
                     xmlns:o="urn:schemas-microsoft-com:office:office">
                    <w:pict>
                        <v:shapetype id="_x0000_t136" coordsize="21600,21600" o:spt="136" adj="10800" path="m@7,l@8,m@5,21600l@6,21600e">
                            <v:path textpathok="t" o:connecttype="custom" o:connectlocs="@9,0;@10,10800;@11,21600;@12,10800" o:connectangles="270,180,90,0"/>
                            <v:textpath on="t" fitshape="t"/>
                            <o:lock v:ext="edit" text="t" shapetype="t"/>
                        </v:shapetype>
                        <v:shape id="VerticalStamp" type="#_x0000_t136" 
                                 style="position:absolute;left:0;text-align:center;margin-left:15pt;margin-top:0pt;width:15pt;height:550pt;rotation:270;z-index:-251657216;mso-position-horizontal:left;mso-position-vertical:center;mso-position-horizontal-relative:page;mso-position-vertical-relative:page" 
                                 fillcolor="#8B0000" stroked="f">
                            <v:textpath style="font-family:'Courier New';font-size:7pt;font-weight:bold" string="{vertical_stamp_text}"/>
                        </v:shape>
                    </w:pict>
                </w:r>
                '''
                vml_run = parse_xml(vml_xml)
                header_p._p.append(vml_run)
            except Exception as e:
                print(f"Failed to inject VML vertical watermark: {e}")
            
            word_doc.save(output_stream)
            content_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

        elif file_extension in ['xlsx', 'xls']:
            wb = openpyxl.load_workbook(io.BytesIO(raw_bytes))
            wb.properties.creator = officer_signature
            wb.properties.lastModifiedBy = officer_signature
            wb.properties.keywords = keywords_str
            wb.properties.description = comments_str
            wb.properties.category = "RESTRICTED / FORENSIC POLICE RECORD"
            
            for ws in wb.worksheets:
                if hasattr(ws, 'sheet_footer'): 
                    ws.sheet_footer.left.text = vertical_stamp_text
                elif hasattr(ws, 'odd_footer'): 
                    ws.odd_footer.left.text = vertical_stamp_text
            wb.save(output_stream)
            content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

        elif file_extension in ['pptx', 'ppt']:
            try:
                prs = Presentation(io.BytesIO(raw_bytes))
                prs.core_properties.author = officer_signature
                prs.core_properties.last_modified_by = officer_signature
                prs.core_properties.keywords = keywords_str
                prs.core_properties.comments = comments_str
                prs.core_properties.category = "RESTRICTED / FORENSIC POLICE RECORD"
                
                if prs.slides:
                    slide = prs.slides[0]
                    left_box = slide.shapes.add_textbox(Inches(0.1), Inches(1.5), Inches(8), Inches(0.5))
                    left_box.rotation = 270 
                    p_left = left_box.text_frame.add_paragraph()
                    p_left.text = vertical_stamp_text
                    p_left.font.size = PPTXPt(7)
                    p_left.font.bold = True
                    p_left.font.name = 'Courier New'
                    p_left.font.color.rgb = PPTXRGBColor(139, 0, 0)

                prs.save(output_stream)
                content_type = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
            except Exception:
                output_stream.write(raw_bytes)
                content_type = "application/vnd.openxmlformats-officedocument.presentationml.presentation"

        elif file_extension == 'pdf':
            try:
                pdf_doc = pymupdf.open(stream=raw_bytes, filetype="pdf")
                for page in pdf_doc:
                    rect = page.rect
                    # 🟢 Rotated 90-degrees upward on Left Margin
                    page.insert_text(
                        pymupdf.Point(20, rect.height - 100),
                        vertical_stamp_text,
                        fontsize=7,
                        fontname="courier-bold",
                        color=(0.545, 0, 0),
                        rotate=90 
                    )
                stamped_pdf_bytes = pdf_doc.tobytes()
                output_stream = io.BytesIO(stamped_pdf_bytes)
                content_type = "application/pdf"
            except Exception:
                output_stream.write(raw_bytes)
                content_type = "application/pdf"

        else:
            output_stream.write(raw_bytes)

        output_stream.seek(0)
        final_bytes = output_stream.getvalue()

        # 🟢 If it's a "Read" request, return the JSON URL so the frontend can open it in the viewer
        if return_url:
            # 🟢 FIX: Clean the filename of spaces and special chars. 
            # Google Docs Viewer throws a 'Network Error' if the S3 URL contains spaces that get double URL-encoded (%2520).
            safe_file_name = file_name.replace(" ", "_").replace("%20", "_")
            temp_s3_key = f"forensic_cache/{stamp_id}_{safe_file_name}"
            
            s3_client.put_object(
                Bucket=BUCKET_NAME,
                Key=temp_s3_key,
                Body=final_bytes,
                ContentType=content_type,
                ContentDisposition="inline", # 🟢 FIX: Force inline rendering on the S3 Object
                ServerSideEncryption="AES256"
            )
            
            presigned_url = s3_client.generate_presigned_url(
                'get_object',
                Params={
                    'Bucket': BUCKET_NAME, 
                    'Key': temp_s3_key,
                    'ResponseContentDisposition': 'inline', # 🟢 FIX: Guarantee the URL does not trigger a 'Save As' dialogue
                    'ResponseContentType': content_type
                },
                ExpiresIn=3600
            )
            return JSONResponse(content={"url": presigned_url})

        # 🟢 If it's a "Download" request, return the streaming attachment
        return StreamingResponse(
            io.BytesIO(final_bytes),
            media_type=content_type,
            headers={
                "Content-Disposition": f'attachment; filename="{file_name}"'
            }
        )

    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Template Download Error: {str(e)}")