# routers/establishments.py
import json
import csv
import io
import os
from datetime import datetime
from typing import Optional, List, Union
from fastapi import APIRouter, Depends, HTTPException, status, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import func, or_, text
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import pyzipper

from app import models
from app.database import get_db, get_logs_db
from auth import get_current_user
from routers.activity_logger import record_neon_activity

router = APIRouter(prefix="/api/v1", tags=["Establishments"])

# 🟢 Enriched hierarchy ensuring both "REGION HEADQUARTERS" and "REGION" designations exist
REGIONAL_HIERARCHY = {
    "KMP NORTH": ["KMP NORTH HEADQUARTERS", "KMP NORTH", "KAWEMPE", "KAKIRI", "KASANGATI", "MATUGGA", "NANSANA", "OLD KAMPALA", "WAKISO", "WANDEGEYA"],
    "KMP EAST": ["KMP EAST HEADQUARTERS", "KMP EAST", "JINJA ROAD", "KIRA", "KIRA DIV", "KIRA ROAD", "MUKONO", "NAGGALAMA", "SEETA"],
    "KMP SOUTH": ["KMP SOUTH HEADQUARTERS", "KMP SOUTH", "NATEETE", "CPS KAMPALA", "PARLIAMENT", "ENTEBBE", "KABALAGALA", "KAJJANSI", "KASENYI", "KATWE", "KYENGERA", "NSANGI"],
    "KMP HEADQUARTERS": ["KMP HEADQUARTERS", "KMP CID", "KMP TRAFFIC", "KMP ICT", "KMP FLYING SQUAD", "KMP CRIME INTELLIGENCE"],
    "POLICE HEADQUARTERS": ["NAGURU", "OPERATIONS", "CRIME INTELLIGENCE", "CID", "LOGISTICS & ENGINEERING", "ICT", "CT", "FIRE & RESCUE"]
}

def get_officer_signature(user):
    if not user:
        return "UNKNOWN COMMANDER"
    fnum = (user.fnum or "").strip()
    rank = (user.rank or "").strip()
    name = (user.name or "").strip()
    return f"{fnum} {rank} {name}".strip().upper()

def get_est_model():
    model = getattr(models, 'Establishments', getattr(models, 'establishments', None))
    if not model:
        raise HTTPException(status_code=500, detail="Establishments database model not configured.")
    return model

def serialize_row(row):
    if not row:
        return {}
    d = row.__dict__.copy()
    d.pop('_sa_instance_state', None)
    if 'sn' in d and 'id' not in d:
        d['id'] = d['sn']
    for k, v in d.items():
        if isinstance(v, datetime):
            d[k] = v.strftime("%Y-%m-%d %H:%M:%S")
        elif hasattr(v, 'isoformat'):
            d[k] = v.isoformat()
    return d

def apply_opsec_scope(current_user, query, ModelClass):
    if not ModelClass:
        return query
        
    user_role = str(current_user.role).strip().upper() if current_user.role else ""
    user_pos = str(current_user.position).strip().upper() if current_user.position else ""
    user_reg = str(current_user.region).strip().upper() if current_user.region else ""
    user_stn = str(current_user.station).strip().upper() if current_user.station else ""

    perms = current_user.permissions or {}
    if isinstance(perms, str):
        try: perms = json.loads(perms)
        except Exception: perms = {}

    is_absolute_global = (
        user_role in ["SUPER_ADMIN", "ADMIN", "ASSISTANT_SUPER_ADMIN", "RPC", "DEPUTY COMMANDER"] or
        "KMP COMMANDER" in user_pos or
        "DEPUTY KMP COMMANDER" in user_pos or
        "KMP ADMIN" in user_pos or
        perms.get("view_global_roster") is True or
        perms.get("global_observer") is True
    )

    is_kmp_sys_mgr = (
        user_role == "SYSTEM_MANAGER" and
        user_reg in ["KMP HEADQUARTERS", "POLICE HEADQUARTERS"] and
        "KMP" in user_pos
    )

    is_regional_command = (
        user_role in ["RPC", "DEPUTY_RPC", "SYSTEM_MANAGER", "ASSISTANT_SYSTEM_MANAGER", "REGIONAL_ADMIN", "ASSISTANT_REGIONAL_ADMIN", "DIVISION_ADMIN"] or
        "HR" in user_pos or
        not is_absolute_global
    )

    if is_absolute_global or is_kmp_sys_mgr:
        pass 
    elif user_reg in REGIONAL_HIERARCHY:
        conds = [func.upper(ModelClass.region) == user_reg]
        
        expanded_stns = set()
        for s in REGIONAL_HIERARCHY[user_reg]:
            expanded_stns.add(s)
            expanded_stns.add(s.replace(' HEADQUARTERS', '').replace(' HQ', ''))
            expanded_stns.add(s + ' HEADQUARTERS')
            expanded_stns.add(s + ' HQ')
            
        conds.append(func.upper(ModelClass.station).in_(list(expanded_stns)))
        if hasattr(ModelClass, 'division'):
            conds.append(func.upper(ModelClass.division).in_(list(expanded_stns)))
            
        query = query.filter(or_(*conds))
    else:
        query = query.filter(func.upper(ModelClass.station) == user_stn)

    return query

def build_filtered_query(query, ModelClass, region, station, search):
    if region and region.upper() not in ["ALL REGIONS", "ALL"]:
        if hasattr(ModelClass, 'region'):
            query = query.filter(func.upper(ModelClass.region) == region.upper())
            
    if station and station.upper() not in ["ALL STATIONS", "ALL"]:
        if hasattr(ModelClass, 'station'):
            clean_stn = station.upper().replace(' HEADQUARTERS', '').replace(' HQ', '')
            query = query.filter(
                or_(
                    func.upper(ModelClass.station) == station.upper(),
                    func.upper(ModelClass.station) == clean_stn,
                    func.upper(ModelClass.station) == f"{clean_stn} HEADQUARTERS"
                )
            )

    if search:
        term = f"%{search.strip().upper()}%"
        search_conds = []
        if hasattr(ModelClass, 'station'): search_conds.append(ModelClass.station.ilike(term))
        if hasattr(ModelClass, 'region'): search_conds.append(ModelClass.region.ilike(term))
        if hasattr(ModelClass, 'division'): search_conds.append(ModelClass.division.ilike(term))
        if hasattr(ModelClass, 'last_updated_by'): search_conds.append(ModelClass.last_updated_by.ilike(term))
        if search_conds:
            query = query.filter(or_(*search_conds))
            
    return query

# ====================================================================
# 1. RETRIEVE ESTABLISHMENTS
# ====================================================================
@router.get("/establishments")
def get_all_establishments(
    region: Optional[str] = Query(default=None),
    station: Optional[str] = Query(default=None),
    search: Optional[str] = Query(default=None),
    db: Session = Depends(get_db), 
    logs_db: Session = Depends(get_logs_db),
    current_user: models.Users = Depends(get_current_user)
):
    EstModel = get_est_model()
    query = db.query(EstModel)
    
    query = apply_opsec_scope(current_user, query, EstModel)
    query = build_filtered_query(query, EstModel, region, station, search)

    pk_col = getattr(EstModel, 'id', getattr(EstModel, 'sn', None))
    if pk_col is not None:
        records = query.order_by(pk_col.desc()).all()
    else:
        records = query.all()
        
    serialized_records = [serialize_row(r) for r in records]

    if search:
        summary_text = f"{current_user.fnum} {current_user.rank} {current_user.name} searched establishments for query: \"{search}\" (Returned {len(serialized_records)} matches)."
    else:
        summary_text = f"{current_user.fnum} {current_user.rank} {current_user.name} accessed Regional Establishments directory (Fetched {len(serialized_records)} records)."

    record_neon_activity(
        logs_db=logs_db,
        fnum=current_user.fnum,
        action_type="VIEW",
        module="ESTABLISHMENTS",
        target_id=search if search else "ALL_ESTABLISHMENTS",
        changes_summary=summary_text
    )

    return serialized_records

# ====================================================================
# 1B. EXPORT ESTABLISHMENTS TO SECURE AES-256 ZIP (.XLSX)
# ====================================================================
@router.get("/establishments/export")
def export_establishments_excel(
    region: Optional[str] = Query(default=None),
    station: Optional[str] = Query(default=None),
    search: Optional[str] = Query(default=None),
    db: Session = Depends(get_db),
    logs_db: Session = Depends(get_logs_db),
    current_user: models.Users = Depends(get_current_user)
):
    EstModel = get_est_model()
    query = db.query(EstModel)

    query = apply_opsec_scope(current_user, query, EstModel)
    query = build_filtered_query(query, EstModel, region, station, search)

    pk_col = getattr(EstModel, 'id', getattr(EstModel, 'sn', None))
    if pk_col is not None:
        records = query.order_by(pk_col.desc()).all()
    else:
        records = query.all()

    # 1. Build Excel Workbook via openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Establishments Ledger"
    ws.views.sheetView[0].showGridLines = True

    # 2. Forensic Stamping in Document Properties
    officer_name = (current_user.name or "UNKNOWN").strip().upper()
    officer_fnum = (current_user.fnum or "UNKNOWN").strip().upper()
    officer_rank = (current_user.rank or "").strip().upper()
    officer_station = (current_user.station or "HQ").strip().upper()
    timestamp_str = datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S UTC")

    wb.properties.creator = f"{officer_fnum} {officer_rank} {officer_name}"
    wb.properties.lastModifiedBy = f"KMP-CSDMS SECURITY AUDIT ENGINE [{officer_station}]"
    wb.properties.title = "Official Law Enforcement Restricted Establishments Export"
    wb.properties.subject = f"Establishments Data Extract — Downloaded by {officer_fnum} ({officer_station}) at {timestamp_str}"
    wb.properties.description = f"SECURE OPSEC RECORD. Cryptographically keyed and audited to Force Number: {officer_fnum}."

    # Styling definitions
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    cell_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'), right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'), bottom=Side(style='thin', color='D1D5DB')
    )

    # 3. Insert Forensic Audit Banner Row at the Top
    ws.merge_cells("A1:J1")
    banner_cell = ws.cell(row=1, column=1)
    banner_cell.value = f"RESTRICTED LAW ENFORCEMENT RECORD — FORENSICALLY STAMPED TO OFFICER: {officer_fnum} {officer_rank} {officer_name} ({officer_station}) — TIMESTAMP: {timestamp_str}"
    banner_cell.font = Font(name="Calibri", size=9, bold=True, color="991B1B")
    banner_cell.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    banner_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 22

    # 4. Table Headers (Row 3)
    headers = [
        "ID", "Region", "Division", "Station", 
        "Personnel In Station", "Personnel In Sub-Station", "Personnel In Post", 
        "Booths", "Personnel In Booth", "Last Updated By"
    ]
    ws.row_dimensions[3].height = 26
    for col_idx, header_text in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_idx, value=header_text)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    total_records = len(records)

    # 5. Populate Data Rows starting at Row 4
    for idx, r in enumerate(records):
        row_data = [
            getattr(r, 'id', getattr(r, 'sn', '')),
            getattr(r, 'region', ''),
            getattr(r, 'division', ''),
            getattr(r, 'station', ''),
            getattr(r, 'personnel_in_station', 0),
            getattr(r, 'personnel_in_sub_station', 0),
            getattr(r, 'personnel_in_post', 0),
            getattr(r, 'booths', 0),
            getattr(r, 'personnel_in_booth', 0),
            getattr(r, 'last_updated_by', '')
        ]
        ws.append(row_data)
        current_row = idx + 4
        ws.row_dimensions[current_row].height = 22

        for col_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.font = cell_font
            cell.border = thin_border
            if col_idx in [5, 6, 7, 8, 9]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    column_widths = {'A': 8, 'B': 18, 'C': 18, 'D': 20, 'E': 22, 'F': 22, 'G': 20, 'H': 12, 'I': 20, 'J': 30}
    for col_letter, width in column_widths.items():
        ws.column_dimensions[col_letter].width = width

    excel_buffer = io.BytesIO()
    wb.save(excel_buffer)
    excel_buffer.seek(0)

    # 6. Compress and Encrypt into AES-256 ZIP Archive using Officer's Force Number
    zip_buffer = io.BytesIO()
    with pyzipper.AESZipFile(zip_buffer, 'w', compression=pyzipper.ZIP_DEFLATED, encryption=pyzipper.WZ_AES) as zf:
        zf.setpassword(officer_fnum.encode('utf-8'))
        file_label = region.replace(" ", "_").upper() if region and region not in ["ALL REGIONS", "ALL"] else "ALL_REGIONS"
        timestamp = datetime.now().strftime("%Y-%m-%d")
        excel_filename = f"KMP_ESTABLISHMENTS_{file_label}_{timestamp}.xlsx"
        zf.writestr(excel_filename, excel_buffer.getvalue())

    zip_buffer.seek(0)
    zip_filename = f"KMP_ESTABLISHMENTS_SECURE_EXPORT_{timestamp}.zip"

    record_neon_activity(
        logs_db=logs_db,
        fnum=current_user.fnum,
        action_type="EXPORT",
        module="ESTABLISHMENTS",
        target_id="AES_ZIP_EXPORT",
        changes_summary=f"Officer generated AES-256 encrypted archive containing {total_records} establishment records (Keyed to FNum: {officer_fnum})."
    )

    return StreamingResponse(
        zip_buffer,
        media_type="application/zip",
        headers={"Content-Disposition": f"attachment; filename={zip_filename}"}
    )

# ====================================================================
# 2. CREATE ESTABLISHMENT
# ====================================================================
@router.post("/establishments")
def create_establishment(
    data: dict, 
    db: Session = Depends(get_db), 
    logs_db: Session = Depends(get_logs_db),
    current_user: models.Users = Depends(get_current_user)
):
    EstModel = get_est_model()
    try:
        data.pop('sn', None)
        data.pop('id', None)
        
        numeric_fields = ['personnel_in_station', 'personnel_in_sub_station', 'personnel_in_post', 'booths', 'personnel_in_booth']
        for field in numeric_fields:
            if field in data:
                try:
                    data[field] = int(data[field]) if data[field] != "" and data[field] is not None else 0
                except (ValueError, TypeError):
                    data[field] = 0

        perms = current_user.permissions or {}
        can_assign_jurisdiction = (
            current_user.role in ["SUPER_ADMIN", "RPC", "ADMIN"] or
            perms.get("global_observer") is True
        )
        
        if not can_assign_jurisdiction:
            data["region"] = current_user.region
            data["division"] = getattr(current_user, 'division', current_user.region)
            data["station"] = current_user.station
            
        new_est = EstModel(**data)
        new_est.last_updated_by = get_officer_signature(current_user)
        
        db.add(new_est)
        db.commit()
        db.refresh(new_est)
        
        assigned_id = getattr(new_est, 'id', getattr(new_est, 'sn', 1))

        record_neon_activity(
            logs_db=logs_db,
            fnum=current_user.fnum,
            action_type="REGISTER",
            module="ESTABLISHMENTS",
            target_id=str(assigned_id),
            changes_summary=f"{current_user.fnum} {current_user.rank} {current_user.name} registered new regional establishment for station [{getattr(new_est, 'station', 'N/A')}]."
        )

        return serialize_row(new_est)
    except Exception as e:
        db.rollback()
        logs_db.rollback()
        print(f"Establishment creation error: {str(e)}")
        raise HTTPException(status_code=500, detail=f"Failed to record establishment: {str(e)}")

# ====================================================================
# 3. UPDATE ESTABLISHMENT
# ====================================================================
@router.put("/establishments/{est_id}")
def update_establishment(
    est_id: int, 
    est_update: dict, 
    db: Session = Depends(get_db), 
    logs_db: Session = Depends(get_logs_db),
    current_user: models.Users = Depends(get_current_user)
):
    EstModel = get_est_model()
    
    pk_col = getattr(EstModel, 'id', getattr(EstModel, 'sn', None))
    existing_est = db.query(EstModel).filter(pk_col == est_id).first()
    
    if not existing_est:
        raise HTTPException(status_code=404, detail="Establishment record not found.")

    est_update.pop('sn', None) 
    est_update.pop('id', None) 
    
    numeric_fields = ['personnel_in_station', 'personnel_in_sub_station', 'personnel_in_post', 'booths', 'personnel_in_booth']
    for field in numeric_fields:
        if field in est_update:
            try:
                est_update[field] = int(est_update[field]) if est_update[field] != "" and est_update[field] is not None else 0
            except (ValueError, TypeError):
                est_update[field] = 0

    perms = current_user.permissions or {}
    can_reassign = (
        current_user.role in ["SUPER_ADMIN", "RPC", "ADMIN"] or
        perms.get("global_observer") is True
    )
    
    if not can_reassign:
        est_update.pop("region", None)
        est_update.pop("division", None)
        est_update.pop("station", None)

    for key, value in est_update.items():
        if hasattr(existing_est, key):
            setattr(existing_est, key, value)

    existing_est.last_updated_by = get_officer_signature(current_user)
    
    try:
        db.commit()
        db.refresh(existing_est)

        record_neon_activity(
            logs_db=logs_db,
            fnum=current_user.fnum,
            action_type="UPDATE",
            module="ESTABLISHMENTS",
            target_id=str(est_id),
            changes_summary=f"{current_user.fnum} {current_user.rank} {current_user.name} modified establishment record for station [{getattr(existing_est, 'station', 'N/A')}]."
        )

        return serialize_row(existing_est)
    except Exception as e:
        db.rollback()
        logs_db.rollback()
        print(f"Establishment update error: {str(e)}")
        raise HTTPException(status_code=500, detail=f"Failed to update establishment: {str(e)}")

# ====================================================================
# 4. DELETE ESTABLISHMENT
# ====================================================================
@router.delete("/establishments/{est_id}")
def delete_establishment(
    est_id: int, 
    db: Session = Depends(get_db), 
    logs_db: Session = Depends(get_logs_db),
    current_user: models.Users = Depends(get_current_user)
):
    EstModel = get_est_model()
    
    pk_col = getattr(EstModel, 'id', getattr(EstModel, 'sn', None))
    existing_est = db.query(EstModel).filter(pk_col == est_id).first()
    
    if not existing_est:
        raise HTTPException(status_code=404, detail="Establishment record not found.")
        
    perms = current_user.permissions or {}
    if current_user.role not in ["SUPER_ADMIN", "ADMIN", "RPC"] and not perms.get("delete_records", False):
        raise HTTPException(status_code=403, detail="Clearance Denied: Record deletion requires admin authorization.")

    try:
        db.delete(existing_est)
        db.commit()

        record_neon_activity(
            logs_db=logs_db,
            fnum=current_user.fnum,
            action_type="DELETE",
            module="ESTABLISHMENTS",
            target_id=str(est_id),
            changes_summary=f"{current_user.fnum} {current_user.rank} {current_user.name} permanently purged establishment record ID {est_id}."
        )

        return {"status": "success", "message": f"Establishment record {est_id} deleted successfully."}
    except Exception as e:
        db.rollback()
        logs_db.rollback()
        raise HTTPException(status_code=500, detail=f"Failed to delete establishment: {str(e)}")