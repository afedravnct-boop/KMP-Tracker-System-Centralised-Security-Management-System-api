# routers/general_documents.py
import io
import os
import json
import base64
import urllib.parse
import re
from datetime import datetime
from typing import Optional, List
import hmac
import hashlib

import boto3
import openpyxl
import pymupdf
import pytz
from docx import Document
from docx.shared import Pt
from docx.oxml import parse_xml
from pptx import Presentation
from pptx.util import Inches, Pt as PPTXPt
from pptx.dml.color import RGBColor as PPTXRGBColor

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form, Request
from fastapi.responses import StreamingResponse, JSONResponse
from sqlalchemy.orm import Session
from sqlalchemy import text

from app import models
from app.database import get_db, get_logs_db
from routers.activity_logger import record_neon_activity
from routers.auth import get_current_user  # 🟢 Properly imported at top level

router = APIRouter(prefix="/api/v1", tags=["General Documents"])

AWS_REGION = os.getenv("AWS_REGION", "eu-central-1")
BUCKET_NAME = os.getenv("AWS_BUCKET_NAME", "kmp-centralised-security-storage")
s3_client = boto3.client(
    "s3", 
    aws_access_key_id=os.getenv("AWS_ACCESS_KEY_ID"), 
    aws_secret_access_key=os.getenv("AWS_SECRET_ACCESS_KEY"), 
    region_name=AWS_REGION
)

def get_general_doc_model():
    for name in ['GeneralDocument', 'GeneralDocuments', 'General_Documents', 'general_documents']:
        if hasattr(models, name):
            return getattr(models, name)
    return None  # Return None safely instead of throwing 500

@router.get("/general-documents")
@router.get("/general-docs/list")
def get_general_documents(
    db: Session = Depends(get_db), 
    logs_db: Session = Depends(get_logs_db),
    current_user: models.Users = Depends(get_current_user)  # 🟢 Properly injected via Depends
):
    try:
        Model = get_general_doc_model()
        if not Model:
            return []

        docs = db.query(Model).order_by(Model.id.desc()).all()
        
        serialized_docs = []
        for d in docs:
            try:
                raw_date = getattr(d, 'upload_date', getattr(d, 'created_at', ''))
                if isinstance(raw_date, datetime):
                    date_str = raw_date.strftime("%Y-%m-%d")
                else:
                    date_str = str(raw_date).split(' ')[0] if raw_date else 'N/A'

                serialized_docs.append({
                    "id": getattr(d, 'id', 0),
                    "name": getattr(d, 'file_name', getattr(d, 'name', 'Document')),
                    "type": getattr(d, 'doc_type', getattr(d, 'type', 'General Document')),
                    "date": date_str,
                    "size": getattr(d, 'file_size', getattr(d, 'size', 'N/A')),
                    "file_path": getattr(d, 'file_path', getattr(d, 'url', '')),
                    "region": getattr(d, 'region', 'KMP HEADQUARTERS'),
                    "station": getattr(d, 'station', 'HQ')
                })
            except Exception as row_err:
                print(f"Skipping malformed document row: {row_err}")
                continue

        # Non-blocking activity logging
        if current_user:
            try:
                record_neon_activity(
                    logs_db=logs_db,
                    fnum=current_user.fnum,
                    action_type="VIEW",
                    module="GENERAL_DOCUMENTS",
                    target_id="ALL_DOCUMENTS",
                    changes_summary=f"{current_user.fnum} {current_user.rank} {current_user.name} accessed General Documents repository (Fetched {len(serialized_docs)} files)."
                )
            except Exception:
                pass

        return serialized_docs

    except Exception as db_err:
        import traceback
        traceback.print_exc()
        return []

@router.post("/general-documents/upload")
@router.post("/general-docs/upload")
async def upload_general_document(
    file: Optional[UploadFile] = File(None),
    files: Optional[List[UploadFile]] = File(None),
    target_region: Optional[str] = Form(None),
    target_station: Optional[str] = Form(None),
    db: Session = Depends(get_db),
    logs_db: Session = Depends(get_logs_db),
    current_user: models.Users = Depends(get_current_user)  # 🟢 Properly injected via Depends
):
    Model = get_general_doc_model()
    file_list = [f for f in [file, *(files or [])] if f is not None]
    if not file_list:
        raise HTTPException(status_code=400, detail="No files received.")

    uploaded_count = 0
    try:
        for f in file_list:
            contents = await f.read()
            file_size_kb = max(1, round(len(contents) / 1024))
            size_str = f"{file_size_kb} KB" if file_size_kb < 1024 else f"{round(file_size_kb / 1024, 1)} MB"

            s3_key = f"general_docs/{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}_{f.filename.replace(' ', '_')}"
            s3_client.put_object(
                Bucket=BUCKET_NAME, 
                Key=s3_key, 
                Body=contents, 
                ContentType=f.content_type or "application/octet-stream", 
                ServerSideEncryption="AES256"
            )
            url = f"https://{BUCKET_NAME}.s3.{AWS_REGION}.amazonaws.com/{s3_key}"

            new_doc = Model(
                file_name=f.filename,
                doc_type="General Document",
                file_size=size_str,
                file_path=url,
                region=target_region or current_user.region or "KMP HEADQUARTERS",
                station=target_station or current_user.station or "HQ",
                uploaded_by=current_user.fnum
            )
            db.add(new_doc)
            uploaded_count += 1

            record_neon_activity(
                logs_db=logs_db,
                fnum=current_user.fnum,
                action_type="REGISTER",
                module="GENERAL_DOCUMENTS",
                target_id=f.filename,
                changes_summary=f"{current_user.fnum} {current_user.rank} {current_user.name} uploaded document titled \"{f.filename}\"."
            )
            
        db.commit()

        return {"status": "success", "message": f"Successfully uploaded {uploaded_count} general document(s)."}
    except Exception as e:
        db.rollback()
        logs_db.rollback()
        print(f"Fetch error: {e}")
        raise HTTPException(status_code=500, detail=f"Database Error: {str(e)}")


@router.get("/general-documents/download/{doc_id}")
@router.get("/general-docs/download/{doc_id}")
def download_general_document(
    doc_id: int, 
    return_url: bool = False,
    db: Session = Depends(get_db),
    logs_db: Session = Depends(get_logs_db),
    current_user: models.Users = Depends(get_current_user)  # 🟢 Properly injected via Depends
):
    Model = get_general_doc_model()
    doc_record = db.query(Model).filter(Model.id == doc_id).first()

    if not doc_record:
        raise HTTPException(status_code=404, detail="General Document not found in system database.")
        
    file_path = getattr(doc_record, 'file_path', getattr(doc_record, 'url', ''))
    file_name = getattr(doc_record, 'file_name', getattr(doc_record, 'name', 'document'))

    if not str(file_path).startswith("http"):
        raise HTTPException(status_code=404, detail="File path invalid or missing S3 storage link.")

    parsed_url = urllib.parse.urlparse(file_path)
    original_s3_key = parsed_url.path.lstrip('/') 
    file_extension = file_name.lower().split('.')[-1] if '.' in file_name else "bin"

    try:
        file_stream = io.BytesIO()
        s3_client.download_fileobj(BUCKET_NAME, original_s3_key, file_stream)
        file_stream.seek(0)
        raw_bytes = file_stream.getvalue()

        eat_tz = pytz.timezone("Africa/Nairobi")
        eat_time = datetime.now(eat_tz).replace(tzinfo=None)
        timestamp_eat = eat_time.strftime("%Y-%m-%d %H:%M:%S EAT")

        officer_fnum = (current_user.fnum or "HQ-UNKNOWN").strip().upper()
        officer_rank = (current_user.rank or "OFFICER").strip().upper()
        officer_name = (current_user.name or "UNKNOWN").strip().upper()
        officer_signature = f"{officer_fnum} {officer_rank} {officer_name}"
        command_post = f"{current_user.station or 'KMP HEADQUARTERS'}, {current_user.region or 'KMP HEADQUARTERS'}"
        stamp_id = f"KMP-GEN-{officer_fnum}-{eat_time.strftime('%Y%m%d%H%M%S')}"

        FORENSIC_SECRET = os.getenv("FORENSIC_SIGNING_SECRET", "super_secret_kmp_key_2026").encode('utf-8')
        
        compact_payload = {"f": officer_fnum, "s": stamp_id}
        payload_str = json.dumps(compact_payload)
        
        signature = hmac.new(FORENSIC_SECRET, payload_str.encode('utf-8'), hashlib.sha256).hexdigest()
        secure_token_data = f"{payload_str}::{signature}"
        encoded_token = base64.b64encode(secure_token_data.encode('utf-8')).decode('utf-8')

        keywords_str = f"KMP_AUDIT;{encoded_token}"[:250]
        comments_str = f"Export: {officer_signature} [{command_post}]. ID: {stamp_id}"

        vertical_stamp_text = f"RESTRICTED LAW ENFORCEMENT RECORD  |  ACCESS BY: {officer_signature}  |  STAMP ID: {stamp_id}  |  TIME: {timestamp_eat}"

        output_stream = io.BytesIO()
        content_type = "application/octet-stream"

        if file_extension == 'docx':
            word_doc = Document(io.BytesIO(raw_bytes))
            core_props = word_doc.core_properties
            core_props.author = officer_signature
            core_props.last_modified_by = officer_signature
            core_props.keywords = keywords_str
            core_props.comments = comments_str
            core_props.category = "RESTRICTED / LAW ENFORCEMENT RECORD"

            section = word_doc.sections[0]
            header = section.header
            
            for p in header.paragraphs:
                p.text = ""

            header_p = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
            
            try:
                if return_url:
                    vml_xml = f'''
                    <w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" 
                         xmlns:v="urn:schemas-microsoft-com:vml" 
                         xmlns:o="urn:schemas-microsoft-com:office:office">
                        <w:pict>
                            <v:rect id="LiveViewStamp" 
                                    style="position:absolute;left:40pt;top:72pt;width:25pt;height:650pt;z-index:251657216;mso-position-horizontal-relative:page;mso-position-vertical-relative:page" 
                                    strokecolor="#CC0000" stroked="f" filled="f">
                            <v:textbox style="mso-layout-flow-alt:bottom-to-top" inset="0,0,0,0">
                                <w:txbxContent>
                                    <w:p>
                                        <w:pPr>
                                            <w:textDirection w:val="btLr"/>
                                            <w:jc w:val="center"/>
                                        </w:pPr>
                                        <w:r>
                                            <w:rPr>
                                                <w:color w:val="CC0000"/>
                                                <w:rFonts w:ascii="Courier New" w:hAnsi="Courier New"/>
                                                <w:sz w:val="14"/>
                                                <w:b/>
                                            </w:rPr>
                                            <w:t>{vertical_stamp_text}</w:t>
                                        </w:r>
                                    </w:p>
                                </w:txbxContent>
                            </v:textbox>
                        </v:rect>
                        </w:pict>
                    </w:r>
                    '''
                else:
                    vml_xml = f'''
                    <w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" 
                         xmlns:v="urn:schemas-microsoft-com:vml" 
                         xmlns:o="urn:schemas-microsoft-com:office:office">
                        <w:pict>
                            <v:rect id="DownloadStamp" 
                                    style="position:absolute;left:40pt;top:72pt;width:25pt;height:650pt;z-index:251657216;mso-position-horizontal-relative:page;mso-position-vertical-relative:page" 
                                    strokecolor="#CC0000" stroked="f" filled="f">
                            <v:textbox style="mso-layout-flow-alt:bottom-to-top" inset="0,0,0,0">
                                <w:txbxContent>
                                    <w:p>
                                        <w:pPr>
                                            <w:jc w:val="center"/>
                                        </w:pPr>
                                        <w:r>
                                            <w:rPr>
                                                <w:color w:val="CC0000"/>
                                                <w:rFonts w:ascii="Courier New" w:hAnsi="Courier New"/>
                                                <w:sz w:val="14"/>
                                                <w:b/>
                                            </w:rPr>
                                            <w:t>{vertical_stamp_text}</w:t>
                                        </w:r>
                                    </w:p>
                                </w:txbxContent>
                            </v:textbox>
                        </v:rect>
                        </w:pict>
                    </w:r>
                    '''

                vml_run = parse_xml(vml_xml)
                header_p._p.append(vml_run)
                
            except Exception as e:
                print(f"Failed to inject DOCX header watermark: {e}")
            
            word_doc.save(output_stream)
            content_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

        elif file_extension == 'pdf':
            try:
                pdf_doc = pymupdf.open(stream=raw_bytes, filetype="pdf")
                for page in pdf_doc:
                    rect = page.rect
                    page.insert_text(
                        pymupdf.Point(20, rect.height - 50),
                        vertical_stamp_text,
                        fontsize=7,
                        fontname="courier-bold",
                        color=(0.545, 0, 0), 
                        rotate=90,
                        overlay=True
                    )
                stamped_pdf_bytes = pdf_doc.tobytes()
                output_stream = io.BytesIO(stamped_pdf_bytes)
                content_type = "application/pdf"
            except Exception as pdf_err:
                print(f"PDF watermarking notice: {pdf_err}")
                output_stream.write(raw_bytes)
                content_type = "application/pdf"

        elif file_extension in ['xlsx', 'xls']:
            wb = openpyxl.load_workbook(io.BytesIO(raw_bytes))
            wb.properties.creator = officer_signature
            wb.properties.lastModifiedBy = officer_signature
            wb.properties.keywords = keywords_str
            wb.properties.description = comments_str
            wb.properties.category = "RESTRICTED / FORENSIC POLICE RECORD"
            
            for ws in wb.worksheets:
                if hasattr(ws, 'sheet_footer'): 
                    ws.sheet_footer.left.text = f"&\"Courier New,Bold\"&7&KCC0000{vertical_stamp_text}"
                elif hasattr(ws, 'odd_footer'): 
                    ws.odd_footer.left.text = f"&\"Courier New,Bold\"&7&KCC0000{vertical_stamp_text}"
            wb.save(output_stream)
            content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

        elif file_extension in ['pptx', 'ppt']:
            try:
                prs = Presentation(io.BytesIO(raw_bytes))
                prs.core_properties.author = officer_signature
                prs.core_properties.last_modified_by = officer_signature
                prs.core_properties.keywords = keywords_str
                prs.core_properties.comments = comments_str
                prs.core_properties.category = "RESTRICTED / FORENSIC POLICE RECORD"
                
                if prs.slides:
                    slide = prs.slides[0]
                    left_box = slide.shapes.add_textbox(Inches(0.2), Inches(0.2), Inches(8), Inches(1))
                    left_box.rotation = 270 
                    p_left = left_box.text_frame.add_paragraph()
                    p_left.text = vertical_stamp_text
                    p_left.font.size = PPTXPt(6.5)
                    p_left.font.bold = True
                    p_left.font.name = 'Courier New'
                    p_left.font.color.rgb = PPTXRGBColor(139, 0, 0)

                prs.save(output_stream)
                content_type = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
            except Exception:
                output_stream.write(raw_bytes)
                content_type = "application/vnd.openxmlformats-officedocument.presentationml.presentation"

        else:
            output_stream.write(raw_bytes)

        output_stream.seek(0)
        final_bytes = output_stream.getvalue()

        record_neon_activity(
            logs_db=logs_db,
            fnum=current_user.fnum,
            action_type="UPDATE",
            module="GENERAL_DOCUMENTS",
            target_id=str(doc_id),
            changes_summary=f"{current_user.fnum} {current_user.rank} {current_user.name} securely downloaded and watermarked document titled \"{file_name}\"."
        )

        if return_url:
            safe_file_name = re.sub(r'[^a-zA-Z0-9.]', '_', file_name)
            safe_file_name = re.sub(r'_+', '_', safe_file_name)
            
            temp_s3_key = f"forensic_cache/{stamp_id}_{safe_file_name}"
            
            s3_client.put_object(
                Bucket=BUCKET_NAME,
                Key=temp_s3_key,
                Body=final_bytes,
                ContentType=content_type,
                ContentDisposition=f'inline; filename="{safe_file_name}"', 
                ServerSideEncryption="AES256"
            )
            
            try:
                waiter = s3_client.get_waiter('object_exists')
                waiter.wait(Bucket=BUCKET_NAME, Key=temp_s3_key, WaiterConfig={'Delay': 1, 'MaxAttempts': 5})
            except Exception:
                import time
                time.sleep(1.5)
            
            presigned_url = s3_client.generate_presigned_url(
                'get_object',
                Params={
                    'Bucket': BUCKET_NAME, 
                    'Key': temp_s3_key,
                    'ResponseContentDisposition': f'inline; filename="{safe_file_name}"',
                    'ResponseContentType': content_type
                },
                ExpiresIn=3600
            )
            return JSONResponse(content={"url": presigned_url})

        return StreamingResponse(
            io.BytesIO(final_bytes),
            media_type=content_type,
            headers={"Content-Disposition": f'attachment; filename="{file_name}"'}
        )

    except Exception as e:
        raise HTTPException(status_code=500, detail=f"General Document Download Error: {str(e)}")