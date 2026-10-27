import io
import re
from datetime import datetime, date, timedelta
from typing import Optional, List

import pandas as pd
import openpyxl
from openpyxl.styles import Alignment, PatternFill, Font, Border, Side
import pyzipper

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import or_, and_, func, text

from app import models
from app.database import get_db, get_logs_db
from auth import get_current_user
from routers.activity_logger import record_neon_activity

router = APIRouter(prefix="/api/v1", tags=["Consolidated Ledger"])

def clean_str(val) -> str:
    if pd.isna(val) or not val: return ''
    return str(val).strip().upper()

def safe_int(val) -> int:
    if pd.isna(val) or val is None: return 0
    if isinstance(val, (int, float)): return int(val)
    s = str(val).strip()
    if not s: return 0
    digits = re.sub(r'[^0-9]', '', s)
    return int(digits) if digits else 0

# 🟢 ROBUST DATE PARSER
def parse_safe_date(val) -> Optional[date]:
    """Safely converts any weird date string, timestamp, or excel float into a pure Date object."""
    if pd.isna(val) or val is None: return None
    if isinstance(val, date) and not isinstance(val, datetime):
        return val if 1900 <= val.year <= 2100 else None
    if isinstance(val, datetime):
        return val.date() if 1900 <= val.year <= 2100 else None
    if type(val).__name__ == 'Timestamp':
        try: return val.date() if 1900 <= val.year <= 2100 else None
        except Exception: return None
    
    val_str = str(val).strip()
    if val_str.lower() in ['nan', 'nat', 'none', 'null', '', '-', 'n/a', 'nil', '0', 'undefined', 'all time', 'all']: return None
    if ' ' in val_str: val_str = val_str.split(' ')[0]
    if 'T' in val_str: val_str = val_str.split('T')[0]
    val_str = val_str.strip()
    if not val_str: return None
    
    try:
        if val_str.replace('.', '', 1).isdigit():
            float_val = float(val_str)
            if 1000 < float_val < 73050:
                try:
                    dt = pd.to_datetime(float_val, unit='D', origin='1899-12-30', errors='coerce')
                    if pd.notna(dt) and 1900 <= dt.year <= 2100: return dt.date()
                except Exception: pass
    except Exception: pass
    
    clean_str_val = re.sub(r'[\./\\]', '-', val_str)
    parts = clean_str_val.split('-')
    if len(parts) == 3:
        p0, p1, p2 = parts[0].strip(), parts[1].strip(), parts[2].strip()
        if p0.isdigit() and p1.isdigit() and p2.isdigit():
            if len(p0) == 4:
                try:
                    res_date = date(int(p0), int(p1), int(p2))
                    if 1900 <= res_date.year <= 2100: return res_date
                except Exception: pass
            elif len(p2) == 4:
                try:
                    res_date = date(int(p2), int(p1), int(p0))
                    if 1900 <= res_date.year <= 2100: return res_date
                except Exception: pass
            elif len(p0) <= 2 and len(p2) <= 2:
                for (day_val, month_val, year_val) in [(int(p0), int(p1), int(p2)), (int(p2), int(p1), int(p0))]:
                    y = year_val
                    if y < 100: y = 1900 + y if y > 40 else 2000 + y
                    try:
                        res_date = date(y, month_val, day_val)
                        if 1900 <= res_date.year <= 2100: return res_date
                    except Exception: continue
                    
    for fmt in ('%Y-%m-%d', '%d-%m-%Y', '%d/%m/%Y', '%Y/%m/%d', '%d.%m.%Y', '%m/%d/%Y', '%m-%d-%Y', '%d-%b-%y', '%d-%b-%Y'):
        try:
            d = datetime.strptime(val_str, fmt).date()
            if 1900 <= d.year <= 2100: return d
        except Exception: continue
        
    try:
        parsed = pd.to_datetime(val_str, dayfirst=True, errors='coerce')
        if pd.notna(parsed):
            py_dt = parsed.to_pydatetime()
            if 1900 <= py_dt.year <= 2100: return py_dt.date()
    except Exception: pass
    return None

def get_row_date(row) -> Optional[date]:
    """Inspects all possible date or timestamp attributes on a model instance."""
    for attr in ('date', 'created_at', 'timestamp', 'created', 'updated_at'):
        if hasattr(row, attr):
            val = getattr(row, attr, None)
            parsed = parse_safe_date(val)
            if parsed:
                return parsed
    return None

def normalize_region_name(region_str) -> str:
    if pd.isna(region_str) or not region_str: return "KMP HEADQUARTERS"
    r = clean_str(region_str)
    if r in ["KMPN", "KMP/N", "KMP NORTH"]: return "KMP NORTH"
    if r in ["KMPS", "KMP/S", "KMP SOUTH"]: return "KMP SOUTH"
    if r in ["KMPE", "KMP/E", "KMP EAST"]: return "KMP EAST"
    if r in ["KMP HQ", "KMP HEADQUARTERS"]: return "KMP HEADQUARTERS"
    if r in ["POLICE HQ", "POLICE HEADQUARTERS", "NAGURU"]: return "POLICE HEADQUARTERS"
    return r

REGIONAL_HIERARCHY = {
    "KMP NORTH": ["KMP NORTH HEADQUARTERS", "KMP NORTH", "KAWEMPE", "KAKIRI", "KASANGATI", "MATUGGA", "NANSANA", "OLD KAMPALA", "WAKISO", "WANDEGEYA"],
    "KMP EAST": ["KMP EAST HEADQUARTERS", "KMP EAST", "JINJA ROAD", "KIRA", "KIRA DIV", "KIRA ROAD", "MUKONO", "NAGGALAMA", "SEETA"],
    "KMP SOUTH": ["KMP SOUTH HEADQUARTERS", "KMP SOUTH", "NATEETE", "CPS KAMPALA", "PARLIAMENT", "ENTEBBE", "KABALAGALA", "KAJJANSI", "KASENYI", "KATWE", "KYENGERA", "NSANGI"],
    "KMP HEADQUARTERS": ["KMP HEADQUARTERS", "KMP CID", "KMP TRAFFIC", "KMP ICT", "KMP FLYING SQUAD", "KMP CRIME INTELLIGENCE"],
    "POLICE HEADQUARTERS": ["NAGURU", "OPERATIONS", "CRIME INTELLIGENCE", "CID", "LOGISTICS & ENGINEERING", "ICT", "CT", "FIRE & RESCUE"]
}

def get_base_station_name(stn: str) -> str:
    if pd.isna(stn) or not stn: return "UNKNOWN"
    val = clean_str(stn)
    if val in ["OLD KLA", "OK", "OLD K", "O/K", "OLD KAMPALA DIVISION", "OLD KAMPALA"]: return "OLD KAMPALA"
    if val in ["CPS KLA", "CPS KAMPALA"]: return "CPS KAMPALA"
    val = re.sub(r'\s+(P\.?S\.?|P/S|PS|P\s*STN|P/STN|POL\s*STN|POLICE\s*STATION|DIVISION|P\.?P\.?|P/P|P/POST|P\s*POST|PP|P|POLICE\s*POST|HEADQUARTERS|HQ)$', '', val)
    return val.strip()

def is_station_equivalent(stat_a, stat_b) -> bool:
    base_a = get_base_station_name(stat_a)
    base_b = get_base_station_name(stat_b)
    if not base_a or not base_b: return False
    return base_a == base_b

def get_official_region_for_station(station_name: str, db_region: Optional[str] = None) -> str:
    base_station = get_base_station_name(station_name)
    clean_db_region = normalize_region_name(db_region)

    if REGIONAL_HIERARCHY.get(clean_db_region):
        base_hier = [get_base_station_name(s) for s in REGIONAL_HIERARCHY[clean_db_region]]
        if base_station in base_hier:
            return clean_db_region

    for region_name, stations_list in REGIONAL_HIERARCHY.items():
        base_hier = [get_base_station_name(s) for s in stations_list]
        if base_station in base_hier:
            return region_name

    return clean_db_region if clean_db_region in REGIONAL_HIERARCHY else (clean_db_region or 'KMP GENERAL')

def get_model_safe(*names):
    for name in names:
        if hasattr(models, name):
            return getattr(models, name)
    return None

def apply_opsec_scope(current_user, query, ModelClass):
    if not ModelClass or not current_user:
        return query.filter(text("1=0"))
    
    user_role = str(current_user.role).strip().upper() if current_user.role else ""
    user_pos = str(current_user.position).strip().upper() if current_user.position else ""
    user_reg = str(current_user.region).strip().upper() if current_user.region else ""
    user_stn = str(current_user.station).strip().upper() if current_user.station else ""

    perms = current_user.permissions or {}
    if isinstance(perms, str):
        try: perms = json.loads(perms)
        except Exception: perms = {}

    is_absolute_global = (
        user_role in ["SUPER_ADMIN", "ADMIN", "ASSISTANT_SUPER_ADMIN"] or
        "KMP COMMANDER" in user_pos or
        "DEPUTY KMP COMMANDER" in user_pos or
        "KMP ADMIN" in user_pos or
        perms.get("view_global_roster") is True or
        perms.get("global_observer") is True
    )

    is_kmp_sys_mgr = (user_role == "SYSTEM_MANAGER" and user_reg in ["KMP HEADQUARTERS", "POLICE HEADQUARTERS"] and "KMP" in user_pos)
    is_kmp_specialist = (user_role == "ASSISTANT_SYSTEM_MANAGER" and user_reg in ["KMP HEADQUARTERS", "POLICE HEADQUARTERS"] and "KMP" in user_pos)
    
    is_regional_command = (
        user_role in ["RPC", "DEPUTY_RPC", "SYSTEM_MANAGER", "ASSISTANT_SYSTEM_MANAGER", "REGIONAL_ADMIN", "ASSISTANT_REGIONAL_ADMIN", "DIVISION_ADMIN"] and
        not is_kmp_sys_mgr and not is_kmp_specialist
    )

    if is_absolute_global or is_kmp_sys_mgr:
        return query
        
    elif is_kmp_specialist and hasattr(ModelClass, 'region'):
        return query.filter(func.upper(ModelClass.region) == user_reg)
        
    elif is_regional_command or user_reg in REGIONAL_HIERARCHY:
        conds = []
        if hasattr(ModelClass, 'region'):
            conds.append(func.upper(ModelClass.region) == user_reg)
        
        if hasattr(ModelClass, 'station') and user_reg in REGIONAL_HIERARCHY:
            expanded_stns = set()
            for s in REGIONAL_HIERARCHY[user_reg]:
                expanded_stns.add(s)
                expanded_stns.add(s.replace(' HEADQUARTERS', '').replace(' HQ', ''))
                expanded_stns.add(s + ' HEADQUARTERS')
                expanded_stns.add(s + ' HQ')
            conds.append(func.upper(ModelClass.station).in_(list(expanded_stns)))
            
        if conds: return query.filter(or_(*conds))
        return query.filter(text("1=0"))
        
    elif hasattr(ModelClass, 'station') and hasattr(ModelClass, 'region'):
        clean_user_stn = get_base_station_name(user_stn)
        return query.filter(
            and_(
                func.upper(ModelClass.region) == user_reg,
                or_(
                    func.upper(ModelClass.station) == user_stn,
                    func.upper(ModelClass.station).ilike(f"{clean_user_stn}%")
                )
            )
        )
    return query.filter(text("1=0"))

def get_ledger_aggregated_data(db: Session, start_dt: Optional[date], end_dt: Optional[date], filter_region: str, filter_station: str, current_user: models.Users):
    region_order = ["KMP NORTH", "KMP EAST", "KMP SOUTH", "KMP HEADQUARTERS", "POLICE HEADQUARTERS"]

    crime_regional, crime_general = {}, {}
    crime_regional_totals = {}
    grand_crime_cases, grand_crime_suspects = 0, 0
    ops_regional = {}
    ops_general = {k: 0 for k in ['arrested', 'given_bond', 'cautioned', 'pending_court', 'taken_to_court', 'released', 'remanded', 'convicted']}
    story_regional = {}
    grand_stories = 0

    station_cache = {}
    def fast_get_region(stn_str, db_reg_str):
        key = f"{stn_str}_{db_reg_str}"
        if key not in station_cache:
            station_cache[key] = get_official_region_for_station(stn_str, db_reg_str)
        return station_cache[key]

    # 🟢 1. CRIME REPORTS (Query Model Instances directly with robust date validation)
    try:
        CrimeModel = get_model_safe('Crime_Reports', 'CrimeReports', 'crime_reports', 'Reports', 'reports')
        if CrimeModel:
            query = db.query(CrimeModel)
            query = apply_opsec_scope(current_user, query, CrimeModel)
            crimes = query.all()

            for row in crimes:
                row_date = get_row_date(row)
                if start_dt and end_dt:
                    if not row_date or row_date < start_dt or row_date > end_dt:
                        continue
                elif start_dt and not end_dt:
                    if not row_date or row_date < start_dt:
                        continue
                elif not start_dt and end_dt:
                    if not row_date or row_date > end_dt:
                        continue

                stn = clean_str(getattr(row, 'station', ''))
                reg = fast_get_region(stn, getattr(row, 'region', ''))

                if filter_region and filter_region != "ALL REGIONS" and reg != filter_region: continue
                if filter_station and filter_station != "ALL STATIONS" and not is_station_equivalent(stn, filter_station): continue

                category = clean_str(getattr(row, 'category', 'GENERAL CRIMES'))
                offence = clean_str(getattr(row, 'offence', 'UNSPECIFIED INCIDENT'))
                off = f"AGRIC: {offence}" if category == 'AGRIC_CRIME' else offence
                if "LOCK-UP TOTAL" in off: continue

                suspects = safe_int(getattr(row, 'suspects', 0))

                if reg not in crime_regional: crime_regional[reg] = {}
                if off not in crime_regional[reg]: crime_regional[reg][off] = {"cases": 0, "suspects": 0}
                
                crime_regional[reg][off]["cases"] += 1
                crime_regional[reg][off]["suspects"] += suspects

                if reg not in crime_regional_totals: crime_regional_totals[reg] = {"cases": 0, "suspects": 0}
                crime_regional_totals[reg]["cases"] += 1
                crime_regional_totals[reg]["suspects"] += suspects

                if off not in crime_general: crime_general[off] = {"cases": 0, "suspects": 0}
                crime_general[off]["cases"] += 1
                crime_general[off]["suspects"] += suspects

                grand_crime_cases += 1
                grand_crime_suspects += suspects
    except Exception as e:
        print(f"CRIMES FETCH ERROR: {e}")

    # 🟢 2. OPERATIONAL STATISTICS (Query Model Instances directly with robust date validation)
    try:
        OpsModel = get_model_safe('Operational_Statistics', 'OperationalStatistics', 'operationl_statistics', 'Disruptive_Operations', 'Stats')
        if OpsModel:
            query = db.query(OpsModel)
            query = apply_opsec_scope(current_user, query, OpsModel)
            stats = query.all()
            
            ops_keys = ['arrested', 'given_bond', 'cautioned', 'pending_court', 'taken_to_court', 'released', 'remanded', 'convicted']

            for row in stats:
                row_date = get_row_date(row)
                if start_dt and end_dt:
                    if not row_date or row_date < start_dt or row_date > end_dt:
                        continue
                elif start_dt and not end_dt:
                    if not row_date or row_date < start_dt:
                        continue
                elif not start_dt and end_dt:
                    if not row_date or row_date > end_dt:
                        continue

                stn = clean_str(getattr(row, 'station', ''))
                reg = fast_get_region(stn, getattr(row, 'region', ''))

                if filter_region and filter_region != "ALL REGIONS" and reg != filter_region: continue
                if filter_station and filter_station != "ALL STATIONS" and not is_station_equivalent(stn, filter_station): continue

                if reg not in ops_regional: ops_regional[reg] = {k: 0 for k in ops_general.keys()}
                
                for key in ops_keys:
                    val = safe_int(getattr(row, key, 0))
                    ops_regional[reg][key] += val
                    ops_general[key] += val
    except Exception as e:
        print(f"OPS FETCH ERROR: {e}")

    # 🟢 3. SUCCESS STORIES (Query Model Instances directly with robust date validation)
    try:
        StoryModel = get_model_safe('Success_Stories', 'SuccessStories', 'success_stories', 'Stories', 'stories')
        if StoryModel:
            query = db.query(StoryModel)
            query = apply_opsec_scope(current_user, query, StoryModel)
            stories = query.all()

            for row in stories:
                row_date = get_row_date(row)
                if start_dt and end_dt:
                    if not row_date or row_date < start_dt or row_date > end_dt:
                        continue
                elif start_dt and not end_dt:
                    if not row_date or row_date < start_dt:
                        continue
                elif not start_dt and end_dt:
                    if not row_date or row_date > end_dt:
                        continue

                stn = clean_str(getattr(row, 'station', ''))
                reg = fast_get_region(stn, getattr(row, 'region', ''))

                if filter_region and filter_region != "ALL REGIONS" and reg != filter_region: continue
                if filter_station and filter_station != "ALL STATIONS" and not is_station_equivalent(stn, filter_station): continue

                story_regional[reg] = story_regional.get(reg, 0) + 1
                grand_stories += 1
    except Exception as e:
        print(f"STORIES FETCH ERROR: {e}")

    # Ensure all regions exist in mapping even if empty
    for reg in region_order:
        if reg not in crime_regional: crime_regional[reg] = {}
        if reg not in crime_regional_totals: crime_regional_totals[reg] = {"cases": 0, "suspects": 0}
        if reg not in ops_regional: ops_regional[reg] = {k: 0 for k in ops_general.keys()}
        if reg not in story_regional: story_regional[reg] = 0

    return {
        "crimeRegional": crime_regional,
        "crimeRegionalTotals": crime_regional_totals,
        "crimeGeneral": crime_general,
        "sortedCrimeRegions": region_order,
        "grandCrimeCases": grand_crime_cases,
        "grandCrimeSuspects": grand_crime_suspects,
        "opsRegional": ops_regional,
        "opsGeneral": ops_general,
        "sortedOpsRegions": region_order,
        "storyRegional": story_regional,
        "sortedStoryRegions": region_order,
        "grandStories": grand_stories
    }

@router.get("/reports/consolidated-ledger")
def get_consolidated_ledger(
    start_date: Optional[str] = Query(default=None),
    end_date: Optional[str] = Query(default=None),
    region: Optional[str] = Query("ALL REGIONS"),
    station: Optional[str] = Query("ALL STATIONS"),
    db: Session = Depends(get_db),
    logs_db: Session = Depends(get_logs_db),
    current_user: models.Users = Depends(get_current_user)
):
    try:
        start_dt = parse_safe_date(start_date)
        end_dt = parse_safe_date(end_date)
        is_all_time = not start_dt or not end_dt

        data = get_ledger_aggregated_data(db, start_dt, end_dt, region, station, current_user)

        time_range = "ALL TIME" if is_all_time else f"{start_dt.strftime('%d %b %Y')} to {end_dt.strftime('%d %b %Y')}"
        record_neon_activity(
            logs_db=logs_db,
            fnum=current_user.fnum,
            action_type="VIEW",
            module="CONSOLIDATED_LEDGER",
            target_id="MASTER_SUMMARY",
            changes_summary=f"{current_user.fnum} {current_user.rank} compiled consolidated ledger ({time_range})."
        )

        return {"status": "success", "data": data}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Consolidated Ledger Data Fetch Failed: {str(e)}")

@router.get("/reports/consolidated-ledger/export")
def export_consolidated_ledger(
    start_date: Optional[str] = Query(default=None),
    end_date: Optional[str] = Query(default=None),
    region: Optional[str] = Query("ALL REGIONS"),
    station: Optional[str] = Query("ALL STATIONS"),
    db: Session = Depends(get_db),
    logs_db: Session = Depends(get_logs_db),
    current_user: models.Users = Depends(get_current_user)
):
    try:
        start_dt = parse_safe_date(start_date)
        end_dt = parse_safe_date(end_date)
        is_all_time = not start_dt or not end_dt

        data = get_ledger_aggregated_data(db, start_dt, end_dt, region, station, current_user)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Consolidated Ledger"
        ws.views.sheetView[0].showGridLines = True

        wb.properties.creator = f"{current_user.rank} {current_user.name} (FNUM: {current_user.fnum})"
        wb.properties.title = "KMP Consolidated Master Ledger"
        time_range = "ALL TIME" if is_all_time else f"{start_dt.strftime('%d %b %Y')} to {end_dt.strftime('%d %b %Y')}"
        wb.properties.subject = f"Jurisdiction: {region} / {station} ({time_range})"
        wb.properties.company = "Uganda Police Force - Kampala Metropolitan Police"

        thin_border = Border(
            left=Side(style='thin', color='A6A6A6'),
            right=Side(style='thin', color='A6A6A6'),
            top=Side(style='thin', color='A6A6A6'),
            bottom=Side(style='thin', color='A6A6A6')
        )

        row_idx = 1

        # --- SECTION 1: CRIMES & AGRICULTURAL THEFTS ---
        ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=5)
        cell = ws.cell(row=row_idx, column=1, value=f"SUMMARY OF CRIMES & AGRICULTURAL THEFTS REGISTERED ({time_range})")
        cell.fill = PatternFill(start_color="B4C6E7", end_color="B4C6E7", fill_type="solid")
        cell.font = Font(bold=True, size=11, color="000000")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        row_idx += 1

        headers_s1 = ["REGION", "NO", "OFFENCE / PRODUCE THEFT", "NUMBER OF CASES / QTY", "SUSPECTS IN CUSTODY"]
        for col_num, h in enumerate(headers_s1, 1):
            c = ws.cell(row=row_idx, column=col_num, value=h)
            c.fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
            c.font = Font(bold=True, size=10)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = thin_border
        row_idx += 1

        romans = ['i', 'ii', 'iii', 'iv', 'v', 'vi', 'vii', 'viii', 'ix', 'x', 'xi', 'xii', 'xiii', 'xiv', 'xv']
        for reg in data["sortedCrimeRegions"]:
            offences = list(data["crimeRegional"][reg].keys())
            reg_cases, reg_suspects = 0, 0

            if not offences:
                ws.cell(row=row_idx, column=1, value=reg).border = thin_border
                ws.cell(row=row_idx, column=2, value="-").border = thin_border
                ws.cell(row=row_idx, column=3, value="NO RECORDS").border = thin_border
                ws.cell(row=row_idx, column=4, value=0).border = thin_border
                ws.cell(row=row_idx, column=5, value=0).border = thin_border
                row_idx += 1
            else:
                for idx, off in enumerate(offences):
                    item = data["crimeRegional"][reg][off]
                    reg_cases += item["cases"]
                    reg_suspects += item["suspects"]

                    ws.cell(row=row_idx, column=1, value=reg if idx == 0 else "").border = thin_border
                    ws.cell(row=row_idx, column=2, value=f"{romans[idx]}." if idx < len(romans) else f"{idx+1}.").alignment = Alignment(horizontal="center")
                    ws.cell(row=row_idx, column=2).border = thin_border
                    
                    c_off = ws.cell(row=row_idx, column=3, value=off)
                    c_off.alignment = Alignment(horizontal="left")
                    c_off.border = thin_border

                    c_cases = ws.cell(row=row_idx, column=4, value=item["cases"])
                    c_cases.alignment = Alignment(horizontal="center")
                    c_cases.border = thin_border

                    c_susp = ws.cell(row=row_idx, column=5, value=item["suspects"])
                    c_susp.alignment = Alignment(horizontal="center")
                    c_susp.border = thin_border
                    row_idx += 1

            ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=3)
            sub_lbl = ws.cell(row=row_idx, column=1, value="SUB-TOTAL")
            sub_lbl.alignment = Alignment(horizontal="left", indent=2)
            sub_lbl.font = Font(bold=True)
            for c_i in range(1, 4): ws.cell(row=row_idx, column=c_i).border = thin_border

            c_sub_c = ws.cell(row=row_idx, column=4, value=reg_cases)
            c_sub_c.font = Font(bold=True)
            c_sub_c.alignment = Alignment(horizontal="center")
            c_sub_c.border = thin_border

            c_sub_s = ws.cell(row=row_idx, column=5, value=reg_suspects)
            c_sub_s.font = Font(bold=True)
            c_sub_s.alignment = Alignment(horizontal="center")
            c_sub_s.border = thin_border
            for c_i in range(1, 6): ws.cell(row=row_idx, column=c_i).fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
            row_idx += 2

        # --- SECTION 2: GENERAL CRIME & AGGREGATE SUMMARY ---
        ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=4)
        cell = ws.cell(row=row_idx, column=1, value="GENERAL CRIME & AGRICULTURAL PRODUCE SUMMARY")
        cell.fill = PatternFill(start_color="C6E0B4", end_color="C6E0B4", fill_type="solid")
        cell.font = Font(bold=True, size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        row_idx += 1

        headers_s2 = ["NO", "OFFENCES / AGRICULTURAL THEFTS", "TOTAL CASES / QTY", "SUSPECTS IN CUSTODY"]
        for col_num, h in enumerate(headers_s2, 1):
            c = ws.cell(row=row_idx, column=col_num, value=h)
            c.fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
            c.font = Font(bold=True, size=10)
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = thin_border
        row_idx += 1

        if data["crimeGeneral"]:
            for idx, (off, item) in enumerate(data["crimeGeneral"].items(), 1):
                ws.cell(row=row_idx, column=1, value=idx).alignment = Alignment(horizontal="center")
                ws.cell(row=row_idx, column=1).border = thin_border
                ws.cell(row=row_idx, column=2, value=off).border = thin_border
                ws.cell(row=row_idx, column=3, value=item["cases"]).alignment = Alignment(horizontal="center")
                ws.cell(row=row_idx, column=3).border = thin_border
                ws.cell(row=row_idx, column=4, value=item["suspects"]).alignment = Alignment(horizontal="center")
                ws.cell(row=row_idx, column=4).border = thin_border
                row_idx += 1
        else:
            ws.cell(row=row_idx, column=1, value=1).border = thin_border
            ws.cell(row=row_idx, column=2, value="NO GENERAL CRIMES RECORDED").border = thin_border
            ws.cell(row=row_idx, column=3, value=0).border = thin_border
            ws.cell(row=row_idx, column=4, value=0).border = thin_border
            row_idx += 1

        ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=2)
        gt_lbl = ws.cell(row=row_idx, column=1, value="GRAND TOTAL")
        gt_lbl.font = Font(bold=True)
        for c_i in range(1, 3): ws.cell(row=row_idx, column=c_i).border = thin_border

        gt_c = ws.cell(row=row_idx, column=3, value=data["grandCrimeCases"])
        gt_c.font = Font(bold=True)
        gt_c.alignment = Alignment(horizontal="center")
        gt_c.border = thin_border

        gt_s = ws.cell(row=row_idx, column=4, value=data["grandCrimeSuspects"])
        gt_s.font = Font(bold=True)
        gt_s.alignment = Alignment(horizontal="center")
        gt_s.border = thin_border
        for c_i in range(1, 5): ws.cell(row=row_idx, column=c_i).fill = PatternFill(start_color="A9D08E", end_color="A9D08E", fill_type="solid")
        row_idx += 2

        # --- SECTION 3: DISRUPTIVE OPERATIONS ---
        ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=9)
        cell = ws.cell(row=row_idx, column=1, value=f"SUMMARY OF DISRUPTIVE OPERATIONS ({time_range})")
        cell.fill = PatternFill(start_color="FFD966", end_color="FFD966", fill_type="solid")
        cell.font = Font(bold=True, size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        row_idx += 1

        headers_ops = ["REGION", "ARRESTED", "GIVEN BOND", "CAUTIONED", "PENDING COURT", "TAKEN TO COURT", "RELEASED", "REMANDED", "CONVICTED"]
        for col_num, h in enumerate(headers_ops, 1):
            c = ws.cell(row=row_idx, column=col_num, value=h)
            c.fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
            c.font = Font(bold=True, size=9)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = thin_border
        row_idx += 1

        for reg in data["sortedOpsRegions"]:
            item = data["opsRegional"][reg]
            ws.cell(row=row_idx, column=1, value=reg).border = thin_border
            vals = [item["arrested"], item["given_bond"], item["cautioned"], item["pending_court"], item["taken_to_court"], item["released"], item["remanded"], item["convicted"]]
            for c_idx, val in enumerate(vals, 2):
                c = ws.cell(row=row_idx, column=c_idx, value=val)
                c.alignment = Alignment(horizontal="center")
                c.border = thin_border
            row_idx += 1

        ws.cell(row=row_idx, column=1, value="SELECTED JURISDICTION").font = Font(bold=True)
        ws.cell(row=row_idx, column=1).border = thin_border
        gen_vals = [data["opsGeneral"]["arrested"], data["opsGeneral"]["given_bond"], data["opsGeneral"]["cautioned"], data["opsGeneral"]["pending_court"], data["opsGeneral"]["taken_to_court"], data["opsGeneral"]["released"], data["opsGeneral"]["remanded"], data["opsGeneral"]["convicted"]]
        for c_idx, val in enumerate(gen_vals, 2):
            c = ws.cell(row=row_idx, column=c_idx, value=val)
            c.font = Font(bold=True)
            c.alignment = Alignment(horizontal="center")
            c.border = thin_border
        for c_i in range(1, 10): ws.cell(row=row_idx, column=c_i).fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
        row_idx += 2

        # --- SECTION 4: SUCCESS STORIES ---
        ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=3)
        cell = ws.cell(row=row_idx, column=1, value=f"SUMMARY OF SUCCESSFUL OPERATIONS & MILESTONES ({time_range})")
        cell.fill = PatternFill(start_color="FFE699", end_color="FFE699", fill_type="solid")
        cell.font = Font(bold=True, size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        row_idx += 1

        headers_story = ["NO", "REGION", "TOTAL SUCCESS STORIES LOGGED"]
        for col_num, h in enumerate(headers_story, 1):
            c = ws.cell(row=row_idx, column=col_num, value=h)
            c.fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
            c.font = Font(bold=True, size=10)
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = thin_border
        row_idx += 1

        for idx, reg in enumerate(data["sortedStoryRegions"], 1):
            ws.cell(row=row_idx, column=1, value=idx).alignment = Alignment(horizontal="center")
            ws.cell(row=row_idx, column=1).border = thin_border
            ws.cell(row=row_idx, column=2, value=reg).border = thin_border
            ws.cell(row=row_idx, column=3, value=data["storyRegional"][reg]).alignment = Alignment(horizontal="center")
            ws.cell(row=row_idx, column=3).border = thin_border
            row_idx += 1

        ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=2)
        st_lbl = ws.cell(row=row_idx, column=1, value="GRAND TOTAL MILESTONES")
        st_lbl.font = Font(bold=True)
        st_lbl.alignment = Alignment(horizontal="right")
        for c_i in range(1, 3): ws.cell(row=row_idx, column=c_i).border = thin_border

        st_v = ws.cell(row=row_idx, column=3, value=data["grandStories"])
        st_v.font = Font(bold=True)
        st_v.alignment = Alignment(horizontal="center")
        st_v.border = thin_border
        for c_i in range(1, 4): ws.cell(row=row_idx, column=c_i).fill = PatternFill(start_color="FFD966", end_color="FFD966", fill_type="solid")

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

        excel_buffer = io.BytesIO()
        wb.save(excel_buffer)
        excel_buffer.seek(0)

        password_str = str(current_user.fnum).strip()
        zip_buffer = io.BytesIO()

        with pyzipper.AESZipFile(zip_buffer, 'w', compression=pyzipper.ZIP_DEFLATED, encryption=pyzipper.WZ_AES) as zf:
            zf.setpassword(password_str.encode('utf-8'))
            dl_time = "ALL_TIME" if is_all_time else f"{start_dt}_{end_dt}"
            zf.writestr(f"KMP_Consolidated_Ledger_{dl_time}.xlsx", excel_buffer.getvalue())

        zip_buffer.seek(0)

        record_neon_activity(
            logs_db=logs_db,
            fnum=current_user.fnum,
            action_type="EXPORT",
            module="CONSOLIDATED_LEDGER",
            target_id="MASTER_SUMMARY_ZIP",
            changes_summary=f"{current_user.fnum} {current_user.rank} exported encrypted consolidated ledger ({time_range})."
        )

        filename_time = "ALL_TIME" if is_all_time else f"{start_dt}_{end_dt}"
        return StreamingResponse(
            zip_buffer,
            media_type="application/zip",
            headers={"Content-Disposition": f"attachment; filename=KMP_Consolidated_Ledger_{filename_time}.zip"}
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Consolidated Ledger Export Failed: {str(e)}")