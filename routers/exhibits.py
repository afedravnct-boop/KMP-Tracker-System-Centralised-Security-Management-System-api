# routers/exhibits.py
import json
import re
import io
from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, PatternFill, Font, Border, Side
import pyzipper

from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import or_, func, text
from typing import Optional, List

from app import models
from app.database import get_db, get_logs_db
from auth import get_current_user  
from routers.activity_logger import record_neon_activity

router = APIRouter(
    prefix="/api/v1/exhibits",
    tags=["Impounded Fleet & Exhibits Registry"]
)

REGIONAL_HIERARCHY = {
    "KMP NORTH": ["KMP NORTH HEADQUARTERS", "KMP NORTH", "KAWEMPE", "KAKIRI", "KASANGATI", "MATUGGA", "NANSANA", "OLD KAMPALA", "WAKISO", "WANDEGEYA"],
    "KMP EAST": ["KMP EAST HEADQUARTERS", "KMP EAST", "JINJA ROAD", "KIRA", "KIRA DIV", "KIRA ROAD", "MUKONO", "NAGGALAMA", "SEETA"],
    "KMP SOUTH": ["KMP SOUTH HEADQUARTERS", "KMP SOUTH", "NATEETE", "CPS KAMPALA", "PARLIAMENT", "ENTEBBE", "KABALAGALA", "KAJJANSI", "KASENYI", "KATWE", "KYENGERA", "NSANGI"],
    "KMP HEADQUARTERS": ["KMP HEADQUARTERS", "KMP CID", "KMP TRAFFIC", "KMP ICT", "KMP FLYING SQUAD", "KMP CRIME INTELLIGENCE"],
    "POLICE HEADQUARTERS": ["NAGURU", "OPERATIONS", "CRIME INTELLIGENCE", "CID", "LOGISTICS & ENGINEERING", "ICT", "CT", "FIRE & RESCUE"]
}

def strip_html_tags(text_str: str) -> str:
    if not text_str:
        return ""
    return re.sub('<.*?>', '', str(text_str))

def get_model_safe(*names):
    """Safe helper to resolve database models across varying naming conventions and tablenames."""
    for name in names:
        if hasattr(models, name):
            return getattr(models, name)
    target_slug = names[-1].lower()
    for attr_name, cls in models.__dict__.items():
        if isinstance(cls, type) and hasattr(cls, '__tablename__'):
            if cls.__tablename__.lower() in [n.lower() for n in names] or cls.__tablename__.lower() == target_slug:
                return cls
    return None

# 🟢 Local serializer to prevent circular imports with api_backend
def serialize_model_row(row):
    if not row:
        return {}
    if isinstance(row, dict):
        d = row.copy()
    elif hasattr(row, '__dict__'):
        d = row.__dict__.copy()
        d.pop('_sa_instance_state', None)
    else:
        return {}
    for k, v in d.items():
        if isinstance(v, datetime):
            d[k] = v.strftime("%Y-%m-%d %H:%M:%S")
        elif hasattr(v, 'isoformat'):
            d[k] = v.isoformat()
        elif hasattr(v, '__float__') and not isinstance(v, (int, float, str, bool)):
            d[k] = float(v)
    return d

def apply_opsec_scope(current_user, query, ModelClass):
    if not ModelClass or not current_user:
        return query.filter(text("1=0"))
    
    user_role = str(current_user.role).strip().upper() if current_user.role else ""
    user_pos = str(current_user.position).strip().upper() if current_user.position else ""
    user_reg = str(current_user.region).strip().upper() if current_user.region else ""
    user_stn = str(current_user.station).strip().upper() if current_user.station else ""

    perms = current_user.permissions or {}
    if isinstance(perms, str):
        try: perms = json.loads(perms)
        except Exception: perms = {}

    is_absolute_global = (
        user_role in ["SUPER_ADMIN", "ADMIN", "ASSISTANT_SUPER_ADMIN", "RPC", "DEPUTY COMMANDER"] or
        "KMP COMMANDER" in user_pos or
        "DEPUTY KMP COMMANDER" in user_pos or
        "KMP ADMIN" in user_pos or
        perms.get("view_global_roster") is True or
        perms.get("global_observer") is True
    )

    is_kmp_sys_mgr = (
        user_role == "SYSTEM_MANAGER" and
        user_reg in ["KMP HEADQUARTERS", "POLICE HEADQUARTERS"] and
        "KMP" in user_pos
    )

    is_regional_command = (
        user_role in ["RPC", "DEPUTY_RPC", "SYSTEM_MANAGER", "ASSISTANT_SYSTEM_MANAGER", "REGIONAL_ADMIN", "ASSISTANT_REGIONAL_ADMIN", "DIVISION_ADMIN", "STATION_ADMIN"] or
        "HR" in user_pos
    )

    if is_absolute_global or is_kmp_sys_mgr:
        return query
        
    elif is_regional_command or user_reg in REGIONAL_HIERARCHY:
        conds = []
        if hasattr(ModelClass, 'region'):
            conds.append(func.upper(ModelClass.region) == user_reg)
        
        if hasattr(ModelClass, 'station') and user_reg in REGIONAL_HIERARCHY:
            expanded_stns = set()
            for s in REGIONAL_HIERARCHY[user_reg]:
                expanded_stns.add(s)
                expanded_stns.add(s.replace(' HEADQUARTERS', '').replace(' HQ', ''))
                expanded_stns.add(s + ' HEADQUARTERS')
                expanded_stns.add(s + ' HQ')
            
            conds.append(func.upper(ModelClass.station).in_(list(expanded_stns)))
            
        if conds:
            return query.filter(or_(*conds))
        return query.filter(text("1=0"))
        
    elif hasattr(ModelClass, 'station'):
        if perms.get("acc_documents") is True or perms.get("global_observer") is True:
            return query
        return query.filter(func.upper(ModelClass.station) == user_stn)
        
    return query.filter(text("1=0"))

# ====================================================================
# 1. RETRIEVE EXHIBITS
# ====================================================================
@router.get("")
def get_exhibits(
    region: Optional[str] = Query(default=None), 
    station: Optional[str] = Query(default=None), 
    case_type: Optional[str] = Query(default=None),
    status_filter: Optional[str] = Query(default=None),
    unit: Optional[str] = Query(default=None),
    start_date: Optional[str] = Query(default=None),
    end_date: Optional[str] = Query(default=None),
    search: Optional[str] = Query(default=None), 
    limit: int = Query(default=300, le=1000), 
    db: Session = Depends(get_db), 
    logs_db: Session = Depends(get_logs_db),
    current_user = Depends(get_current_user)
):
    try:
        Model = get_model_safe('impounded_exhibits', 'Impounded_Exhibits', 'Exhibits', 'ImpoundedExhibits')
        if not Model: return []
        
        query = db.query(Model)
        query = apply_opsec_scope(current_user, query, Model)
            
        if region and region.upper() not in ['ALL REGIONS', 'ALL']:
            query = query.filter(func.upper(Model.region) == region.upper())
            
        if station and station.upper() not in ['ALL STATIONS', 'ALL']:
            clean_stn = station.upper().replace(' HEADQUARTERS', '').replace(' HQ', '')
            query = query.filter(
                or_(
                    func.upper(Model.station) == station.upper(),
                    func.upper(Model.station) == clean_stn,
                    func.upper(Model.station) == f"{clean_stn} HEADQUARTERS"
                )
            )

        if case_type and case_type.upper() not in ['ALL', 'ANY CASE TYPE']:
            base_case_type = case_type.replace(':', '').upper()
            query = query.filter(func.upper(Model.case_no).startswith(base_case_type))

        if status_filter and status_filter.upper() not in ['ALL', 'ANY STATUS']:
            query = query.filter(func.upper(Model.status) == status_filter.upper())

        if unit and unit.upper() not in ['ALL', 'ANY UNIT RESPONSIBLE']:
            query = query.filter(func.upper(Model.unit_responsible) == unit.upper())

        if start_date and hasattr(Model, 'date_impounded'):
            query = query.filter(Model.date_impounded >= start_date)
            
        if end_date and hasattr(Model, 'date_impounded'):
            query = query.filter(Model.date_impounded <= end_date)
            
        if search:
            term = f"%{search.strip().upper()}%"
            if hasattr(Model, 'category'):
                query = query.filter(or_(
                    Model.reg_no.ilike(term),
                    Model.category.ilike(term),
                    Model.type_make.ilike(term),
                    Model.case_no.ilike(term),
                    Model.reason.ilike(term),
                    Model.status.ilike(term)
                ))
            else:
                query = query.filter(or_(
                    Model.reg_no.ilike(term),
                    Model.type_make.ilike(term),
                    Model.case_no.ilike(term),
                    Model.reason.ilike(term),
                    Model.status.ilike(term)
                ))
            
        pk_col = getattr(Model, 'id', getattr(Model, 'sn', getattr(Model, 's_no', None)))
        if pk_col is not None and hasattr(pk_col, 'desc'):
            records = query.order_by(pk_col.desc()).limit(limit).all()
        else:
            records = query.limit(limit).all()

        serialized_records = []
        for r in records:
            item = {}
            if hasattr(Model, '__table__'):
                for col in Model.__table__.columns:
                    col_name = col.name
                    val = getattr(r, col_name, None)
                    if isinstance(val, datetime):
                        item[col_name] = val.strftime("%Y-%m-%d %H:%M:%S")
                    elif hasattr(val, 'isoformat'):
                        item[col_name] = val.isoformat()
                    else:
                        item[col_name] = val
            else:
                item = serialize_model_row(r)
            serialized_records.append(item)

        if search:
            summary_text = f"{current_user.fnum} {current_user.rank} {current_user.name} searched exhibits ledger for query: \"{search}\" (Returned {len(serialized_records)} matches)."
        else:
            summary_text = f"{current_user.fnum} {current_user.rank} {current_user.name} accessed Exhibits & Impounded Fleet Ledger (Fetched {len(serialized_records)} items)."

        record_neon_activity(
            logs_db=logs_db,
            fnum=current_user.fnum,
            action_type="VIEW",
            module="EXHIBITS_REGISTRY",
            target_id=search if search else "ALL_RECORDS",
            changes_summary=summary_text
        )

        return serialized_records
    except Exception as e:
        print(f"Error fetching exhibits: {e}")
        return []

# ====================================================================
# 1B. SECURE EXPORT EXHIBITS TO ENCRYPTED ZIP (EXCEL)
# ====================================================================
@router.get("/export")
def export_exhibits_secure(
    region: Optional[str] = Query(default=None),
    station: Optional[str] = Query(default=None),
    case_type: Optional[str] = Query(default=None),
    status_filter: Optional[str] = Query(default=None),
    unit: Optional[str] = Query(default=None),
    start_date: Optional[str] = Query(default=None),
    end_date: Optional[str] = Query(default=None),
    search: Optional[str] = Query(default=None),
    db: Session = Depends(get_db),
    logs_db: Session = Depends(get_logs_db),
    current_user: models.Users = Depends(get_current_user)
):
    try:
        Model = get_model_safe('impounded_exhibits', 'Impounded_Exhibits', 'Exhibits', 'ImpoundedExhibits')
        if not Model:
            raise HTTPException(status_code=500, detail="Exhibits model not found")

        query = db.query(Model)
        query = apply_opsec_scope(current_user, query, Model)

        if region and region.upper() not in ['ALL REGIONS', 'ALL']:
            query = query.filter(func.upper(Model.region) == region.upper())
            
        if station and station.upper() not in ['ALL STATIONS', 'ALL']:
            clean_stn = station.upper().replace(' HEADQUARTERS', '').replace(' HQ', '')
            query = query.filter(
                or_(
                    func.upper(Model.station) == station.upper(),
                    func.upper(Model.station) == clean_stn,
                    func.upper(Model.station) == f"{clean_stn} HEADQUARTERS"
                )
            )

        if case_type and case_type.upper() not in ['ALL', 'ANY CASE TYPE']:
            base_case_type = case_type.replace(':', '').upper()
            query = query.filter(func.upper(Model.case_no).startswith(base_case_type))

        if status_filter and status_filter.upper() not in ['ALL', 'ANY STATUS']:
            query = query.filter(func.upper(Model.status) == status_filter.upper())

        if unit and unit.upper() not in ['ALL', 'ANY UNIT RESPONSIBLE']:
            query = query.filter(func.upper(Model.unit_responsible) == unit.upper())

        if start_date and hasattr(Model, 'date_impounded'):
            query = query.filter(Model.date_impounded >= start_date)
            
        if end_date and hasattr(Model, 'date_impounded'):
            query = query.filter(Model.date_impounded <= end_date)

        if search:
            term = f"%{search.strip().upper()}%"
            if hasattr(Model, 'category'):
                query = query.filter(or_(
                    Model.reg_no.ilike(term),
                    Model.category.ilike(term),
                    Model.type_make.ilike(term),
                    Model.case_no.ilike(term),
                    Model.reason.ilike(term),
                    Model.status.ilike(term)
                ))
            else:
                query = query.filter(or_(
                    Model.reg_no.ilike(term),
                    Model.type_make.ilike(term),
                    Model.case_no.ilike(term),
                    Model.reason.ilike(term),
                    Model.status.ilike(term)
                ))

        pk_col = getattr(Model, 'id', getattr(Model, 'sn', getattr(Model, 's_no', None)))
        if pk_col is not None and hasattr(pk_col, 'desc'):
            records = query.order_by(pk_col.desc()).all()
        else:
            records = query.all()
            
        total_records = len(records)

        # --- 🟢 1. Initialize Workbook & Inject Metadata ---
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Impounded Exhibits"
        ws.views.sheetView[0].showGridLines = True

        wb.properties.creator = f"{current_user.rank} {current_user.name} (FNUM: {current_user.fnum})"
        wb.properties.title = "KMP Secure Exhibits & Impounded Fleet Registry"
        wb.properties.subject = f"Jurisdiction: {region or 'ALL'} / {station or 'ALL'}"
        wb.properties.company = "Uganda Police Force - Kampala Metropolitan Police"

        thin_border = Border(
            left=Side(style='thin', color='A6A6A6'), right=Side(style='thin', color='A6A6A6'),
            top=Side(style='thin', color='A6A6A6'), bottom=Side(style='thin', color='A6A6A6')
        )

        # --- 🟢 2. Build Headers ---
        headers = ["ID", "IDENTIFIER (REG/SERIAL)", "CATEGORY", "DESCRIPTION", "COLOUR/FEATURES", 
                   "DATE IMPOUNDED", "CASE NO", "REASON", "STATUS", "UNIT RESPONSIBLE", 
                   "ASSORTED ITEMS", "IMPOUNDED BY", "ENTERED BY", "COMMENT", "PHOTO URL"]
        
        for col_num, h in enumerate(headers, 1):
            c = ws.cell(row=1, column=col_num, value=h)
            c.fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
            c.font = Font(bold=True, size=10)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = thin_border

        # --- 🟢 3. Populate Rows ---
        for row_idx, r in enumerate(records, 2):
            cat = getattr(r, 'category', 'MOTOR VEHICLE')
            desc = getattr(r, 'type_make', '')
            
            if not getattr(r, 'category', None) and ' - ' in desc:
                parts = desc.split(' - ')
                cat = parts[0].strip()
                desc = ' - '.join(parts[1:]).strip()

            impounded_by_str = ""
            if getattr(r, 'impounded_by_rank', None):
                impounded_by_str = f"{getattr(r, 'impounded_by_rank')} {getattr(r, 'impounded_by_name', '')} ({getattr(r, 'impounded_by_fnum', '')})"

            row_data = [
                getattr(r, 'id', ''),
                getattr(r, 'reg_no', ''),
                cat,
                desc,
                getattr(r, 'colour', ''),
                str(getattr(r, 'date_impounded', '')),
                getattr(r, 'case_no', ''),
                getattr(r, 'reason', ''),
                getattr(r, 'status', ''),
                getattr(r, 'unit_responsible', ''),
                getattr(r, 'assorted_items', ''),
                impounded_by_str.strip(),
                getattr(r, 'entered_by', ''),
                getattr(r, 'comment', ''),
                getattr(r, 'photo_url', '') # 🟢 photo_url explicitly included in exports
            ]

            for col_num, val in enumerate(row_data, 1):
                c = ws.cell(row=row_idx, column=col_num, value=val)
                c.border = thin_border
                c.alignment = Alignment(vertical="center", horizontal="left" if col_num in [4, 8, 11, 14, 15] else "center")

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

        # --- 🟢 4. Write to Buffer & Encrypt ---
        excel_buffer = io.BytesIO()
        wb.save(excel_buffer)
        excel_buffer.seek(0)

        password_str = str(current_user.fnum).strip()
        zip_buffer = io.BytesIO()

        file_label = status_filter.replace(" ", "_").upper() if status_filter and status_filter not in ["ALL", "ANY STATUS"] else "ALL_EXHIBITS"
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        excel_filename = f"KMP_EXHIBITS_{file_label}_{timestamp}.xlsx"

        with pyzipper.AESZipFile(zip_buffer, 'w', compression=pyzipper.ZIP_DEFLATED, encryption=pyzipper.WZ_AES) as zf:
            zf.setpassword(password_str.encode('utf-8'))
            zf.writestr(excel_filename, excel_buffer.getvalue())

        zip_buffer.seek(0)
        
        # --- 🟢 5. Log & Stream Response ---
        record_neon_activity(
            logs_db=logs_db,
            fnum=current_user.fnum,
            action_type="EXPORT",
            module="EXHIBITS_REGISTRY",
            target_id="SECURE_ZIP_EXPORT",
            changes_summary=f"{current_user.fnum} {current_user.rank} exported highly secure encrypted AES-256 ZIP containing {total_records} exhibit records (Filter: {file_label})."
        )

        return StreamingResponse(
            zip_buffer,
            media_type="application/zip",
            headers={"Content-Disposition": f"attachment; filename=KMP_EXHIBITS_ENCRYPTED_{timestamp}.zip"}
        )
    except Exception as e:
        print(f"Secure Export error: {e}")
        raise HTTPException(status_code=500, detail="Failed to generate secure ZIP export.")

# ====================================================================
# 2. IMPOUND NEW EXHIBIT
# ====================================================================
@router.post("")
def create_exhibit(
    data: dict, 
    db: Session = Depends(get_db), 
    logs_db: Session = Depends(get_logs_db),
    current_user = Depends(get_current_user)
):
    try:
        Model = get_model_safe('impounded_exhibits', 'Impounded_Exhibits', 'Exhibits', 'ImpoundedExhibits')
        if not Model: raise HTTPException(status_code=500, detail="Exhibits model not initialized.")
        
        model_kwargs = {
            "reg_no": data.get("reg_no", "NIL"),
            "type_make": data.get("type_make", "UNKNOWN"),
            "colour": data.get("colour"),
            "case_no": data.get("case_no"),
            "reason": data.get("reason"),
            "status": data.get("status", "COURT"),
            "unit_responsible": data.get("unit_responsible", "CID"),
            "assorted_items": data.get("assorted_items", "NIL"),
            "comment": data.get("comment", "NIL"),
            "photo_url": data.get("photo_url"), # 🟢 Added photo_url parsing from payload
            "region": data.get("region") or getattr(current_user, 'region', 'KMP HEADQUARTERS'),
            "station": data.get("station") or getattr(current_user, 'station', 'KMP HEADQUARTERS'),
            "date_impounded": data.get("date_impounded"),
            "impounded_by_fnum": data.get("impounded_by_fnum"),
            "impounded_by_rank": data.get("impounded_by_rank"),
            "impounded_by_name": data.get("impounded_by_name"),
            "date_cleared": data.get("date_cleared"),
            "entered_by": data.get("entered_by")
        }

        cat = data.get("category", "MOTOR VEHICLE")
        if hasattr(Model, 'category'):
            model_kwargs['category'] = cat
        else:
            model_kwargs["type_make"] = f"{cat} - {model_kwargs['type_make']}"

        new_item = Model(**model_kwargs)
        db.add(new_item)
        db.commit()
        db.refresh(new_item)

        record_neon_activity(
            logs_db=logs_db,
            fnum=current_user.fnum,
            action_type="REGISTER",
            module="EXHIBITS_REGISTRY",
            target_id=str(new_item.id),
            changes_summary=f"{current_user.fnum} {current_user.rank} {current_user.name} registered new exhibit: Case No [{new_item.case_no}], Type: [{new_item.type_make}]."
        )

        return serialize_model_row(new_item)
    except Exception as e:
        db.rollback()
        logs_db.rollback()
        raise HTTPException(status_code=500, detail=str(e))

# ====================================================================
# 3. UPDATE EXHIBIT RECORD
# ====================================================================
@router.put("/{item_id}")
def update_exhibit(
    item_id: int, 
    data: dict, 
    db: Session = Depends(get_db), 
    logs_db: Session = Depends(get_logs_db),
    current_user = Depends(get_current_user)
):
    try:
        Model = get_model_safe('impounded_exhibits', 'Impounded_Exhibits', 'Exhibits', 'ImpoundedExhibits')
        item = db.query(Model).filter(Model.id == item_id).first()
        if not item: raise HTTPException(status_code=404, detail="Exhibit record not found.")
        
        cat = data.get("category", "MOTOR VEHICLE")
        if hasattr(Model, 'category'):
            setattr(item, 'category', cat)
        else:
            data['type_make'] = f"{cat} - {data.get('type_make', 'UNKNOWN')}"
        
        for k, v in data.items():
            if hasattr(item, k) and k not in ['id', 'category']:
                setattr(item, k, v)
                
        db.commit()
        db.refresh(item)

        record_neon_activity(
            logs_db=logs_db,
            fnum=current_user.fnum,
            action_type="UPDATE",
            module="EXHIBITS_REGISTRY",
            target_id=str(item_id),
            changes_summary=f"{current_user.fnum} {current_user.rank} {current_user.name} modified exhibit record for Case No [{getattr(item, 'case_no', 'N/A')}]."
        )

        return serialize_model_row(item)
    except Exception as e:
        db.rollback()
        logs_db.rollback()
        raise HTTPException(status_code=500, detail=str(e))