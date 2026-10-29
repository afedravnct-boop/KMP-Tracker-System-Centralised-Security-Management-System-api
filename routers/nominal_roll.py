import io
import os
import math
import re
import json
import base64
from datetime import datetime, date
from typing import Optional, List, Union
from urllib.parse import unquote

import numpy as np
import pandas as pd
import pytz
import pyzipper
import openpyxl
from openpyxl.styles import Alignment, PatternFill, Font, Border, Side

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form, status, Request, Query
from fastapi.responses import JSONResponse, StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import func, or_, and_
from sqlalchemy.exc import IntegrityError

from app import models, schemas
from app.database import get_db, get_logs_db
from auth import get_current_user
from routers.activity_logger import record_neon_activity

router = APIRouter(prefix="/api/v1", tags=["Nominal Roll & HR"])

# ====================================================================
# HIERARCHICAL RANK & COMMAND RESPONSIBILITY WEIGHT ENGINE
# ====================================================================

RANK_SENIORITY = {
    "IGP": 1, "DIGP": 2, "AIGP": 3, "SCP": 4, "CP": 5, "ACP": 6,
    "SSP": 7, "SP": 8, "MAJOR": 8.5, "SASP": 9, "ASP": 10, "IP": 11, "AIP": 12,
    "HCM": 13, "HC": 14, "S/SGT": 15, "SSGT": 15, "SGT": 16,
    "CPL": 17, "L/CPL": 18, "LCPL": 18, "PC": 19, "PPC": 20, "SPC": 21, "CIVILIAN": 50
}

REGION_SORT_PRIORITY = {
    "KMP HEADQUARTERS": 1,
    "KMP NORTH HEADQUARTERS": 2,
    "KMP NORTH": 3,
    "KMP EAST HEADQUARTERS": 4,
    "KMP EAST": 5,
    "KMP SOUTH HEADQUARTERS": 6,
    "KMP SOUTH": 7,
    "POLICE HEADQUARTERS": 8
}

REGIONAL_HIERARCHY = {
    "KMP NORTH": ["KMP NORTH HEADQUARTERS", "KAWEMPE", "KAKIRI", "KASANGATI", "MATUGGA", "NANSANA", "OLD KAMPALA", "WAKISO", "WANDEGEYA"],
    "KMP EAST": ["KMP EAST HEADQUARTERS", "JINJA ROAD", "KIRA", "KIRA DIV", "KIRA ROAD", "MUKONO", "NAGGALAMA", "SEETA"],
    "KMP SOUTH": ["KMP SOUTH HEADQUARTERS", "NATEETE", "CPS KAMPALA", "PARLIAMENT", "ENTEBBE", "KABALAGALA", "KAJJANSI", "KASENYI", "KATWE", "KYENGERA", "NSANGI"],
    "KMP HEADQUARTERS": ["KMP HEADQUARTERS", "FLYING SQUAD", "CRIME INTELLIGENCE"],
    "POLICE HEADQUARTERS": ["NAGURU"]
}

def clean_str(val) -> str:
    if not val: return ''
    return str(val).strip().upper()

def normalize_region_name(region_str) -> str:
    if not region_str: return "KMP HEADQUARTERS"
    r = clean_str(region_str)
    if r in ["KMPN", "KMP/N", "KMP NORTH"]: return "KMP NORTH"
    if r in ["KMPS", "KMP/S", "KMP SOUTH"]: return "KMP SOUTH"
    if r in ["KMPE", "KMP/E", "KMP EAST"]: return "KMP EAST"
    if r in ["KMP HQ", "KMP HEADQUARTERS"]: return "KMP HEADQUARTERS"
    if r in ["POLICE HQ", "POLICE HEADQUARTERS", "NAGURU"]: return "POLICE HEADQUARTERS"
    return r

def normalize_station_name(station_str) -> str:
    if not station_str: return "UNKNOWN"
    val = clean_str(station_str)
    if val in ["OLD KLA", "OK", "OLD K", "O/K", "OLD KAMPALA DIVISION"]:
        return "OLD KAMPALA"
    val = val.replace(" POLICE STATION", "").replace(" POL STN", "").replace(" P/STN", "").replace(" DIVISION", "")
    if val.endswith(" P") and len(val) > 2:  
        val = val[:-2].strip()
    return val

def normalize_personnel_status(status_str) -> str:
    if not status_str: return "ACTIVE (DEPLOYABLE)"
    val = clean_str(status_str)
    if any(k in val for k in ["SICK", "MATERNITY", "M/LEAVE", "M.LEAVE", "LEAVE", "MEDICAL", "LIGHT DUTY", "UNFIT", "EXCUSED", "NON-DEPLOYABLE"]):
        return "NON-DEPLOYABLE (HEALTH/LEAVE)"
    return val

def calculate_age_bracket(dob_val) -> str:
    if not dob_val:
        return "UNKNOWN / UNRECORDED"
    try:
        birth_date = dob_val if isinstance(dob_val, date) else parse_safe_date(dob_val)
        if not birth_date:
            return "UNKNOWN / UNRECORDED"
        current_year = datetime.now().year
        age = current_year - birth_date.year
        if 18 <= age <= 29: return "18-29 YEARS"
        if 30 <= age <= 39: return "30-39 YEARS"
        if 40 <= age <= 49: return "40-49 YEARS"
        if age >= 50: return "50+ YEARS"
        return "UNKNOWN / UNRECORDED"
    except Exception:
        return "UNKNOWN / UNRECORDED"

def normalize_section_unit(pos_str) -> str:
    if not pos_str: return "GENERAL DUTY (GD)"
    val = clean_str(pos_str)
    if any(k in val for k in ["CID", "HOMICIDE", "C.I.D"]): return "CID"
    if any(k in val for k in ["CRIME INTEL", "CRIME INTELLIG", "INTELLIGENCE", " C.I.", " CI "]) or val.startswith("CI "): return "CRIME INTELLIGENCE"
    if any(k in val for k in ["COUNTER", "TERRORISM", " C.T.", "CT "]) or val == "CT": return "COUNTER-TERRORISM"
    if any(k in val for k in ["TRAFFIC", "TRAFF", "TRAF", "T&RSA", "TRAS", "TRSA", "TRAFFIC&RS", "TRAFFIC & RS"]): return "TRAFFIC & RS"
    if any(k in val for k in ["CFPU", "CPC", "FAMILY", "CLO", "GBV", "GENDER"]): return "CFPU / GENDER"
    if any(k in val for k in ["LOGISTIC", "LOGS", "L&E", "DRV", "DRIVING", "FLEET", "CONSTRUCTION", "LOG & ENG"]): return "LOGISTICS & ENGINEERING"
    if any(k in val for k in ["SOCO", "FORENSIC"]): return "FORENSIC / SOCO"
    if any(k in val for k in ["PRO", "MEDIA", "PUBLIC RELATIONS"]): return "PRO / MEDIA"
    if any(k in val for k in ["ICT", "CCTV", "COMM", "SIGNALLER", "COMMUNICATOR"]): return "ICT & COMMUNICATIONS"
    if any(k in val for k in ["FIRE", "F&RS", "RESCUE"]): return "FIRE & RESCUE"
    if any(k in val for k in ["GD", "G.D", "G/D", "OPERATIONS", "OPS"]): return "OPERATIONS / GD"
    return val

def get_station_priority_weight(station, region) -> int:
    stn = clean_str(station)
    reg = clean_str(region)
    if 'KMP HEADQUARTERS' in stn or 'KMP HEADQUARTERS' in reg or stn == 'HQ':
        return 0
    if 'HEADQUARTERS' in stn or 'RPC' in stn:
        return 1
    return 2

def get_command_weight(officer) -> int:
    try:
        pos = clean_str(getattr(officer, 'position', ''))
        if 'KMP COMMANDER' in pos or (any(k in pos for k in ['COMD', 'COMDR', 'COM', 'COMMANDER', 'CDR']) and 'KMP' in pos and not ('DEP' in pos or 'DEPUTY' in pos)):
            return 1
        if 'DEPUTY KMP' in pos or 'DEP KMP' in pos or 'DEP. KMP' in pos or (('DEP' in pos or 'DEPUTY' in pos) and 'KMP' in pos):
            return 2
        if 'ADMIN KMP' in pos or 'ADMIN. KMP' in pos or 'KMP ADMIN OFFICER' in pos:
            return 3
        if 'RPC' in pos and not ('DEPUTY' in pos or 'DEP' in pos):
            return 1
        if 'DEPUTY RPC' in pos or 'DEP RPC' in pos or 'D/RPC' in pos:
            return 2
        if 'DPC' in pos or 'DIVISION COMMANDER' in pos or 'DIV COMDR' in pos:
            return 4
        if 'DEPUTY DPC' in pos or 'DEP DPC' in pos:
            return 5
        if (any(k in pos for k in ['COM', 'COMD', 'COMDR', 'COMMANDER', 'CDR'])) and not ('DEP' in pos or 'DEPUTY' in pos):
            return 5
        if 'DEP' in pos or 'DEPUTY' in pos:
            return 6
        if 'OC STATION' in pos or 'OC DIV' in pos or ('OC' in pos and not ('CID' in pos or 'CI' in pos)):
            return 7
        if 'OC CID' in pos or 'HEAD CID' in pos:
            return 8
        if 'OC CI' in pos or 'CRIME INTELLIGENCE' in pos:
            return 9
        if 'OC POST' in pos or 'O/C POST' in pos or 'IC POST' in pos:
            return 10
        if '2I/C' in pos or 'DEPUTY OC' in pos or 'I/C' in pos:
            return 11
    except Exception:
        pass
    return 50

def get_rank_weight(rank_str: str) -> int:
    if not rank_str: return 99
    r = clean_str(rank_str)
    if 'DRV' in r or 'DRIVER' in r:
        if 'SGT' in r or 'SERGEANT' in r: r = 'SGT'
        elif 'CPL' in r or 'CORPORAL' in r: r = 'CPL'
        else: r = 'PC'
    if r == 'DC' or r.startswith('D/C'):
        r = 'PC'
    elif r.startswith('D/') or r.startswith('D-') or r.startswith('D '):
        r = r.replace('D/', '').replace('D-', '').replace('D ', '').strip()
        if r == 'C': r = 'PC'
    if not r: r = 'PC'
    return RANK_SENIORITY.get(r, 40)

def hierarchical_sort_key(officer):
    try:
        stn = getattr(officer, 'station', '')
        reg = getattr(officer, 'region', '')
        r_name = getOfficialRegionForStation(stn, reg)
        reg_pri = REGION_SORT_PRIORITY.get(r_name, 50)
        rank_w = get_rank_weight(getattr(officer, 'rank', ''))
        cmd = get_command_weight(officer)
        fnum = clean_str(getattr(officer, 'f_num', getattr(officer, 'fnum', '')))
        return (reg_pri, rank_w, cmd, fnum)
    except Exception:
        return (50, 99, 50, '')

def require_export_privilege(current_user: models.Users = Depends(get_current_user)):
    user_role = str(current_user.role).strip().upper() if current_user.role else ""
    perms = current_user.permissions or {}
    if isinstance(perms, str):
        try: perms = json.loads(perms)
        except Exception: perms = {}
    if (
        user_role not in ["ADMIN", "SUPER_ADMIN", "RPC", "SYSTEM_MANAGER", "ASSISTANT_SYSTEM_MANAGER", "REGIONAL_ADMIN"] and 
        not perms.get("export_data", False) and 
        not perms.get("global_observer", False) and
        not perms.get("view_global_roster", False)
    ):
        raise HTTPException(status_code=403, detail="Clearance Denied: Data Export Privileges Required.")
    return current_user

def aggressive_clean_text(val):
    if pd.isna(val) or val is None: return None
    s = str(val)
    if s.lower() in ['nan', 'nat', 'none', 'null', '']: return None
    s = re.sub(r"[,!?'\"]", "", s)
    s = re.sub(r'\s+', ' ', s)
    s = s.strip('. -/\\')
    if not s: return None
    return s.upper()

def clean_nin(val):
    if pd.isna(val) or val is None: return None
    s = str(val).strip().upper()
    if s.lower() in ['nan', 'nat', 'none', 'null', '', '0', 'N/A', 'NIL']: return None
    if s.endswith('.0'): s = s[:-2]
    s = re.sub(r'[^A-Z0-9]', '', s)
    return s if len(s) > 0 else None

def clean_numeric(val):
    if pd.isna(val) or val is None: return None
    s = str(val).strip()
    if s.lower() in ['nan', 'nat', 'none', 'null', '']: return None
    s = re.sub(r'[^\d.]', '', s) 
    if s.endswith('.0'): s = s[:-2]
    return s

def format_phone_number(val):
    if pd.isna(val) or val is None: return None
    s = str(val).strip()
    if s.lower() in ['nan', 'nat', 'none', 'null', '']: return None
    parts = re.split(r'[,/&|;]|\band\b', s, flags=re.IGNORECASE)
    valid_numbers = []
    for part in parts:
        cleaned = re.sub(r'[^\d]', '', part)
        if not cleaned: continue
        if cleaned.startswith('256'):
            cleaned = '0' + cleaned[3:]
        elif cleaned.startswith('7') and len(cleaned) == 9:
            cleaned = '0' + cleaned
        if len(cleaned) >= 10:
            valid_numbers.append(cleaned)
    if not valid_numbers: return None
    return ' / '.join(valid_numbers)

def normalize_sex(val):
    clean_val = aggressive_clean_text(val)
    if not clean_val: return "MALE"
    if clean_val.startswith('F'): return "FEMALE"
    if clean_val.startswith('M'): return "MALE"
    return clean_val

def normalize_education_level(educ_str):
    cleaned = aggressive_clean_text(educ_str)
    if not cleaned: return None
    if any(term in cleaned for term in ['S.1', 'S1', 'S.2', 'S2', 'S.3', 'S3', 'SENIOR 1', 'SENIOR 2', 'SENIOR 3']): return cleaned
    if any(term in cleaned for term in ['UACE', 'A-LEVEL', 'A LEVEL', 'S.6', 'S6', 'SENIOR 6']): return "UACE"
    if any(term in cleaned for term in ['UCE', 'O-LEVEL', 'O LEVEL', 'S.4', 'S4', 'SENIOR 4', 'PLE', 'P.7', 'P7']): return "UCE"
    return cleaned

def is_uniformed_rank(rank_str: str) -> bool:
    if not rank_str: return False
    r = str(rank_str).strip().upper()
    uniformed_ranks = {
        'IGP', 'DIGP', 'AIGP', 'SCP', 'CP', 'ACP', 'SSP', 'SP', 'MAJOR', 'SASP', 'ASP',
        'IP', 'AIP', 'HCM', 'HC', 'S/SGT', 'SSGT', 'SGT', 'CPL', 'L/CPL', 'LCPL',
        'PC', 'PPC', 'SPC', 'DC', 'D/C'
    }
    return r in uniformed_ranks

def is_invalid_roster_entry(r) -> bool:
    name = str(getattr(r, 'name', '') or '').strip().upper()
    rank = str(getattr(r, 'rank', '') or '').strip().upper()
    fnum = str(getattr(r, 'fnum', getattr(r, 'f_num', '')) or '').strip().upper()
    combined = f"{fnum} {rank} {name}"
    if "---" in combined or "___" in combined or "===" in combined: return True
    forbidden_terms = [
        "POLICE STATION", "CANINE UNIT", "ATTACHED", "DEPARTMENT", 
        "DIVISION HEADQUARTERS", "REGIONAL HQ", "POST", "BARRACKS", 
        "SECTION", "OC STATION", "OC CID", "CRIME INVESTIGATION"
    ]
    if any(term in combined for term in forbidden_terms): return True
    if any(char.isdigit() for char in name) and not any(r_term in rank for r_term in ['PC', 'SGT', 'CPL', 'ASP', 'IP']):
        return True
    if name in ["UNKNOWN", "", "NONE", "NAN"] and not is_uniformed_rank(rank): return True
    return False

def parse_safe_date(val) -> Optional[date]:
    if pd.isna(val) or val is None: return None
    if isinstance(val, date) and not isinstance(val, datetime):
        return val if 1900 <= val.year <= 2100 else None
    if isinstance(val, datetime):
        return val.date() if 1900 <= val.year <= 2100 else None
    if type(val).__name__ == 'Timestamp':
        try: return val.date() if 1900 <= val.year <= 2100 else None
        except Exception: return None
    val_str = str(val).strip()
    if val_str.lower() in ['nan', 'nat', 'none', 'null', '', '-', 'n/a', 'nil', '0', 'undefined']: return None
    if ' ' in val_str: val_str = val_str.split(' ')[0]
    if 'T' in val_str: val_str = val_str.split('T')[0]
    val_str = val_str.strip()
    if not val_str: return None
    try:
        if val_str.replace('.', '', 1).isdigit():
            float_val = float(val_str)
            if 1000 < float_val < 73050:
                try:
                    dt = pd.to_datetime(float_val, unit='D', origin='1899-12-30', errors='coerce')
                    if pd.notna(dt) and 1900 <= dt.year <= 2100: return dt.date()
                except Exception: pass
    except Exception: pass
    clean_str_val = re.sub(r'[\./\\]', '-', val_str)
    parts = clean_str_val.split('-')
    if len(parts) == 3:
        p0, p1, p2 = parts[0].strip(), parts[1].strip(), parts[2].strip()
        if p0.isdigit() and p1.isdigit() and p2.isdigit():
            if len(p0) == 4:
                try:
                    res_date = date(int(p0), int(p1), int(p2))
                    if 1900 <= res_date.year <= 2100: return res_date
                except Exception: pass
            elif len(p2) == 4:
                try:
                    res_date = date(int(p2), int(p1), int(p0))
                    if 1900 <= res_date.year <= 2100: return res_date
                except Exception: pass
            elif len(p0) <= 2 and len(p2) <= 2:
                for (day_val, month_val, year_val) in [(int(p0), int(p1), int(p2)), (int(p2), int(p1), int(p0))]:
                    y = year_val
                    if y < 100: y = 1900 + y if y > 40 else 2000 + y
                    try:
                        res_date = date(y, month_val, day_val)
                        if 1900 <= res_date.year <= 2100: return res_date
                    except Exception: continue
    for fmt in ('%Y-%m-%d', '%d-%m-%Y', '%d/%m/%Y', '%Y/%m/%d', '%d.%m.%Y', '%m/%d/%Y', '%m-%d-%Y'):
        try:
            d = datetime.strptime(val_str, fmt).date()
            if 1900 <= d.year <= 2100: return d
        except Exception: continue
    try:
        parsed = pd.to_datetime(val_str, dayfirst=True, errors='coerce')
        if pd.notna(parsed):
            py_dt = parsed.to_pydatetime()
            if 1900 <= py_dt.year <= 2100: return py_dt.date()
    except Exception: pass
    return None

def get_officer_signature(user):
    if not user: return "UNKNOWN COMMANDER"
    fnum = (user.fnum or "").strip()
    rank = (user.rank or "").strip()
    name = (user.name or "").strip()
    return f"{fnum} {rank} {name}".strip().upper()

def get_active_model():
    model = getattr(models, 'NominalRoll', getattr(models, 'Nominal_Roll', getattr(models, 'nominal_roll', None)))
    if not model: raise HTTPException(status_code=500, detail="Nominal Roll database model not configured.")
    return model

def get_archive_model():
    model = getattr(models, 'NominalRollArchive', getattr(models, 'Nominal_Roll_Archive', getattr(models, 'nominal_roll_archive', None)))
    if not model: raise HTTPException(status_code=500, detail="Nominal Roll Archive database model not configured.")
    return model

STATION_GEO_MAP = {
    "KMP NORTH HEADQUARTERS": {"region": "KMP NORTH", "district": "KAMPALA"},
    "KAWEMPE": {"region": "KMP NORTH", "district": "KAMPALA"},
    "WANDEGEYA": {"region": "KMP NORTH", "district": "KAMPALA"},
    "OLD KAMPALA": {"region": "KMP NORTH", "district": "KAMPALA"},
    "MATUGGA": {"region": "KMP NORTH", "district": "WAKISO"},
    "NANSANA": {"region": "KMP NORTH", "district": "WAKISO"},
    "KASANGATI": {"region": "KMP NORTH", "district": "WAKISO"},
    "KAKIRI": {"region": "KMP NORTH", "district": "WAKISO"},
    "WAKISO": {"region": "KMP NORTH", "district": "WAKISO"},
    "KMP SOUTH HEADQUARTERS": {"region": "KMP SOUTH", "district": "KAMPALA"},
    "NATEETE": {"region": "KMP SOUTH", "district": "WAKISO"},
    "CPS KAMPALA": {"region": "KMP SOUTH", "district": "KAMPALA"},
    "PARLIAMENT": {"region": "KMP SOUTH", "district": "KAMPALA"},
    "ENTEBBE": {"region": "KMP SOUTH", "district": "WAKISO"},
    "KABALAGALA": {"region": "KMP SOUTH", "district": "KAMPALA"},
    "KAJJANSI": {"region": "KMP SOUTH", "district": "KAMPALA"},
    "NSANGI": {"region": "KMP SOUTH", "district": "WAKISO"},
    "KASENYI": {"region": "KMP SOUTH", "district": "WAKISO"},
    "KYENGERA": {"region": "KMP SOUTH", "district": "WAKISO"},
    "KATWE": {"region": "KMP SOUTH", "district": "KAMPALA"},
    "KMP EAST HEADQUARTERS": {"region": "KMP EAST", "district": "KAMPALA"},
    "JINJA ROAD": {"region": "KMP EAST", "district": "KAMPALA"},
    "MUKONO": {"region": "KMP EAST", "district": "MUKONO"},
    "KIRA ROAD": {"region": "KMP EAST", "district": "KAMPALA"},
    "KIRA DIV": {"region": "KMP EAST", "district": "WAKISO"},
    "KIRA": {"region": "KMP EAST", "district": "WAKISO"},
    "NAGGALAMA": {"region": "KMP EAST", "district": "MUKONO"},
    "SEETA": {"region": "KMP EAST", "district": "MUKONO"}
}

def parseStationHierarchy(rawStation):
    cleaned = normalize_station_name(rawStation)
    if not cleaned: return {"division": "UNASSIGNED", "station": "UNASSIGNED", "post": "", "type": "POST"}
    if '/' in cleaned or '\\' in cleaned:
        parts = cleaned.replace('\\', '/').split('/')
        parent = parts[0].strip()
        child = parts[1].strip() if len(parts) > 1 else ""
        type_str = "POST"
        if "SUB-STATION" in child or "SUBSTATION" in child:
            return {"division": parent, "station": parent, "post": "", "type": "STATION"} 
        elif "STATION" in child:
            type_str = "STATION"
        return {"division": parent, "station": parent, "post": type_str == "POST" and child or "", "type": type_str}
    return {"division": cleaned, "station": cleaned, "post": "", "type": "STATION"}

def getOfficialRegionForStation(station_name: str, current_region: Optional[str] = None) -> str:
    norm_reg = normalize_region_name(current_region)
    station_norm = normalize_station_name(station_name)
    
    # 🟢 Precision Check via REGIONAL_HIERARCHY mapping (ensures regional headquarters roll up to parent region)
    for reg_name, stationsList in REGIONAL_HIERARCHY.items():
        if station_norm in stationsList:
            return reg_name

    if "KMP EAST" in station_norm or "KMP EAST" in norm_reg:
        return "KMP EAST"
    if "KMP NORTH" in station_norm or "KMP NORTH" in norm_reg:
        return "KMP NORTH"
    if "KMP SOUTH" in station_norm or "KMP SOUTH" in norm_reg:
        return "KMP SOUTH"
        
    if station_norm in STATION_GEO_MAP:
        return STATION_GEO_MAP[station_norm]["region"]
        
    return norm_reg

def auto_infer_geography(station_name: str, current_region: Optional[str] = None, current_district: Optional[str] = None):
    station_norm = normalize_station_name(station_name)
    region_norm = normalize_region_name(current_region)
    district_norm = clean_str(current_district)

    for reg_name, stationsList in REGIONAL_HIERARCHY.items():
        if station_norm in stationsList:
            region_norm = reg_name
            break

    if station_norm in STATION_GEO_MAP:
        geo = STATION_GEO_MAP[station_norm]
        if not region_norm or region_norm in ["KMP HEADQUARTERS", "UNKNOWN"]:
            region_norm = geo["region"]
        if not district_norm or district_norm in ["UNKNOWN"]:
            district_norm = geo["district"]

    if not district_norm:
        district_norm = "KAMPALA"

    return region_norm, district_norm

def get_scoped_nominal_query(db: Session, current_user: models.Users, Model):
    query = db.query(Model)
    user_role = (current_user.role or "").upper()
    user_region = normalize_region_name(current_user.region)
    user_station = normalize_station_name(current_user.station)
    
    perms = current_user.permissions or {}
    if isinstance(perms, str):
        try: perms = json.loads(perms)
        except Exception: perms = {}
    
    is_absolute_global = (
        user_role in ["SUPER_ADMIN", "ADMIN", "ASSISTANT_SUPER_ADMIN"] or
        perms.get("view_global_roster") is True or
        perms.get("global_observer") is True
    )

    is_regional_command = user_role in ["RPC", "DEPUTY_RPC", "REGIONAL_ADMIN", "ASSISTANT_REGIONAL_ADMIN"]

    if is_absolute_global:
        return query
    elif is_regional_command:
        # 🟢 Gather all stations under the regional hierarchy including regional headquarters
        allowed_stations = REGIONAL_HIERARCHY.get(user_region, [])
        return query.filter(
            or_(
                func.upper(Model.region) == user_region,
                func.upper(Model.station).in_([s.upper() for s in allowed_stations]),
                func.upper(Model.station).ilike(f"%{user_region}%")
            )
        )
    elif user_station:
        station_pattern = f"%{user_station}%"
        return query.filter(
            or_(
                func.upper(Model.station) == user_station,
                func.upper(Model.station).ilike(station_pattern)
            )
        )
    return query.filter(Model.id == -1)

@router.get("/nominal-roll")
def get_Nominal_Rolls(
    search: Optional[str] = Query(default=None),
    status_filter: Optional[str] = Query(default=None),
    category_filter: Optional[str] = Query(default=None),
    db: Session = Depends(get_db), 
    logs_db: Session = Depends(get_logs_db),
    current_user: models.Users = Depends(get_current_user)
):
    ActiveModel = get_active_model()
    ArchiveModel = get_archive_model()
    
    active_query = get_scoped_nominal_query(db, current_user, ActiveModel)
    archive_query = get_scoped_nominal_query(db, current_user, ArchiveModel)
    
    if search:
        term = f"%{search.strip().upper()}%"
        active_conds = []
        archive_conds = []
        for M, C in [(ActiveModel, active_conds), (ArchiveModel, archive_conds)]:
            if hasattr(M, 'name'): C.append(M.name.ilike(term))
            if hasattr(M, 'fnum'): C.append(M.fnum.ilike(term))
            if hasattr(M, 'f_num'): C.append(M.f_num.ilike(term))
            if hasattr(M, 'rank'): C.append(M.rank.ilike(term))
            if hasattr(M, 'station'): C.append(M.station.ilike(term))
            if hasattr(M, 'ipps'): C.append(M.ipps.ilike(term))
        if active_conds: active_query = active_query.filter(or_(*active_conds))
        if archive_conds: archive_query = archive_query.filter(or_(*archive_conds))

    active_records = sorted(active_query.all(), key=hierarchical_sort_key)
    archive_records = sorted(archive_query.all(), key=hierarchical_sort_key)
    
    clean_results = []
    sequence_counter = 1

    for r in active_records:
        if is_invalid_roster_entry(r): continue
        r_dict = r.__dict__.copy()
        r_dict.pop("_sa_instance_state", None)
        
        fnum_val = r_dict.get('f_num') or r_dict.get('fnum') or ''
        r_dict['fnum'] = fnum_val
        r_dict['f_num'] = fnum_val
        r_dict['station'] = normalize_station_name(r_dict.get('station'))
        r_dict['region'] = getOfficialRegionForStation(r_dict.get('station'), r_dict.get('region'))
        r_dict['normalized_status'] = normalize_personnel_status(r_dict.get('status'))
        r_dict['normalized_section'] = normalize_section_unit(r_dict.get('position') or r_dict.get('section'))
        r_dict['age_bracket'] = calculate_age_bracket(r_dict.get('dob'))
        r_dict['educ_level'] = normalize_education_level(r_dict.get('educlevel') or r_dict.get('educ_level'))

        if status_filter and status_filter.upper() != "ALL" and r_dict['normalized_status'].upper() != status_filter.upper():
            continue
        if category_filter and category_filter.upper() != "ALL" and r_dict['normalized_section'].upper() != category_filter.upper():
            continue

        r_dict['sn'] = sequence_counter
        r_dict['dbAuditId'] = getattr(r, 'id', sequence_counter)
        r_dict['is_archived'] = False
        r_dict['status'] = r_dict.get('status') or 'ACTIVE'
        
        clean_results.append(r_dict)
        sequence_counter += 1

    for r in archive_records:
        if is_invalid_roster_entry(r): continue
        r_dict = r.__dict__.copy()
        r_dict.pop("_sa_instance_state", None)
        
        fnum_val = r_dict.get('fnum') or r_dict.get('f_num') or ''
        r_dict['fnum'] = fnum_val
        r_dict['f_num'] = fnum_val
        r_dict['station'] = normalize_station_name(r_dict.get('station'))
        r_dict['region'] = getOfficialRegionForStation(r_dict.get('station'), r_dict.get('region'))
        r_dict['normalized_status'] = normalize_personnel_status(r_dict.get('status'))
        r_dict['normalized_section'] = normalize_section_unit(r_dict.get('position') or r_dict.get('section'))
        r_dict['age_bracket'] = calculate_age_bracket(r_dict.get('dob'))
        r_dict['educ_level'] = normalize_education_level(r_dict.get('educlevel') or r_dict.get('educ_level'))

        if status_filter and status_filter.upper() != "ALL" and r_dict['normalized_status'].upper() != status_filter.upper():
            continue
        if category_filter and category_filter.upper() != "ALL" and r_dict['normalized_section'].upper() != category_filter.upper():
            continue

        r_dict['sn'] = sequence_counter
        r_dict['dbAuditId'] = f"ARC-{getattr(r, 'id', sequence_counter)}"
        r_dict['is_archived'] = True
        r_dict['status'] = r_dict.get('status') or 'ARCHIVED'
        
        clean_results.append(r_dict)
        sequence_counter += 1

    return clean_results

@router.post("/nominal-roll/bulk-upload")
@router.post("/nominal-roll/upload")
async def bulk_upload_nominal_roll(
    file: Optional[UploadFile] = File(None),
    files: Optional[List[UploadFile]] = File(None),
    db: Session = Depends(get_db),
    logs_db: Session = Depends(get_logs_db),
    current_user: models.Users = Depends(get_current_user)
):
    ActiveModel = get_active_model()
    ArchiveModel = get_archive_model()
    
    file_list = []
    if files: file_list.extend(files)
    if file: file_list.append(file)

    if not file_list:
        raise HTTPException(status_code=400, detail="No valid file uploaded.")

    inserted_count = 0
    updated_count = 0
    skipped_archived = []
    skipped_blank = []
    officer_sig = get_officer_signature(current_user)

    try:
        all_dfs = []
        for single_file in file_list:
            contents = await single_file.read()
            filename = (single_file.filename or "").lower()

            if filename.endswith(".csv"): 
                df = pd.read_csv(io.BytesIO(contents), dtype=str)
            elif filename.endswith((".xls", ".xlsx")): 
                df = pd.read_excel(io.BytesIO(contents), dtype=str)
            else: 
                continue

            def standardize_header(h):
                h = str(h).lower().strip()
                h = h.replace("f/no", "fnum").replace("f-no", "fnum").replace("force no", "fnum")
                h = h.replace("d.o.b", "dob").replace("d.o.e", "doe").replace("d.o.p", "dopost")
                return re.sub(r'[^a-z0-9]', '', h)
                
            df.columns = [standardize_header(col) for col in df.columns]
            df['__source_file'] = single_file.filename
            all_dfs.append(df)

        if not all_dfs:
            raise HTTPException(status_code=400, detail="No valid CSV or Excel files found in upload batch.")

        combined_df = pd.concat(all_dfs, ignore_index=True)

        date_columns = ['dob', 'dateofbirth', 'doe', 'dateofenlistment', 'dopost', 'dop', 'dopro', 'dateofpromotion']
        for col in date_columns:
            if col in combined_df.columns:
                combined_df[col] = combined_df[col].apply(parse_safe_date)

        for idx, row in combined_df.iterrows():
            fnum_val = aggressive_clean_text(row.get("fnum") or row.get("forceno") or row.get("forcenumber") or row.get("fileno") or row.get("fno"))
            ipps_val = clean_numeric(row.get("ipps") or row.get("ippsno") or row.get("ippsnumber"))
            nin_val = clean_nin(row.get("nin") or row.get("nationalid") or row.get("ninno"))
            rank_val = aggressive_clean_text(row.get("rank"))
            name_val = aggressive_clean_text(row.get("name"))
            source_filename = row.get("__source_file", "Batch Upload")

            row_text_signature = f"{fnum_val or ''} {rank_val or ''} {name_val or ''}".upper()
            
            if (
                not fnum_val and not rank_val and (not name_val or name_val == "UNKNOWN")
            ) or any(term in row_text_signature for term in ["DEPARTMENT", "POL. POST", "POLICE POST", "SECTION", "DIV HEADQUARTERS", "OC STATION", "---", "___", "CANINE UNIT", "BUSEGA"]):
                continue

            if rank_val and any(stn_term in rank_val for stn_term in ["STATION", "POST", "DIV", "UNIT", "HQ", "HEADQUARTERS"]):
                continue

            if not fnum_val:
                if is_uniformed_rank(rank_val):
                    skipped_blank.append(f"[{source_filename}] Row {idx+2}: {rank_val} {name_val} (Missing F/No)")
                    continue
                elif ipps_val: fnum_val = f"CIV-IPPS-{ipps_val}"
                elif nin_val: fnum_val = f"CIV-NIN-{nin_val}"
                else: 
                    skipped_blank.append(f"[{source_filename}] Row {idx+2}: {name_val or 'Unknown'} (Missing F/No, IPPS, & NIN)")
                    continue 

            clean_fnum = fnum_val
            stn_val = normalize_station_name(row.get("station") or current_user.station or "HQ")
            reg_val, dist_val = auto_infer_geography(stn_val, aggressive_clean_text(row.get("region")), aggressive_clean_text(row.get("district")))

            dob_val = row.get("dob") if isinstance(row.get("dob"), date) else parse_safe_date(row.get("dob") or row.get("dateofbirth"))
            doe_val = row.get("doe") if isinstance(row.get("doe"), date) else parse_safe_date(row.get("doe") or row.get("dateofenlistment"))
            dopost_val = row.get("dopost") if isinstance(row.get("dopost"), date) else parse_safe_date(row.get("dopost") or row.get("dop"))
            dopro_val = row.get("dopro") if isinstance(row.get("dopro"), date) else parse_safe_date(row.get("dopro") or row.get("dateofpromotion"))

            officer_payload = {
                "rank": rank_val or "CIVILIAN",
                "name": name_val or "UNKNOWN",
                "sex": normalize_sex(row.get("sex") or row.get("gender")),
                "position": aggressive_clean_text(row.get("position") or row.get("title") or "GENERAL DUTIES"),
                "dob": dob_val,
                "doe": doe_val,
                "do_post": dopost_val,
                "do_pro": dopro_val,
                "contact": format_phone_number(row.get("contact") or row.get("phone") or row.get("phonenumber")),
                "educ_level": normalize_education_level(row.get("educ_level") or row.get("educlevel") or row.get("education")),
                "ipps": ipps_val,
                "tin": clean_numeric(row.get("tin") or row.get("tinno") or row.get("tinnumber")),
                "nin": nin_val,
                "home_dist": aggressive_clean_text(row.get("homedist") or row.get("homedistrict")),
                "tribe": aggressive_clean_text(row.get("tribe")),
                "acc_no": clean_numeric(row.get("accno") or row.get("accountno") or row.get("accountnumber")),
                "bank_branch": aggressive_clean_text(row.get("bankbranch") or row.get("bank")),
                "station": stn_val,
                "district": dist_val,
                "region": reg_val,
                "section": aggressive_clean_text(row.get("section")),
                "dir": aggressive_clean_text(row.get("dir") or row.get("directorate")),
                "status": aggressive_clean_text(row.get("status") or "ACTIVE"),
                "last_updated_by": officer_sig
            }

            for key, value in list(officer_payload.items()):
                if isinstance(value, str) and not value.strip():
                    officer_payload[key] = None
                elif isinstance(value, float) and math.isnan(value):
                    officer_payload[key] = None

            if hasattr(ActiveModel, 'f_num'): officer_payload['f_num'] = clean_fnum
            if hasattr(ActiveModel, 'fnum'): officer_payload['fnum'] = clean_fnum

            fnum_filter = []
            if hasattr(ActiveModel, 'f_num'): fnum_filter.append(func.trim(func.upper(ActiveModel.f_num)) == clean_fnum)
            if hasattr(ActiveModel, 'fnum'): fnum_filter.append(func.trim(func.upper(ActiveModel.fnum)) == clean_fnum)

            existing = db.query(ActiveModel).filter(or_(*fnum_filter)).first()

            if existing:
                for k, v in officer_payload.items():
                    if hasattr(existing, k) and v is not None:
                        setattr(existing, k, v)
                updated_count += 1
            else:
                arc_filter = []
                if hasattr(ArchiveModel, 'f_num'): arc_filter.append(func.trim(func.upper(ArchiveModel.f_num)) == clean_fnum)
                if hasattr(ArchiveModel, 'fnum'): arc_filter.append(func.trim(func.upper(ArchiveModel.fnum)) == clean_fnum)
                
                is_archived = db.query(ArchiveModel).filter(or_(*arc_filter)).first()
                
                if is_archived:
                    safe_payload_json = {}
                    for k, v in officer_payload.items():
                        if isinstance(v, (date, datetime)): safe_payload_json[k] = v.isoformat()
                        else: safe_payload_json[k] = v
                        
                    entry_obj = {
                        "display": f"{officer_payload['rank']} {officer_payload['name']} ({clean_fnum})",
                        "fnum": clean_fnum,
                        "payload": safe_payload_json
                    }
                    
                    if not any(isinstance(x, dict) and x.get('fnum') == clean_fnum for x in skipped_archived):
                        skipped_archived.append(entry_obj)
                    continue
                    
                valid_cols = [c.key for c in ActiveModel.__table__.columns]
                safe_payload = {k: v for k, v in officer_payload.items() if k in valid_cols}
                new_entry = ActiveModel(**safe_payload)
                db.add(new_entry)
                inserted_count += 1
            
            db.flush()

        db.commit()

        record_neon_activity(
            logs_db=logs_db,
            fnum=current_user.fnum,
            action_type="REGISTER",
            module="NOMINAL_ROLL",
            target_id="BULK_UPLOAD",
            changes_summary=f"{current_user.fnum} {current_user.rank} {current_user.name} executed bulk personnel roster upload ({inserted_count} added, {updated_count} updated)."
        )

        return {
            "status": "warning" if (skipped_archived or skipped_blank) else "success",
            "message": f"Batch process complete across {len(file_list)} files. {inserted_count} new personnel recorded, {updated_count} updated.",
            "skipped": skipped_archived,
            "skipped_blank": skipped_blank
        }

    except Exception as e:
        db.rollback()
        logs_db.rollback()
        raise HTTPException(status_code=500, detail=f"Bulk Nominal Roll Upload Failed: {str(e)}")

@router.post("/nominal-roll")
def create_Nominal_Roll(
    data: dict, 
    db: Session = Depends(get_db), 
    logs_db: Session = Depends(get_logs_db),
    current_user: models.Users = Depends(get_current_user)
):
    ActiveModel = get_active_model()
    ArchiveModel = get_archive_model()
    
    try:
        reintegration_reason = data.pop('reintegration_reason', None)
        previous_fnum = data.pop('previous_fnum', None)
        
        data.pop('sn', None) 
        data.pop('id', None)
        
        clean_data = {}
        for k, v in data.items():
            clean_data[k] = None if v == "" else v

        test_rank = str(clean_data.get('rank', '')).upper()
        if any(stn_term in test_rank for stn_term in ["STATION", "POST", "DIV", "UNIT", "HQ", "HEADQUARTERS"]):
            raise HTTPException(status_code=400, detail="Validation Error: Station or unit names cannot be entered into the Rank field.")

        perms = current_user.permissions or {}
        user_role = (current_user.role or "").upper()
        is_global_user = (
            user_role in ["SUPER_ADMIN", "ADMIN", "RPC", "DEPUTY COMMANDER"] or
            "HR" in (current_user.position or "").upper() or
            perms.get("view_global_roster") is True or
            perms.get("global_observer") is True
        )

        if not is_global_user:
            clean_data["region"] = normalize_region_name(current_user.region)
            clean_data["station"] = normalize_station_name(current_user.station)

        if 'contact' in clean_data and clean_data['contact']:
            clean_data['contact'] = format_phone_number(clean_data['contact'])
            
        if 'name' in clean_data and clean_data['name']:
            clean_data['name'] = aggressive_clean_text(clean_data['name'])

        if 'sex' in clean_data:
            clean_data['sex'] = normalize_sex(clean_data['sex'])
            
        if 'educ_level' in clean_data:
            clean_data['educ_level'] = normalize_education_level(clean_data['educ_level'])

        if 'station' in clean_data and clean_data['station']:
            clean_data['station'] = normalize_station_name(clean_data['station'])

        if 'region' in clean_data and clean_data['region']:
            clean_data['region'] = normalize_region_name(clean_data['region'])

        for text_field in ['position', 'home_dist', 'tribe', 'bank_branch', 'section', 'dir']:
            if text_field in clean_data and clean_data[text_field]:
                clean_data[text_field] = aggressive_clean_text(clean_data[text_field])

        for date_field in ['dob', 'doe', 'do_post', 'do_pro']:
            if date_field in clean_data and clean_data[date_field]:
                clean_data[date_field] = parse_safe_date(clean_data[date_field])

        target_fnum = clean_data.get('f_num') or clean_data.get('fnum')
        if not target_fnum:
            raise HTTPException(status_code=400, detail="Force/File number is mandatory.")

        clean_fnum = aggressive_clean_text(target_fnum)
        if hasattr(ActiveModel, 'f_num'): clean_data['f_num'] = clean_fnum
        if hasattr(ActiveModel, 'fnum'): clean_data['fnum'] = clean_fnum

        fnum_filter = []
        if hasattr(ActiveModel, 'f_num'): fnum_filter.append(func.trim(func.upper(ActiveModel.f_num)) == clean_fnum)
        if hasattr(ActiveModel, 'fnum'): fnum_filter.append(func.trim(func.upper(ActiveModel.fnum)) == clean_fnum)

        active_officer = db.query(ActiveModel).filter(or_(*fnum_filter)).first()
        if active_officer:
            raise HTTPException(status_code=400, detail="Duplicate Entry: This Force Number or File Number is currently active.")

        search_fnum = aggressive_clean_text(previous_fnum) if previous_fnum else clean_fnum
        arc_filter = []
        if hasattr(ArchiveModel, 'fnum'): arc_filter.append(func.trim(func.upper(ArchiveModel.fnum)) == search_fnum)
        if hasattr(ArchiveModel, 'f_num'): arc_filter.append(func.trim(func.upper(ArchiveModel.f_num)) == search_fnum)

        archived_officer = db.query(ArchiveModel).filter(or_(*arc_filter)).first()
        
        if archived_officer:
            if not reintegration_reason:
                return JSONResponse(
                    status_code=409, 
                    content={
                        "detail": "Officer history found in the archive. Please authorize re-entry.", 
                        "is_archived_returnee": True,
                        "old_rank": getattr(archived_officer, 'rank', 'N/A'),
                        "old_fnum": getattr(archived_officer, 'fnum', getattr(archived_officer, 'f_num', search_fnum))
                    }
                )
            
            clean_data['dob'] = getattr(archived_officer, 'dob', clean_data.get('dob'))
            clean_data['doe'] = getattr(archived_officer, 'doe', clean_data.get('doe'))
            clean_data['ipps'] = getattr(archived_officer, 'ipps', clean_data.get('ipps'))
            clean_data['status'] = "ACTIVE"
            
            valid_cols = [c.key for c in ActiveModel.__table__.columns]
            safe_payload = {k: v for k, v in clean_data.items() if k in valid_cols}
            
            new_record = ActiveModel(**safe_payload)
            new_record.last_updated_by = get_officer_signature(current_user)
            
            db.add(new_record)
            db.commit()
            
            assigned_id = getattr(new_record, 'id', getattr(new_record, 'sn', 1))

            record_neon_activity(
                logs_db=logs_db,
                fnum=current_user.fnum,
                action_type="REGISTER",
                module="NOMINAL_ROLL",
                target_id=clean_fnum,
                changes_summary=f"{current_user.fnum} {current_user.rank} {current_user.name} reintegrated officer {clean_fnum} ({clean_data.get('rank')}) from archives."
            )

            return {"status": "success", "message": f"Officer re-integrated successfully as {clean_data.get('rank')}", "id": assigned_id}

        valid_cols = [c.key for c in ActiveModel.__table__.columns]
        safe_payload = {k: v for k, v in clean_data.items() if k in valid_cols}
        
        new_record = ActiveModel(**safe_payload)
        new_record.last_updated_by = get_officer_signature(current_user)
        
        db.add(new_record)
        db.commit()
        db.refresh(new_record)
        
        assigned_id = getattr(new_record, 'id', getattr(new_record, 'sn', 1))

        record_neon_activity(
            logs_db=logs_db,
            fnum=current_user.fnum,
            action_type="REGISTER",
            module="NOMINAL_ROLL",
            target_id=clean_fnum,
            changes_summary=f"{current_user.fnum} {current_user.rank} {current_user.name} registered new personnel record for officer {clean_fnum}."
        )

        return {"status": "success", "message": "Officer recorded successfully.", "id": assigned_id}
        
    except IntegrityError:
        db.rollback() 
        logs_db.rollback()
        raise HTTPException(status_code=400, detail="Duplicate Entry: Force Number or IPPS already exists in active database.")
    except Exception as e:
        db.rollback()
        logs_db.rollback()
        raise HTTPException(status_code=500, detail=f"Database error: {str(e)}")

@router.put("/nominal-roll/archive-record")
def archive_personnel(
    payload: dict, 
    db: Session = Depends(get_db), 
    logs_db: Session = Depends(get_logs_db),
    current_user: models.Users = Depends(get_current_user)
):
    ActiveModel = get_active_model()
    ArchiveModel = get_archive_model()
    
    try:
        raw_fnum = payload.get("fnum") or payload.get("f_num")
        archive_reason = payload.get("archive_reason", "ADMINISTRATIVE")
        
        if not raw_fnum:
            raise HTTPException(status_code=400, detail="Missing Force Number identifier for archiving.")

        fnum_clean = str(raw_fnum).split('/ARCHIVE')[0].replace('/ARCHIVE', '').strip().upper()
        alt_fnum = fnum_clean.replace('/', '')
        
        query_filters = []
        if hasattr(ActiveModel, 'f_num'):
            query_filters.extend([
                func.trim(func.upper(ActiveModel.f_num)) == fnum_clean,
                func.trim(func.upper(ActiveModel.f_num)) == alt_fnum
            ])
        if hasattr(ActiveModel, 'fnum'):
            query_filters.extend([
                func.trim(func.upper(ActiveModel.fnum)) == fnum_clean,
                func.trim(func.upper(ActiveModel.fnum)) == alt_fnum
            ])
        if hasattr(ActiveModel, 'ipps'):
            query_filters.append(func.trim(func.upper(ActiveModel.ipps)) == fnum_clean)
            
        active_record = db.query(ActiveModel).filter(or_(*query_filters)).first()

        if not active_record:
            raise HTTPException(status_code=404, detail=f"Officer record '{fnum_clean}' not found in active roll.")

        record_data = active_record.__dict__.copy()
        record_data.pop("_sa_instance_state", None) 
        record_data.pop("id", None) 
        record_data.pop("sn", None) 
        
        if hasattr(ArchiveModel, 'fnum'): record_data["fnum"] = fnum_clean
        if hasattr(ArchiveModel, 'f_num'): record_data["f_num"] = fnum_clean
            
        record_data["status"] = "ARCHIVED"
        record_data["archive_reason"] = aggressive_clean_text(archive_reason)
        record_data["archive_date"] = datetime.now().date()
        record_data["last_updated_by"] = get_officer_signature(current_user)

        valid_archive_columns = [c.key for c in ArchiveModel.__table__.columns]
        safe_record_data = {k: v for k, v in record_data.items() if k in valid_archive_columns}

        archived_record = ArchiveModel(**safe_record_data)
        db.add(archived_record)
        db.delete(active_record)
        db.commit()

        record_neon_activity(
            logs_db=logs_db,
            fnum=current_user.fnum,
            action_type="DELETE",
            module="NOMINAL_ROLL",
            target_id=fnum_clean,
            changes_summary=f"{current_user.fnum} {current_user.rank} {current_user.name} archived personnel record {fnum_clean}. Reason: {archive_reason}."
        )
        
        return {"status": "success", "message": "Officer successfully moved to archives."}
        
    except HTTPException:
        raise
    except Exception as e:
        db.rollback()
        logs_db.rollback()
        raise HTTPException(status_code=500, detail=f"Failed to migrate record: {str(e)}")

@router.put("/nominal-roll/{identifier:path}")
def update_Nominal_Roll(
    identifier: str, 
    data: dict, 
    db: Session = Depends(get_db), 
    logs_db: Session = Depends(get_logs_db),
    current_user: models.Users = Depends(get_current_user)
):
    ActiveModel = get_active_model()
    clean_id = unquote(unquote(identifier)).strip().upper()
    
    query_filters = []
    
    if hasattr(ActiveModel, 'fnum'):
        query_filters.append(func.trim(func.upper(ActiveModel.fnum)) == clean_id)
    if hasattr(ActiveModel, 'f_num'):
        query_filters.append(func.trim(func.upper(ActiveModel.f_num)) == clean_id)
        
    alt_id = clean_id.replace('/', '')
    if hasattr(ActiveModel, 'fnum'):
        query_filters.append(func.trim(func.upper(ActiveModel.fnum)) == alt_id)
    if hasattr(ActiveModel, 'f_num'):
        query_filters.append(func.trim(func.upper(ActiveModel.f_num)) == alt_id)

    if clean_id.isdigit():
        pk_col = getattr(ActiveModel, 'id', getattr(ActiveModel, 'sn', None))
        if pk_col is not None:
            query_filters.append(pk_col == int(clean_id))

    officer = db.query(ActiveModel).filter(or_(*query_filters)).first()
    
    if not officer:
        raise HTTPException(status_code=404, detail=f"Officer record '{clean_id}' not found in active Nominal Roll.")

    data.pop('id', None)
    data.pop('sn', None)
    
    test_rank = str(data.get('rank', '')).upper()
    if any(stn_term in test_rank for stn_term in ["STATION", "POST", "DIV", "UNIT", "HQ", "HEADQUARTERS"]):
        raise HTTPException(status_code=400, detail="Validation Error: Station or unit names cannot be entered into the Rank field.")

    if 'educ_level' in data:
        data['educ_level'] = normalize_education_level(data['educ_level'])
        
    if 'contact' in data and data['contact']:
        data['contact'] = format_phone_number(data['contact'])
        
    if 'name' in data and data['name']:
        data['name'] = aggressive_clean_text(data['name'])

    if 'station' in data and data['station']:
        data['station'] = normalize_station_name(data['station'])

    if 'region' in data and data['region']:
        data['region'] = normalize_region_name(data['region'])

    for text_field in ['position', 'home_dist', 'tribe', 'bank_branch', 'section', 'dir']:
        if text_field in data and data[text_field]:
            data[text_field] = aggressive_clean_text(data[text_field])

    for date_field in ['dob', 'doe', 'do_post', 'do_pro']:
        if date_field in data and data[date_field]:
            data[date_field] = parse_safe_date(data[date_field])
    
    perms = current_user.permissions or {}
    user_role = (current_user.role or "").upper()
    is_global_user = (
        user_role in ["SUPER_ADMIN", "ADMIN", "RPC", "DEPUTY COMMANDER"] or
        "HR" in (current_user.position or "").upper() or
        perms.get("view_global_roster") is True or
        perms.get("global_observer") is True
    )

    if not is_global_user:
        data.pop('region', None)
        data.pop('station', None)

    for key, value in data.items():
        if hasattr(officer, key):
            setattr(officer, key, value if value != "" else None)

    officer.last_updated_by = get_officer_signature(current_user)
    
    try:
        db.commit()
        db.refresh(officer)

        record_neon_activity(
            logs_db=logs_db,
            fnum=current_user.fnum,
            action_type="UPDATE",
            module="NOMINAL_ROLL",
            target_id=clean_id,
            changes_summary=f"{current_user.fnum} {current_user.rank} {current_user.name} updated personnel record for officer {clean_id}."
        )

        return {"status": "success", "message": f"Officer record updated successfully."}
    except Exception as e:
        db.rollback()
        logs_db.rollback()
        raise HTTPException(status_code=500, detail=f"Failed to update officer record: {str(e)}")

@router.get("/nominal-roll/archive")
@router.get("/nominal-roll-archive")
def get_archived_personnel(
    db: Session = Depends(get_db), 
    logs_db: Session = Depends(get_logs_db),
    current_user: models.Users = Depends(get_current_user)
):
    try:
        ArchiveModel = get_archive_model()
        query = get_scoped_nominal_query(db, current_user, ArchiveModel)
        
        archives = sorted(query.all(), key=hierarchical_sort_key)
        
        clean_list = []
        for a in archives:
            if is_invalid_roster_entry(a): continue
            d = a.__dict__.copy()
            d.pop("_sa_instance_state", None)
            d['educ_level'] = normalize_education_level(d.get('educlevel') or d.get('educ_level'))
            for k, v in d.items():
                if hasattr(v, 'isoformat'):
                    d[k] = str(v)
            clean_list.append(d)

        record_neon_activity(
            logs_db=logs_db,
            fnum=current_user.fnum,
            action_type="VIEW",
            module="NOMINAL_ROLL_ARCHIVE",
            target_id="ARCHIVED_PERSONNEL",
            changes_summary=f"{current_user.fnum} {current_user.rank} {current_user.name} inspected archived personnel roster (Fetched {len(clean_list)} records)."
        )

        return clean_list
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to fetch archives: {str(e)}")

@router.post("/nominal-roll/bulk-archive")
def bulk_archive_personnel(
    payload: dict,
    db: Session = Depends(get_db),
    logs_db: Session = Depends(get_logs_db),
    current_user: models.Users = Depends(get_current_user)
):
    ActiveModel = get_active_model()
    ArchiveModel = get_archive_model()
    
    fnums = payload.get("fnums", [])
    archive_reason = payload.get("archive_reason", "ADMINISTRATIVE")
    
    if not fnums:
        raise HTTPException(status_code=400, detail="No officers specified for bulk archive.")
        
    success_count = 0
    fail_count = 0
    officer_sig = get_officer_signature(current_user)

    try:
        for fnum in fnums:
            fnum_clean = unquote(unquote(str(fnum))).strip().upper()
            
            query_filters = []
            if hasattr(ActiveModel, 'f_num'):
                query_filters.append(func.trim(func.upper(ActiveModel.f_num)) == fnum_clean)
            if hasattr(ActiveModel, 'fnum'):
                query_filters.append(func.trim(func.upper(ActiveModel.fnum)) == fnum_clean)
            if hasattr(ActiveModel, 'ipps'):
                query_filters.append(func.trim(func.upper(ActiveModel.ipps)) == fnum_clean)
                
            active_record = db.query(ActiveModel).filter(or_(*query_filters)).first()
            
            if not active_record:
                alt_fnum = fnum_clean.replace('/', '')
                query_filters_alt = []
                if hasattr(ActiveModel, 'f_num'): query_filters_alt.append(func.trim(func.upper(ActiveModel.f_num)) == alt_fnum)
                if hasattr(ActiveModel, 'fnum'): query_filters_alt.append(func.trim(func.upper(ActiveModel.fnum)) == alt_fnum)
                active_record = db.query(ActiveModel).filter(or_(*query_filters_alt)).first()

            if active_record:
                record_data = active_record.__dict__.copy()
                record_data.pop("_sa_instance_state", None)
                record_data.pop("id", None)
                record_data.pop("sn", None)
                
                if hasattr(ArchiveModel, 'fnum'): record_data["fnum"] = fnum_clean
                if hasattr(ArchiveModel, 'f_num'): record_data["f_num"] = fnum_clean
                
                record_data["status"] = "ARCHIVED"
                record_data["archive_reason"] = aggressive_clean_text(archive_reason)
                record_data["archive_date"] = datetime.now().date()
                record_data["last_updated_by"] = officer_sig

                valid_archive_columns = [c.key for c in ArchiveModel.__table__.columns]
                safe_record_data = {k: v for k, v in record_data.items() if k in valid_archive_columns}

                archived_record = ArchiveModel(**safe_record_data)
                db.add(archived_record)
                db.delete(active_record)
                success_count += 1
            else:
                fail_count += 1

        db.commit()

        record_neon_activity(
            logs_db=logs_db,
            fnum=current_user.fnum,
            action_type="DELETE",
            module="NOMINAL_ROLL",
            target_id="BULK_ARCHIVE",
            changes_summary=f"{current_user.fnum} {current_user.rank} {current_user.name} executed bulk personnel archive ({success_count} archived, {fail_count} failed)."
        )

        return {
            "status": "success", 
            "success_count": success_count, 
            "fail_count": fail_count,
            "message": f"Bulk archive complete: {success_count} succeeded, {fail_count} failed."
        }
    except Exception as e:
        db.rollback()
        logs_db.rollback()
        raise HTTPException(status_code=500, detail=f"Bulk archive transaction failed: {str(e)}")

@router.get("/nominal-roll/export-analytics")
def export_nominal_roll_analytics(
    region: str = "ALL REGIONS",
    station: str = "ALL STATIONS",
    db: Session = Depends(get_db),
    logs_db: Session = Depends(get_logs_db),
    current_user: models.Users = Depends(require_export_privilege)
):
    try:
        ActiveModel = get_active_model()
        query = get_scoped_nominal_query(db, current_user, ActiveModel)
        
        region_clean = normalize_region_name(region)
        station_clean = normalize_station_name(station)

        records = query.all()
        valid_records = [r for r in records if not is_invalid_roster_entry(r)]

        filtered_records = []
        for r in valid_records:
            parsed_stn = parseStationHierarchy(str(getattr(r, 'station', '')).strip().upper())
            r_stn = parsed_stn['station']
            r_reg = getOfficialRegionForStation(r_stn, str(getattr(r, 'region', '')).strip().upper())

            if region_clean != "ALL REGIONS" and r_reg != region_clean: continue
            if station_clean != "ALL STATIONS" and r_stn != station_clean: continue
            filtered_records.append(r)

        total_personnel = len(filtered_records)
        rank_counts = {}
        sex_counts = {}
        status_counts = {}
        station_counts = {}

        for r in filtered_records:
            rk = clean_str(getattr(r, 'rank', 'UNKNOWN'))
            sx = clean_str(getattr(r, 'sex', 'UNSPECIFIED'))
            st = normalize_personnel_status(getattr(r, 'status', 'ACTIVE'))
            stn = normalize_station_name(getattr(r, 'station', 'HQ'))

            rank_counts[rk] = rank_counts.get(rk, 0) + 1
            sex_counts[sx] = sex_counts.get(sx, 0) + 1
            status_counts[st] = status_counts.get(st, 0) + 1
            station_counts[stn] = station_counts.get(stn, 0) + 1

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Nominal Analytics"
        ws.views.sheetView[0].showGridLines = True

        eat_tz = pytz.timezone("Africa/Nairobi")
        eat_time = datetime.now(eat_tz).replace(tzinfo=None)
        
        officer_name = (current_user.name or "UNKNOWN").strip().upper()
        officer_fnum = (current_user.fnum or "UNKNOWN").strip().upper()
        officer_rank = (current_user.rank or "").strip().upper()
        officer_station = normalize_station_name(current_user.station or "HQ")
        timestamp_str = eat_time.strftime("%Y-%m-%d %H:%M:%S EAT")

        wb.properties.creator = f"{officer_fnum} {officer_rank} {officer_name}"
        wb.properties.lastModifiedBy = f"KMP-CSDMS SECURITY AUDIT ENGINE [{officer_station}]"
        wb.properties.title = "Official Law Enforcement Restricted Nominal Roll Analytics Export"
        wb.properties.subject = f"Nominal Roll Analytics Extract — Downloaded by {officer_fnum} ({officer_station}) at {timestamp_str}"
        wb.properties.description = f"SECURE OPSEC RECORD. Cryptographically keyed and audited to Force Number: {officer_fnum}."

        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell_font = Font(name="Calibri", size=10)
        thin_border = Border(
            left=Side(style='thin', color='D1D5DB'), right=Side(style='thin', color='D1D5DB'),
            top=Side(style='thin', color='D1D5DB'), bottom=Side(style='thin', color='D1D5DB')
        )

        ws.merge_cells("A1:C1")
        banner_cell = ws.cell(row=1, column=1)
        banner_cell.value = f"RESTRICTED LAW ENFORCEMENT RECORD — FORENSICALLY STAMPED TO OFFICER: {officer_fnum} {officer_rank} {officer_name} ({officer_station}) — TIMESTAMP: {timestamp_str}"
        banner_cell.font = Font(name="Calibri", size=9, bold=True, color="991B1B")
        banner_cell.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
        banner_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 22

        ws.cell(row=2, column=1, value=f"UGANDA POLICE FORCE - NOMINAL ROLL ANALYTICS SUMMARY").font = Font(bold=True, size=11)
        ws.cell(row=3, column=1, value=f"Scope: Station [{station_clean}] | Region [{region_clean}] | Total Personnel: {total_personnel}").font = Font(size=10)
        ws.row_dimensions[2].height = 18
        ws.row_dimensions[3].height = 18

        current_row = 5

        def write_summary_table(title, data_dict):
            nonlocal current_row
            ws.cell(row=current_row, column=1, value=title).font = Font(bold=True, size=11, color="002060")
            current_row += 1

            headers = ["Category / Attribute", "Total Personnel", "Percentage (%)"]
            ws.row_dimensions[current_row].height = 24
            for col_idx, h in enumerate(headers, 1):
                cell = ws.cell(row=current_row, column=col_idx, value=h)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_align
                cell.border = thin_border
            current_row += 1

            sorted_data = sorted(data_dict.items(), key=lambda x: x[1], reverse=True)
            for k, v in sorted_data:
                pct = round((v / total_personnel) * 100, 2) if total_personnel > 0 else 0
                ws.row_dimensions[current_row].height = 20
                r_cells = [
                    ws.cell(row=current_row, column=1, value=k),
                    ws.cell(row=current_row, column=2, value=v),
                    ws.cell(row=current_row, column=3, value=f"{pct}%")
                ]
                r_cells[0].alignment = Alignment(horizontal="left", vertical="center")
                r_cells[1].alignment = Alignment(horizontal="center", vertical="center")
                r_cells[2].alignment = Alignment(horizontal="center", vertical="center")
                for c in r_cells:
                    c.font = cell_font
                    c.border = thin_border
                current_row += 1
            current_row += 2

        write_summary_table("RANK STRENGTH DISTRIBUTION", rank_counts)
        write_summary_table("GENDER DEMOGRAPHICS", sex_counts)
        write_summary_table("PERSONNEL STATUS BREAKDOWN", status_counts)
        write_summary_table("STATION DEPLOYMENT CONCENTRATION", station_counts)

        ws.column_dimensions['A'].width = 35
        ws.column_dimensions['B'].width = 18
        ws.column_dimensions['C'].width = 18

        excel_stream = io.BytesIO()
        wb.save(excel_stream)
        excel_stream.seek(0)

        zip_stream = io.BytesIO()
        zip_password = str(current_user.fnum).strip().encode('utf-8')
        fnum_clean = str(current_user.fnum).replace('/', '_').upper()
        excel_filename = f"{fnum_clean}_Nominal_Roll_Analytics_{eat_time.strftime('%Y%m%d')}.xlsx"
        zip_filename = f"SECURE_NOMINAL_ANALYTICS_{eat_time.strftime('%Y%m%d')}.zip"

        with pyzipper.AESZipFile(zip_stream, 'w', compression=pyzipper.ZIP_DEFLATED, encryption=pyzipper.WZ_AES) as zf:
            zf.setpassword(zip_password)
            zf.writestr(excel_filename, excel_stream.getvalue())

        zip_stream.seek(0)

        record_neon_activity(
            logs_db=logs_db,
            fnum=current_user.fnum,
            action_type="EXPORT",
            module="NOMINAL_ROLL_EXPORT",
            target_id="ANALYTICS_EXPORT",
            changes_summary=f"{current_user.fnum} {current_user.rank} {current_user.name} securely generated AES-256 encrypted archive for Nominal Roll Analytics ({total_personnel} personnel audited)."
        )

        return StreamingResponse(
            zip_stream,
            media_type="application/zip",
            headers={
                'Content-Disposition': f'attachment; filename="{zip_filename}"',
                'Access-Control-Expose-Headers': 'Content-Disposition'
            }
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Analytics Export Failed: {str(e)}")

@router.get("/nominal-roll/export-missing-audit")
def export_missing_info_audit(
    region: str = "ALL REGIONS",
    station: str = "ALL STATIONS",
    db: Session = Depends(get_db),
    logs_db: Session = Depends(get_logs_db),
    current_user: models.Users = Depends(require_export_privilege)
):
    try:
        ActiveModel = get_active_model()
        query = get_scoped_nominal_query(db, current_user, ActiveModel)
        
        region_clean = normalize_region_name(region)
        station_clean = normalize_station_name(station)

        records = sorted(query.all(), key=hierarchical_sort_key)
        missing_rows = []

        for r in records:
            if is_invalid_roster_entry(r): continue
            
            parsed_stn = parseStationHierarchy(str(getattr(r, 'station', '')).strip().upper())
            r_stn = parsed_stn['station'] 
            r_reg = getOfficialRegionForStation(r_stn, str(getattr(r, 'region', '')).strip().upper())

            if region_clean != "ALL REGIONS" and r_reg != region_clean: continue
            if station_clean != "ALL STATIONS" and r_stn != station_clean: continue

            rank = str(getattr(r, 'rank', '')).strip().upper()
            is_constable_tier = rank in ['PC', 'DC', 'D/C', 'CONSTABLE', 'C/DRV', 'DRV'] or 'DRV' in rank

            dob = getattr(r, 'dob', None)
            doe = getattr(r, 'doe', None)
            contact = getattr(r, 'contact', None)
            nin = getattr(r, 'nin', None)
            ipps = getattr(r, 'ipps', None)
            dopro = getattr(r, 'do_pro', None)

            missing_fields = []
            if not dob: missing_fields.append("DOB")
            if not doe: missing_fields.append("DOE")
            if not contact: missing_fields.append("Contact")
            if not nin: missing_fields.append("NIN")
            if not ipps: missing_fields.append("IPPS")
            if not is_constable_tier and not dopro: missing_fields.append("DO_PRO")

            if missing_fields:
                missing_rows.append({
                    "Force Number": getattr(r, 'f_num', getattr(r, 'fnum', '')),
                    "Rank": rank,
                    "Name": getattr(r, 'name', ''),
                    "Region": r_reg,
                    "Station": r_stn,
                    "Post": parsed_stn['post'] if parsed_stn['post'] else "N/A",
                    "Missing Fields": " | ".join(missing_fields)
                })

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Missing Info Audit"
        ws.views.sheetView[0].showGridLines = True

        eat_tz = pytz.timezone("Africa/Nairobi")
        eat_time = datetime.now(eat_tz).replace(tzinfo=None)
        
        officer_name = (current_user.name or "UNKNOWN").strip().upper()
        officer_fnum = (current_user.fnum or "UNKNOWN").strip().upper()
        officer_rank = (current_user.rank or "").strip().upper()
        officer_station = normalize_station_name(current_user.station or "HQ")
        timestamp_str = eat_time.strftime("%Y-%m-%d %H:%M:%S EAT")

        wb.properties.creator = f"{officer_fnum} {officer_rank} {officer_name}"
        wb.properties.lastModifiedBy = f"KMP-CSDMS SECURITY AUDIT ENGINE [{officer_station}]"
        wb.properties.title = "Official Law Enforcement Restricted Missing Info Audit Export"
        wb.properties.subject = f"Missing Info Audit Extract — Downloaded by {officer_fnum} ({officer_station}) at {timestamp_str}"
        wb.properties.description = f"SECURE OPSEC RECORD. Cryptographically keyed and audited to Force Number: {officer_fnum}."

        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell_font = Font(name="Calibri", size=10)
        thin_border = Border(
            left=Side(style='thin', color='D1D5DB'), right=Side(style='thin', color='D1D5DB'),
            top=Side(style='thin', color='D1D5DB'), bottom=Side(style='thin', color='D1D5DB')
        )

        ws.merge_cells("A1:H1")
        banner_cell = ws.cell(row=1, column=1)
        banner_cell.value = f"RESTRICTED LAW ENFORCEMENT RECORD — FORENSICALLY STAMPED TO OFFICER: {officer_fnum} {officer_rank} {officer_name} ({officer_station}) — TIMESTAMP: {timestamp_str}"
        banner_cell.font = Font(name="Calibri", size=9, bold=True, color="991B1B")
        banner_cell.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
        banner_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 22

        ws.cell(row=2, column=1, value=f"UGANDA POLICE FORCE - MANPOWER AUDIT: MISSING INFORMATION REPORT").font = Font(bold=True, size=11)
        ws.cell(row=3, column=1, value=f"Station / Unit: {station_clean} (Region: {region_clean})").font = Font(size=10)
        ws.row_dimensions[2].height = 18
        ws.row_dimensions[3].height = 18

        headers = ["SN", "Force Number", "Rank", "Name", "Region", "Station", "Police Post", "Missing Fields"]
        ws.row_dimensions[5].height = 26
        for col_idx, header_text in enumerate(headers, 1):
            cell = ws.cell(row=5, column=col_idx, value=header_text)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = thin_border

        for idx, row in enumerate(missing_rows):
            row_data = [
                idx + 1,
                row["Force Number"],
                row["Rank"],
                row["Name"],
                row["Region"],
                row["Station"],
                row["Post"],
                row["Missing Fields"]
            ]
            ws.append(row_data)
            current_row = idx + 6
            ws.row_dimensions[current_row].height = 22

            for col_idx in range(1, len(row_data) + 1):
                cell = ws.cell(row=current_row, column=col_idx)
                cell.font = cell_font
                cell.border = thin_border
                if col_idx in [1, 2, 3, 5, 6, 7]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

        column_widths = {'A': 8, 'B': 18, 'C': 12, 'D': 25, 'E': 18, 'F': 20, 'G': 18, 'H': 35}
        for col_letter, width in column_widths.items():
            ws.column_dimensions[col_letter].width = width

        excel_stream = io.BytesIO()
        wb.save(excel_stream)
        excel_stream.seek(0)

        zip_stream = io.BytesIO()
        zip_password = str(current_user.fnum).strip().encode('utf-8')
        fnum_clean = str(current_user.fnum).replace('/', '_').upper()
        excel_filename = f"{fnum_clean}_Missing_Info_Audit_{eat_time.strftime('%Y%m%d')}.xlsx"
        zip_filename = f"SECURE_MISSING_AUDIT_{eat_time.strftime('%Y%m%d')}.zip"

        with pyzipper.AESZipFile(zip_stream, 'w', compression=pyzipper.ZIP_DEFLATED, encryption=pyzipper.WZ_AES) as zf:
            zf.setpassword(zip_password)
            zf.writestr(excel_filename, excel_stream.getvalue())

        zip_stream.seek(0)

        record_neon_activity(
            logs_db=logs_db,
            fnum=current_user.fnum,
            action_type="EXPORT",
            module="NOMINAL_ROLL_EXPORT",
            target_id="MISSING_INFO_AUDIT",
            changes_summary=f"{current_user.fnum} {current_user.rank} {current_user.name} securely generated AES-256 encrypted archive for Missing Info Audit ({len(missing_rows)} flagged records)."
        )

        return StreamingResponse(
            zip_stream,
            media_type="application/zip",
            headers={
                'Content-Disposition': f'attachment; filename="{zip_filename}"',
                'Access-Control-Expose-Headers': 'Content-Disposition'
            }
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Audit Export Failed: {str(e)}")

@router.get("/nominal-roll/export-station-ledger")
def export_station_nominal_roll(
    region: str = "ALL REGIONS",
    station: str = "ALL STATIONS",
    db: Session = Depends(get_db),
    logs_db: Session = Depends(get_logs_db),
    current_user: models.Users = Depends(require_export_privilege)
):
    try:
        ActiveModel = get_active_model()
        query = get_scoped_nominal_query(db, current_user, ActiveModel)
        
        region_clean = normalize_region_name(region)
        station_clean = normalize_station_name(station)

        records = sorted(query.all(), key=hierarchical_sort_key)
        station_rows = []

        for r in records:
            if is_invalid_roster_entry(r): continue
            
            parsed_stn = parseStationHierarchy(str(getattr(r, 'station', '')).strip().upper())
            r_stn = parsed_stn['station']
            r_reg = getOfficialRegionForStation(r_stn, str(getattr(r, 'region', '')).strip().upper())

            if region_clean != "ALL REGIONS" and r_reg != region_clean: continue
            if station_clean != "ALL STATIONS" and r_stn != station_clean: continue

            station_rows.append({
                "Force Number": getattr(r, 'f_num', getattr(r, 'fnum', '')),
                "Rank": getattr(r, 'rank', ''),
                "Name": getattr(r, 'name', ''),
                "Sex": getattr(r, 'sex', ''),
                "Position": getattr(r, 'position', ''),
                "Contact": getattr(r, 'contact', ''),
                "IPPS": getattr(r, 'ipps', ''),
                "NIN": getattr(r, 'nin', ''),
                "TIN": getattr(r, 'tin', ''),
                "DOB": getattr(r, 'dob', ''),
                "DOE": getattr(r, 'doe', ''),
                "Date of Post": getattr(r, 'do_post', getattr(r, 'dopost', '')),
                "Date of Promotion": getattr(r, 'do_pro', getattr(r, 'dopro', '')),
                "Education Level": getattr(r, 'educ_level', getattr(r, 'educlevel', '')),
                "Home District": getattr(r, 'home_dist', getattr(r, 'homedist', '')),
                "Tribe": getattr(r, 'tribe', ''),
                "Bank Branch": getattr(r, 'bank_branch', getattr(r, 'bankbranch', '')),
                "Account Number": getattr(r, 'acc_no', getattr(r, 'accno', '')),
                "Section": getattr(r, 'section', ''),
                "Directorate": getattr(r, 'dir', ''),
                "Station": r_stn,
                "Police Post": parsed_stn['post'] if parsed_stn['post'] else "N/A",
                "District": getattr(r, 'district', ''),
                "Region": r_reg,
                "Status": getattr(r, 'status', 'ACTIVE'),
                "Last Updated By": getattr(r, 'last_updated_by', '')
            })

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"Nominal Roll - {station_clean}"
        ws.views.sheetView[0].showGridLines = True

        eat_tz = pytz.timezone("Africa/Nairobi")
        eat_time = datetime.now(eat_tz).replace(tzinfo=None)
        
        officer_name = (current_user.name or "UNKNOWN").strip().upper()
        officer_fnum = (current_user.fnum or "UNKNOWN").strip().upper()
        officer_rank = (current_user.rank or "").strip().upper()
        officer_station = normalize_station_name(current_user.station or "HQ")
        timestamp_str = eat_time.strftime("%Y-%m-%d %H:%M:%S EAT")

        wb.properties.creator = f"{officer_fnum} {officer_rank} {officer_name}"
        wb.properties.lastModifiedBy = f"KMP-CSDMS SECURITY AUDIT ENGINE [{officer_station}]"
        wb.properties.title = "Official Law Enforcement Restricted Full Nominal Roll Export"
        wb.properties.subject = f"Full Station Nominal Roll Extract — Downloaded by {officer_fnum} ({officer_station}) at {timestamp_str}"
        wb.properties.description = f"SECURE OPSEC RECORD. Cryptographically keyed and audited to Force Number: {officer_fnum}."

        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell_font = Font(name="Calibri", size=10)
        thin_border = Border(
            left=Side(style='thin', color='D1D5DB'), right=Side(style='thin', color='D1D5DB'),
            top=Side(style='thin', color='D1D5DB'), bottom=Side(style='thin', color='D1D5DB')
        )

        ws.merge_cells("A1:AA1")
        banner_cell = ws.cell(row=1, column=1)
        banner_cell.value = f"RESTRICTED LAW ENFORCEMENT RECORD — FORENSICALLY STAMPED TO OFFICER: {officer_fnum} {officer_rank} {officer_name} ({officer_station}) — TIMESTAMP: {timestamp_str}"
        banner_cell.font = Font(name="Calibri", size=9, bold=True, color="991B1B")
        banner_cell.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
        banner_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 22

        ws.cell(row=2, column=1, value=f"UGANDA POLICE FORCE - MASTER NOMINAL ROLL LEDGER (FULL DETAILS)").font = Font(bold=True, size=11)
        ws.cell(row=3, column=1, value=f"Station / Unit: {station_clean} (Region: {region_clean})").font = Font(size=10)
        ws.row_dimensions[2].height = 18
        ws.row_dimensions[3].height = 18

        headers = [
            "SN", "Force Number", "Rank", "Name", "Sex", "Position", "Contact", 
            "IPPS", "NIN", "TIN", "DOB", "DOE", "Date of Post", "Date of Promotion", 
            "Education Level", "Home District", "Tribe", "Bank Branch", "Account Number", 
            "Section", "Directorate", "Station", "Police Post", "District", "Region", "Status", "Last Updated By"
        ]
        ws.row_dimensions[5].height = 26
        for col_idx, header_text in enumerate(headers, 1):
            cell = ws.cell(row=5, column=col_idx, value=header_text)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = thin_border

        for idx, row in enumerate(station_rows):
            row_data = [
                idx + 1,
                row["Force Number"],
                row["Rank"],
                row["Name"],
                row["Sex"],
                row["Position"],
                row["Contact"],
                row["IPPS"],
                row["NIN"],
                row["TIN"],
                str(row["DOB"]) if row["DOB"] else "",
                str(row["DOE"]) if row["DOE"] else "",
                str(row["Date of Post"]) if row["Date of Post"] else "",
                str(row["Date of Promotion"]) if row["Date of Promotion"] else "",
                row["Education Level"],
                row["Home District"],
                row["Tribe"],
                row["Bank Branch"],
                row["Account Number"],
                row["Section"],
                row["Directorate"],
                row["Station"],
                row["Police Post"],
                row["District"],
                row["Region"],
                row["Status"],
                row["Last Updated By"]
            ]
            ws.append(row_data)
            current_row = idx + 6
            ws.row_dimensions[current_row].height = 22

            for col_idx in range(1, len(row_data) + 1):
                cell = ws.cell(row=current_row, column=col_idx)
                cell.font = cell_font
                cell.border = thin_border
                if col_idx in [1, 2, 3, 5, 8, 9, 10, 11, 12, 13, 14, 22, 23, 24, 25, 26]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

        for col in ws.columns:
            col_letter = col[0].column_letter
            max_len = max([len(str(cell.value or '')) for cell in col], default=0)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

        excel_stream = io.BytesIO()
        wb.save(excel_stream)
        excel_stream.seek(0)

        zip_stream = io.BytesIO()
        zip_password = str(current_user.fnum).strip().encode('utf-8')
        fnum_clean = str(current_user.fnum).replace('/', '_').upper()
        excel_filename = f"{fnum_clean}_Full_Nominal_Roll_{station_clean.replace(' ', '_')}_{eat_time.strftime('%Y%m%d')}.xlsx"
        zip_filename = f"SECURE_FULL_STATION_LEDGER_{eat_time.strftime('%Y%m%d')}.zip"

        with pyzipper.AESZipFile(zip_stream, 'w', compression=pyzipper.ZIP_DEFLATED, encryption=pyzipper.WZ_AES) as zf:
            zf.setpassword(zip_password)
            zf.writestr(excel_filename, excel_stream.getvalue())

        zip_stream.seek(0)

        record_neon_activity(
            logs_db=logs_db,
            fnum=current_user.fnum,
            action_type="EXPORT",
            module="NOMINAL_ROLL_EXPORT",
            target_id="STATION_LEDGER",
            changes_summary=f"{current_user.fnum} {current_user.rank} {current_user.name} securely generated AES-256 encrypted archive for Full Station Nominal Roll ({len(station_rows)} personnel records for {station_clean})."
        )

        return StreamingResponse(
            zip_stream,
            media_type="application/zip",
            headers={
                'Content-Disposition': f'attachment; filename="{zip_filename}"',
                'Access-Control-Expose-Headers': 'Content-Disposition'
            }
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Full Station Ledger Export Failed: {str(e)}")

@router.get("/nominal-roll/export-filtered-ledger")
def export_filtered_nominal_roll(
    region: str = "ALL REGIONS",
    station: str = "ALL STATIONS",
    filter_type: str = "RANK",
    filter_value: str = "ALL",
    db: Session = Depends(get_db),
    logs_db: Session = Depends(get_logs_db),
    current_user: models.Users = Depends(require_export_privilege)
):
    try:
        ActiveModel = get_active_model()
        query = get_scoped_nominal_query(db, current_user, ActiveModel)
        
        region_clean = normalize_region_name(region)
        station_clean = normalize_station_name(station)
        f_type = filter_type.strip().upper()
        f_val = filter_value.strip().upper()

        records = sorted(query.all(), key=hierarchical_sort_key)
        filtered_rows = []

        for r in records:
            if is_invalid_roster_entry(r): continue
            
            parsed_stn = parseStationHierarchy(str(getattr(r, 'station', '')).strip().upper())
            r_stn = parsed_stn['station']
            r_reg = getOfficialRegionForStation(r_stn, str(getattr(r, 'region', '')).strip().upper())

            if region_clean != "ALL REGIONS" and r_reg != region_clean: continue
            if station_clean != "ALL STATIONS" and r_stn != station_clean: continue

            if f_val != "ALL":
                if f_type == "RANK":
                    r_rank = clean_str(getattr(r, 'rank', ''))
                    if 'DRV' in r_rank: r_rank = 'PC'
                    if r_rank != f_val: continue
                elif f_type == "UNIT":
                    if r_stn != f_val: continue
                elif f_type == "STATUS":
                    r_status = normalize_personnel_status(getattr(r, 'status', ''))
                    if r_status != f_val: continue
                elif f_type == "SEX":
                    r_sex = clean_str(getattr(r, 'sex', ''))
                    if r_sex != f_val: continue
                elif f_type == "AGE":
                    r_age = calculate_age_bracket(getattr(r, 'dob', None))
                    if r_age != f_val: continue
                elif f_type == "SECTION":
                    r_sec = normalize_section_unit(getattr(r, 'position', '') or getattr(r, 'section', ''))
                    if r_sec != f_val: continue

            filtered_rows.append({
                "Force Number": getattr(r, 'f_num', getattr(r, 'fnum', '')),
                "Rank": getattr(r, 'rank', ''),
                "Name": getattr(r, 'name', ''),
                "Sex": getattr(r, 'sex', ''),
                "Position": getattr(r, 'position', ''),
                "Contact": getattr(r, 'contact', ''),
                "IPPS": getattr(r, 'ipps', ''),
                "NIN": getattr(r, 'nin', ''),
                "Station": r_stn,
                "Region": r_reg,
                "Status": getattr(r, 'status', 'ACTIVE')
            })

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"Filtered List - {f_val}"
        ws.views.sheetView[0].showGridLines = True

        eat_tz = pytz.timezone("Africa/Nairobi")
        eat_time = datetime.now(eat_tz).replace(tzinfo=None)
        officer_name = (current_user.name or "UNKNOWN").strip().upper()
        officer_fnum = (current_user.fnum or "UNKNOWN").strip().upper()
        officer_rank = (current_user.rank or "").strip().upper()
        timestamp_str = eat_time.strftime("%Y-%m-%d %H:%M:%S EAT")

        wb.properties.creator = f"{officer_fnum} {officer_rank} {officer_name}"
        wb.properties.title = f"Official Filtered Extract: {f_type} - {f_val}"
        
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
        thin_border = Border(
            left=Side(style='thin', color='D1D5DB'), right=Side(style='thin', color='D1D5DB'),
            top=Side(style='thin', color='D1D5DB'), bottom=Side(style='thin', color='D1D5DB')
        )

        headers = ["SN", "Force Number", "Rank", "Name", "Sex", "Position", "Contact", "IPPS", "NIN", "Station", "Region", "Status"]
        
        ws.append(headers)
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for idx, row in enumerate(filtered_rows):
            row_data = [
                idx + 1, row["Force Number"], row["Rank"], row["Name"], row["Sex"], 
                row["Position"], row["Contact"], row["IPPS"], row["NIN"], 
                row["Station"], row["Region"], row["Status"]
            ]
            ws.append(row_data)

        for col in ws.columns:
            col_letter = col[0].column_letter
            max_len = max([len(str(cell.value or '')) for cell in col], default=0)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

        zip_stream = io.BytesIO()
        zip_password = str(current_user.fnum).strip().encode('utf-8')
        excel_filename = f"{officer_fnum}_Filtered_List_{f_type}_{eat_time.strftime('%Y%m%d')}.xlsx"
        zip_filename = f"SECURE_FILTERED_LEDGER_{eat_time.strftime('%Y%m%d')}.zip"

        with pyzipper.AESZipFile(zip_stream, 'w', compression=pyzipper.ZIP_DEFLATED, encryption=pyzipper.WZ_AES) as zf:
            zf.setpassword(zip_password)
            excel_buffer = io.BytesIO()
            wb.save(excel_buffer)
            zf.writestr(excel_filename, excel_buffer.getvalue())

        zip_stream.seek(0)
        return StreamingResponse(
            zip_stream,
            media_type="application/zip",
            headers={
                'Content-Disposition': f'attachment; filename="{zip_filename}"',
                'Access-Control-Expose-Headers': 'Content-Disposition'
            }
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Filtered Ledger Export Failed: {str(e)}")

@router.get("/nominal-roll/ranks-list")
def get_nominal_roll_ranks(
    db: Session = Depends(get_db),
    current_user: models.Users = Depends(get_current_user)
):
    try:
        ActiveModel = get_active_model()
        query = get_scoped_nominal_query(db, current_user, ActiveModel)
        records = query.all()
        
        ranks_set = set()
        for r in records:
            if is_invalid_roster_entry(r): continue
            rk = clean_str(getattr(r, 'rank', ''))
            if rk:
                ranks_set.add(rk)
                
        return sorted(list(ranks_set))
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to fetch ranks list: {str(e)}")

@router.get("/nominal-roll/audit-missing-preview")
def audit_missing_preview(
    region: str = "ALL REGIONS",
    station: str = "ALL STATIONS",
    db: Session = Depends(get_db),
    current_user: models.Users = Depends(require_export_privilege)
):
    try:
        ActiveModel = get_active_model()
        query = get_scoped_nominal_query(db, current_user, ActiveModel)
        
        region_clean = normalize_region_name(region)
        station_clean = normalize_station_name(station)

        records = sorted(query.all(), key=hierarchical_sort_key)
        missing_rows = []

        for r in records:
            if is_invalid_roster_entry(r): continue
            
            parsed_stn = parseStationHierarchy(str(getattr(r, 'station', '')).strip().upper())
            r_stn = parsed_stn['station']
            r_reg = getOfficialRegionForStation(r_stn, str(getattr(r, 'region', '')).strip().upper())

            if region_clean != "ALL REGIONS" and r_reg != region_clean: continue
            if station_clean != "ALL STATIONS" and r_stn != station_clean: continue

            rank = str(getattr(r, 'rank', '')).strip().upper()
            is_constable_tier = rank in ['PC', 'DC', 'D/C', 'CONSTABLE', 'C/DRV', 'DRV'] or 'DRV' in rank

            dob = getattr(r, 'dob', None)
            doe = getattr(r, 'doe', None)
            contact = getattr(r, 'contact', None)
            nin = getattr(r, 'nin', None)
            ipps = getattr(r, 'ipps', None)
            dopro = getattr(r, 'do_post', getattr(r, 'dopro', None))

            missing_fields = []
            if not dob: missing_fields.append("DOB")
            if not doe: missing_fields.append("DOE")
            if not contact: missing_fields.append("Contact")
            if not nin: missing_fields.append("NIN")
            if not ipps: missing_fields.append("IPPS")
            if not is_constable_tier and not dopro: missing_fields.append("DO_PRO")

            if missing_rows and len(missing_rows) >= 500: break

            if missing_fields:
                missing_rows.append({
                    "force_number": getattr(r, 'f_num', getattr(r, 'fnum', '')),
                    "rank": rank,
                    "name": getattr(r, 'name', ''),
                    "region": r_reg,
                    "station": r_stn,
                    "missing_fields": " | ".join(missing_fields)
                })

        return missing_rows
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Audit Preview Failed: {str(e)}")